"""图谱导入文档解析组件。"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Protocol

logger = logging.getLogger(__name__)

SUPPORTED_GRAPH_IMPORT_SUFFIXES = frozenset(
    {
        ".txt",
        ".md",
        ".pdf",
        ".docx",
        ".doc",
        ".pptx",
        ".ppt",
        ".html",
        ".htm",
        ".jpg",
        ".jpeg",
        ".png",
    }
)
MAX_GRAPH_IMPORT_FILE_BYTES = 50 * 1024 * 1024


class GraphImportCachePort(Protocol):
    """解析组件需要的缓存端口。"""

    def _compute_file_hash(self, file_bytes: bytes) -> str: ...

    def _load_graph_import_cache(
        self,
        file_hash: str,
    ) -> tuple[dict[str, Any] | None, str | None]: ...

    def _write_graph_import_cache(
        self,
        file_hash: str,
        prepared: dict[str, Any],
        chunks: list[str],
    ) -> None: ...

    def _get_chunk_strategy_signature(self) -> str: ...


class GraphImportParser:
    """负责图谱导入文件校验、正文解析、分块和缓存读写。"""

    def __init__(
        self,
        *,
        cache: GraphImportCachePort,
        chunk_limit: int,
        chunk_strategy_version: int,
        log_api_error: Callable[..., None],
    ) -> None:
        self._cache = cache
        self._chunk_limit = chunk_limit
        self._chunk_strategy_version = chunk_strategy_version
        self._log_api_error = log_api_error

    @staticmethod
    def split_text_by_sentence(text: str, limit: int) -> list[str]:
        """优先按句末标点切分文本，超长句按限制硬切。"""
        chunks: list[str] = []
        remaining = text.strip()
        sentence_ends = ("。", ".", "!", "！", "?", "？")
        while remaining:
            if len(remaining) <= limit:
                chunks.append(remaining)
                break
            window = remaining[:limit]
            split_at = max(window.rfind(mark) for mark in sentence_ends)
            split_at = limit if split_at < 0 else split_at + 1
            chunk = remaining[:split_at].strip()
            if not chunk:
                chunk = remaining[:limit]
                split_at = limit
            chunks.append(chunk)
            remaining = remaining[split_at:].strip()
        return chunks

    @staticmethod
    def _extract_text_fallback(raw_bytes: bytes, suffix: str) -> str:
        """Docling 不可用时按文件类型执行内置文本解析。"""
        if suffix == ".pdf":
            import fitz  # type: ignore[import-untyped]

            document = fitz.open(stream=raw_bytes, filetype="pdf")
            try:
                return "\n".join(page.get_text("text") for page in document)
            finally:
                document.close()
        if suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(raw_bytes))
            return "\n\n".join(
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text and paragraph.text.strip()
            )
        if suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            html = raw_bytes.decode("utf-8", errors="replace")
            return BeautifulSoup(html, "html.parser").get_text(
                separator="\n",
                strip=True,
            )
        raise RuntimeError(
            f"当前环境无法使用 Docling，且该格式无内置回退解析器: {suffix}"
        )

    def parse_and_split(
        self,
        file_bytes: bytes,
        filename: str,
        stage_callback: Callable[..., None],
    ) -> dict[str, Any]:
        """执行文件校验、正文解析、缓存恢复和文本分块。"""
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_GRAPH_IMPORT_SUFFIXES:
            supported = ", ".join(sorted(SUPPORTED_GRAPH_IMPORT_SUFFIXES))
            raise RuntimeError(f"不支持的文件格式: {suffix}，仅支持 {supported}")
        if not file_bytes:
            raise RuntimeError("上传文件为空")
        if len(file_bytes) > MAX_GRAPH_IMPORT_FILE_BYTES:
            raise RuntimeError("文件大小不能超过 50MB")

        stage_callback("enter", "validate_file", "正在校验文件")
        stage_callback("complete", "validate_file", "文件校验完成")
        file_hash = self._cache._compute_file_hash(file_bytes)
        logger.info(
            "graph_import.parse_prepare file_hash=%s chunk_limit=%s "
            "chunk_strategy_version=%s file_name=%s",
            file_hash,
            self._chunk_limit,
            self._chunk_strategy_version,
            filename,
        )

        stage_callback("enter", "parse_document", "正在检查本地缓存")
        cached_prepared, cache_invalidated_reason = (
            self._cache._load_graph_import_cache(file_hash)
        )
        if cached_prepared:
            stage_callback(
                "complete",
                "parse_document",
                f"已从本地缓存恢复解析结果，共 {cached_prepared.get('chunk_count') or 0} 段文本",
            )
            logger.info(
                "graph_import.cache_hit file_hash=%s chunk_limit=%s "
                "chunk_count=%s first_chunk_chars=%s",
                file_hash,
                cached_prepared.get("chunk_limit"),
                cached_prepared.get("chunk_count"),
                (
                    len((cached_prepared.get("_chunks") or [""])[0])
                    if cached_prepared.get("_chunks")
                    else 0
                ),
            )
            return cached_prepared
        if cache_invalidated_reason:
            logger.info(
                "graph_import.cache_invalidated file_hash=%s chunk_limit=%s reason=%s",
                file_hash,
                self._chunk_limit,
                cache_invalidated_reason,
            )

        parser_used = "plain_text"
        parse_warning = ""
        parsed: Any = None
        text = ""
        if suffix in {".txt", ".md"}:
            stage_callback("heartbeat", "parse_document", "正在解析纯文本内容")
            text = file_bytes.decode("utf-8", errors="replace")
            stage_callback("complete", "parse_document", "文档正文解析完成")
        else:
            try:
                from app.document_parser.docling_parser import DoclingParser

                parser = DoclingParser(
                    enable_image_ocr=True,
                    strict_image_ocr=True,
                    progress_callback=stage_callback,
                )
                parsed = parser.parse_bytes(file_bytes, filename)
                text = parsed.text
                parser_used = "docling"
            except (
                RuntimeError,
                ValueError,
                TypeError,
                KeyError,
                AttributeError,
                OSError,
                ImportError,
            ) as docling_error:
                if "图片 OCR 失败" in str(docling_error):
                    raise
                self._log_api_error(
                    "/admin/graph/jobs/extract-import",
                    "docling_parse_fallback",
                    docling_error,
                    filename=filename,
                )
                stage_callback(
                    "heartbeat",
                    "parse_document",
                    "Docling 不可用，正在执行回退解析",
                )
                text = self._extract_text_fallback(file_bytes, suffix)
                stage_callback("complete", "parse_document", "回退解析完成")
                parser_used = "fallback"
                parse_warning = f"Docling 不可用，已自动回退解析: {docling_error}"

        if not text.strip():
            raise RuntimeError("文件解析后文本为空，无法抽取图谱")
        if parse_warning:
            stage_callback("warning", detail=parse_warning)

        chunks = self.split_text_by_sentence(text, self._chunk_limit)
        total_chunk_chars = sum(len(chunk) for chunk in chunks)
        ocr_engine = None
        docling_failed_pages: list[int] = []
        pdf_fallback_used = False
        if parsed and parser_used == "docling":
            metadata = parsed.metadata if isinstance(parsed.metadata, dict) else {}
            ocr_engine = metadata.get("ocr_engine")
            docling_failed_pages = list(metadata.get("docling_failed_pages") or [])
            pdf_fallback_used = bool(metadata.get("pdf_fallback_used"))
            if pdf_fallback_used:
                parser_used = "docling_with_pymupdf_fallback"
                if not parse_warning:
                    parse_warning = str(metadata.get("pdf_fallback_reason") or "")

        response_data: dict[str, Any] = {
            "file_name": filename,
            "parsed_text_length": len(text),
            "total_chunk_chars": total_chunk_chars,
            "processed_text_length": 0,
            "chunk_limit": self._chunk_limit,
            "chunk_strategy_version": self._chunk_strategy_version,
            "parser_used": parser_used,
            "ocr_engine": ocr_engine,
            "parse_warning": parse_warning,
            "docling_failed_pages": docling_failed_pages,
            "pdf_fallback_used": pdf_fallback_used,
            "chunk_count": len(chunks),
            "entity_count": 0,
            "relation_count": 0,
            "node_batches": 0,
            "relation_batches": 0,
            "message": f"文档解析完成，共切分为 {len(chunks)} 段文本",
            "cache_invalidated_reason": cache_invalidated_reason or "",
            "_chunks": chunks,
            "cache_used": False,
        }
        if parsed:
            response_data["parse_details"] = {
                "format": parsed.format,
                "pages": parsed.pages,
                "tables_count": len(parsed.tables),
                "images_count": len(parsed.images),
                "metadata": parsed.metadata,
                "text_preview": parsed.text[:2000] if parsed.text else "",
            }

        logger.info(
            "graph_import.chunked file_hash=%s chunk_limit=%s chunk_count=%s "
            "first_chunk_chars=%s current_strategy=%s",
            file_hash,
            self._chunk_limit,
            len(chunks),
            len(chunks[0]) if chunks else 0,
            self._cache._get_chunk_strategy_signature(),
        )
        try:
            self._cache._write_graph_import_cache(file_hash, response_data, chunks)
        except OSError as exc:
            logger.warning(
                "写入图谱导入缓存失败: file_hash=%s error=%s",
                file_hash,
                exc,
            )
        return response_data
