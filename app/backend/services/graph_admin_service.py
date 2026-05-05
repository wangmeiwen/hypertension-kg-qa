"""管理员图谱服务。"""

from __future__ import annotations

import asyncio
import gzip
import hashlib
import json
import logging
import traceback
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Awaitable, Callable, Optional
from uuid import uuid4

from fastapi import UploadFile

from app.common.concurrency import concurrency_guard
from app.common.config import Config
from app.backend.repositories.database import execute, fetch_all, fetch_one

logger = logging.getLogger(__name__)


class GraphAdminService:
    """封装管理员图谱管理与导入嵌入业务。"""

    JOB_STATUS_QUEUED = "queued"
    JOB_STATUS_RUNNING = "running"
    JOB_STATUS_PAUSED = "paused"
    JOB_STATUS_SUCCEEDED = "succeeded"
    JOB_STATUS_FAILED = "failed"
    JOB_STATUS_CANCELLED = "cancelled"

    IMPORT_JOB_TYPE = "graph_extract_import"
    EMBED_JOB_TYPE = "graph_embed_milvus"
    SOFT_TIMEOUT_SECONDS = 300
    SOFT_TIMEOUT_REPEAT_SECONDS = 60
    GRAPH_IMPORT_CACHE_VERSION = 2
    GRAPH_IMPORT_CHUNK_LIMIT = 10000
    GRAPH_IMPORT_CHUNK_STRATEGY_VERSION = 2
    GRAPH_LLM_REQUEST_TIMEOUT_SECONDS = 600
    GRAPH_LLM_MAX_ATTEMPTS = 3
    GRAPH_LLM_MAX_PARALLELISM = 3
    GRAPH_LLM_TRANSIENT_BACKOFF_SECONDS = 15

    def __init__(
        self,
        await_call: Callable[..., Awaitable[Any]],
        neo4j_client: Any,
        get_shared_kg_milvus_store: Callable[[], Any],
        validate_cypher_identifier: Callable[[Any], Optional[str]],
        sanitize_graph_properties: Callable[[Any], dict[str, Any]],
        log_api_error: Callable[..., None],
    ) -> None:
        """初始化实例并注入运行依赖。

        Args:
            await_call: 兼容同步/异步函数的调用适配器。
            neo4j_client: Neo4j 客户端实例，用于执行图数据库查询和写入。
            get_shared_kg_milvus_store: 获取图谱向量库实例的回调。
            validate_cypher_identifier: 校验 Cypher label/type/property 标识符的回调。
            sanitize_graph_properties: 清洗图谱节点和关系属性的回调。
            log_api_error: 统一记录管理员图谱接口错误的回调。

        Returns:
            None: 无返回值。
        """
        self._await_call = await_call
        self._neo4j_client = neo4j_client
        self._get_shared_kg_milvus_store = get_shared_kg_milvus_store
        self._validate_cypher_identifier = validate_cypher_identifier
        self._sanitize_graph_properties = sanitize_graph_properties
        self._log_api_error = log_api_error
        self._config = Config()
        self._graph_llm_lane_configs = list(self._config.GRAPH_LLM_LANE_CONFIGS)
        self._jobs: dict[str, dict[str, Any]] = {}
        self._job_lock = asyncio.Lock()
        self._max_retained_jobs = 100

    @staticmethod
    def _utc_now_iso() -> str:
        """返回统一的 UTC ISO 时间文本。"""
        return datetime.utcnow().isoformat(timespec="seconds") + "Z"

    @classmethod
    def _build_job_response(cls, job: dict[str, Any]) -> dict[str, Any]:
        """裁剪对外 job 响应。"""
        return {
            "job_id": job["job_id"],
            "job_type": job["job_type"],
            "status": job["status"],
            "labels": job.get("labels"),
            "batch_size": job.get("batch_size"),
            "progress_text": job.get("progress_text") or "",
            "created_at": job.get("created_at"),
            "started_at": job.get("started_at"),
            "finished_at": job.get("finished_at"),
            "stage": job.get("stage"),
            "stage_detail": job.get("stage_detail") or "",
            "stage_started_at": job.get("stage_started_at"),
            "last_heartbeat_at": job.get("last_heartbeat_at"),
            "total_chars": job.get("total_chars"),
            "processed_chars": job.get("processed_chars"),
            "total_chunks": job.get("total_chunks"),
            "completed_chunks": job.get("completed_chunks"),
            "current_chunk_index": job.get("current_chunk_index"),
            "current_chunk_chars": job.get("current_chunk_chars"),
            "file_hash": job.get("file_hash"),
            "run_id": job.get("run_id"),
            "is_resumed": bool(job.get("is_resumed")),
            "resume_run_id": job.get("resume_run_id"),
            "resumed_from_chunk": job.get("resumed_from_chunk"),
            "chunk_limit": job.get("chunk_limit"),
            "chunk_strategy_version": job.get("chunk_strategy_version"),
            "cache_used": bool(job.get("cache_used")),
            "cache_invalidated_reason": job.get("cache_invalidated_reason"),
            "resume_compatible": job.get("resume_compatible"),
            "parallelism": int(job.get("parallelism") or 0),
            "active_lanes": int(job.get("active_lanes") or 0),
            "current_attempt": int(job.get("current_attempt") or 0),
            "max_attempts": int(job.get("max_attempts") or cls.GRAPH_LLM_MAX_ATTEMPTS),
            "lane_statuses": cls._clone_lane_statuses(job.get("lane_statuses") or []),
            "warnings": list(job.get("warnings") or []),
            "trace": list(job.get("trace") or []),
            "result": job.get("result"),
            "error": job.get("error"),
            "deduplicated": bool(job.get("deduplicated")),
            "existing_job_id": job.get("existing_job_id"),
            "existing_run_id": job.get("existing_run_id"),
        }

    @staticmethod
    def _parse_utc_iso(value: Optional[str]) -> Optional[datetime]:
        """将统一的 UTC ISO 时间文本解析为 datetime。"""
        if not value:
            return None
        try:
            return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            return None

    @staticmethod
    def _compute_file_hash(file_bytes: bytes) -> str:
        """计算文件内容哈希。"""
        return hashlib.sha256(file_bytes).hexdigest()

    @classmethod
    def _get_chunk_strategy_signature(cls) -> str:
        """返回当前图谱导入分块策略签名。"""
        return f"v{cls.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION}:limit={cls.GRAPH_IMPORT_CHUNK_LIMIT}"

    @classmethod
    def _get_graph_import_cache_dir(cls) -> Path:
        """返回图谱导入本地缓存目录。"""
        return Path(__file__).resolve().parents[3] / "runtime_cache" / "graph_import"

    @classmethod
    def _get_graph_import_cache_path(cls, file_hash: str) -> Path:
        """返回图谱导入缓存文件路径。"""
        return cls._get_graph_import_cache_dir() / f"{file_hash}.json.gz"

    @staticmethod
    def _json_dumps(value: Any) -> str:
        """以稳定格式序列化 JSON 文本。"""
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

    @staticmethod
    def _json_loads(value: Any, default: Any) -> Any:
        """安全反序列化 JSON 文本。"""
        if value in (None, ""):
            return default
        try:
            return json.loads(str(value))
        except (TypeError, ValueError, json.JSONDecodeError):
            return default

    def _load_graph_import_cache(self, file_hash: str) -> tuple[Optional[dict[str, Any]], Optional[str]]:
        """读取本地图谱解析缓存，并返回失效原因。"""
        cache_path = self._get_graph_import_cache_path(file_hash)
        if not cache_path.exists():
            return None, None
        try:
            with concurrency_guard.sync_keyed_lock(f"graph-import-cache:{file_hash}"):
                with gzip.open(cache_path, "rt", encoding="utf-8") as handle:
                    payload = json.load(handle)
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            logger.warning("读取图谱导入缓存失败，将忽略缓存: path=%s error=%s", cache_path, exc)
            return None, f"本地缓存读取失败，已忽略: {exc}"
        if int(payload.get("version") or 0) != self.GRAPH_IMPORT_CACHE_VERSION:
            return None, "本地缓存版本已过期，已重新解析"
        if int(payload.get("chunk_limit") or 0) != self.GRAPH_IMPORT_CHUNK_LIMIT:
            return None, "本地缓存分块大小与当前策略不一致，已重新分块"
        if int(payload.get("chunk_strategy_version") or 0) != self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION:
            return None, "本地缓存分块策略已过期，已重新解析"
        if str(payload.get("file_hash") or "") != file_hash:
            return None, "本地缓存文件哈希不匹配，已重新解析"
        prepared = payload.get("prepared")
        chunks = payload.get("chunks")
        if not isinstance(prepared, dict) or not isinstance(chunks, list):
            return None, "本地缓存内容不完整，已重新解析"
        prepared = dict(prepared)
        prepared["_chunks"] = [str(chunk or "") for chunk in chunks]
        prepared["cache_used"] = True
        prepared["cache_path"] = str(cache_path)
        prepared["cache_invalidated_reason"] = None
        prepared["chunk_limit"] = int(payload.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT)
        prepared["chunk_strategy_version"] = int(
            payload.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION
        )
        return prepared, None

    def _write_graph_import_cache(self, file_hash: str, prepared: dict[str, Any], chunks: list[str]) -> None:
        """写入本地图谱解析缓存。"""
        cache_dir = self._get_graph_import_cache_dir()
        cache_dir.mkdir(parents=True, exist_ok=True)
        cache_path = self._get_graph_import_cache_path(file_hash)
        payload = {
            "version": self.GRAPH_IMPORT_CACHE_VERSION,
            "file_hash": file_hash,
            "chunk_limit": self.GRAPH_IMPORT_CHUNK_LIMIT,
            "chunk_strategy_version": self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
            "cached_at": self._utc_now_iso(),
            "prepared": {
                "file_name": prepared.get("file_name"),
                "parsed_text_length": prepared.get("parsed_text_length") or 0,
                "total_chunk_chars": prepared.get("total_chunk_chars") or 0,
                "processed_text_length": 0,
                "parser_used": prepared.get("parser_used"),
                "ocr_engine": prepared.get("ocr_engine"),
                "parse_warning": prepared.get("parse_warning") or "",
                "docling_failed_pages": list(prepared.get("docling_failed_pages") or []),
                "pdf_fallback_used": bool(prepared.get("pdf_fallback_used")),
                "chunk_count": prepared.get("chunk_count") or 0,
                "entity_count": 0,
                "relation_count": 0,
                "node_batches": 0,
                "relation_batches": 0,
                "message": prepared.get("message") or "",
                "chunk_limit": prepared.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT,
                "chunk_strategy_version": prepared.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
            },
            "chunks": chunks,
        }
        tmp_path = cache_path.with_name(f"{cache_path.name}.{uuid4().hex}.tmp")
        with concurrency_guard.sync_keyed_lock(f"graph-import-cache:{file_hash}"):
            try:
                with gzip.open(tmp_path, "wt", encoding="utf-8") as handle:
                    json.dump(payload, handle, ensure_ascii=False)
                tmp_path.replace(cache_path)
            finally:
                try:
                    if tmp_path.exists():
                        tmp_path.unlink()
                except OSError:
                    pass

    @staticmethod
    def _normalize_chunk_status(status: Any) -> str:
        """将未知状态收敛到块级状态枚举。"""
        status_text = str(status or "").strip().lower()
        if status_text in {"pending", "running", "succeeded", "failed"}:
            return status_text
        return "pending"

    @staticmethod
    def _format_resume_message(chunk_index: Optional[int]) -> str:
        """生成自动续跑提示文案。"""
        if chunk_index and chunk_index > 0:
            return f"检测到未完成导入，正在从第 {chunk_index} 块继续"
        return "检测到未完成导入，正在自动继续执行"

    @staticmethod
    def _first_error_line(error_text: Optional[str]) -> str:
        """提取异常首行，避免前端展示整段堆栈。"""
        text = str(error_text or "").strip()
        if not text:
            return ""
        return next((line.strip() for line in text.splitlines() if line.strip()), "")

    @classmethod
    def _summarize_pause_error(cls, error_text: Optional[str]) -> str:
        """将块级中断异常归一化为适合前端展示的暂停原因。"""
        raw_text = str(error_text or "").strip()
        if not raw_text:
            return "大模型调用中断"
        first_line = cls._first_error_line(raw_text) or raw_text
        lowered = raw_text.lower()
        if "401 unauthorized" in lowered:
            if "dashscope" in lowered:
                return "DashScope 接口返回 401 Unauthorized，请检查 API Key、模型权限或免费额度。"
            if "open.bigmodel.cn" in lowered:
                return "GLM 接口返回 401 Unauthorized，请检查 API Key 或模型权限。"
            return "模型接口返回 401 Unauthorized，请检查 API Key 或模型权限。"
        if "404 not found" in lowered:
            if "dashscope" in lowered:
                return "DashScope 接口返回 404 Not Found，请检查模型名或兼容接口地址配置。"
            if "open.bigmodel.cn" in lowered:
                return "GLM 接口返回 404 Not Found，请检查模型名或接口地址配置。"
            return "模型接口返回 404 Not Found，请检查模型名或接口地址配置。"
        if "timeout" in lowered or "timed out" in lowered:
            return "模型接口调用超时，当前块未完成，重新提交同文件可继续。"
        if "rate limit" in lowered or "too many requests" in lowered or "429" in lowered:
            return "模型接口触发限流，当前块未完成，重新提交同文件可继续。"
        return first_line

    @staticmethod
    def _is_timeout_error(error: BaseException | str | None) -> bool:
        """判断异常是否属于超时类错误。"""
        error_text = str(error or "").strip().lower()
        if not error_text:
            return False
        timeout_markers = (
            "timeout",
            "timed out",
            "readtimeout",
            "read timeout",
            "connecttimeout",
            "connect timeout",
            "apitimeouterror",
            "request timed out",
        )
        return any(marker in error_text for marker in timeout_markers)

    @staticmethod
    def _classify_lane_error(error: BaseException | str | None) -> str:
        """将 lane 错误归类为 permanent 或 transient。"""
        error_text = str(error or "").strip().lower()
        if not error_text:
            return "transient"
        permanent_markers = (
            "401 unauthorized",
            "403 forbidden",
            "404 not found",
        )
        if any(marker in error_text for marker in permanent_markers):
            return "permanent"
        transient_markers = (
            "429",
            "rate limit",
            "too many requests",
            "timeout",
            "timed out",
            "readtimeout",
            "connecttimeout",
            "connection reset",
            "connection aborted",
            "connection error",
            "api connection error",
            "service unavailable",
            "server disconnected",
            "temporarily unavailable",
            "502",
            "503",
            "504",
            "network",
        )
        if any(marker in error_text for marker in transient_markers):
            return "transient"
        return "transient"

    @staticmethod
    def _clone_lane_statuses(lane_statuses: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """复制 lane 状态用于写回任务响应。"""
        return [dict(item) for item in lane_statuses]

    @classmethod
    def _count_active_lanes(cls, lane_statuses: list[dict[str, Any]]) -> int:
        """统计当前可立即提供服务的 lane 数量。"""
        return sum(1 for item in lane_statuses if item.get("status") in {"ready", "running"})

    @classmethod
    def _count_viable_lanes(cls, lane_statuses: list[dict[str, Any]]) -> int:
        """统计当前仍可恢复或继续提供服务的 lane 数量。"""
        return sum(1 for item in lane_statuses if item.get("status") in {"ready", "running", "backoff"})

    @staticmethod
    def _format_lane_label(lane_config: dict[str, Any]) -> str:
        """格式化 lane 展示标签。"""
        lane_id = lane_config.get("lane_id")
        model_name = str(lane_config.get("model_name") or "")
        source = str(lane_config.get("source") or "")
        if source:
            return f"lane{lane_id}:{model_name} ({source})"
        return f"lane{lane_id}:{model_name}"

    @staticmethod
    def _build_paused_job_result(
        *,
        message: str,
        pause_reason: str,
        file_name: str,
        file_hash: str,
        run_id: str,
        resumed_from_chunk: Optional[int],
        prepared: dict[str, Any],
        entity_count: int,
        relation_count: int,
        processed_chars: int,
        completed_chunks: int,
    ) -> dict[str, Any]:
        """生成暂停态任务结果。

        Args:
            message: 对外展示的暂停提示。
            pause_reason: 暂停原因。
            file_name: 导入文件名。
            file_hash: 文件内容哈希。
            run_id: 持久化导入 run 标识。
            resumed_from_chunk: 本次续跑起始块序号。
            prepared: 文件解析和分块元数据。
            entity_count: 已抽取实体数量。
            relation_count: 已抽取关系数量。
            processed_chars: 已处理字符数。
            completed_chunks: 已完成块数。

        Returns:
            dict[str, Any]: 可写入 job.result 的暂停态结果。
        """
        return {
            "_job_status": "paused",
            "message": message,
            "file_name": file_name,
            "file_hash": file_hash,
            "run_id": run_id,
            "is_resumed": bool(prepared.get("is_resumed")),
            "resume_run_id": run_id if prepared.get("is_resumed") else None,
            "resumed_from_chunk": resumed_from_chunk,
            "chunk_limit": prepared.get("chunk_limit"),
            "chunk_strategy_version": prepared.get("chunk_strategy_version"),
            "parsed_text_length": prepared.get("parsed_text_length") or 0,
            "processed_text_length": processed_chars,
            "total_chunk_chars": prepared.get("total_chunk_chars") or 0,
            "chunk_count": prepared.get("chunk_count") or 0,
            "completed_chunks": completed_chunks,
            "entity_count": entity_count,
            "relation_count": relation_count,
            "parser_used": prepared.get("parser_used"),
            "ocr_engine": prepared.get("ocr_engine"),
            "parse_warning": prepared.get("parse_warning"),
            "docling_failed_pages": prepared.get("docling_failed_pages") or [],
            "pdf_fallback_used": bool(prepared.get("pdf_fallback_used")),
            "cache_used": bool(prepared.get("cache_used")),
            "cache_invalidated_reason": prepared.get("cache_invalidated_reason") or "",
            "resume_compatible": prepared.get("resume_compatible"),
            "pause_reason": pause_reason,
            "last_error": pause_reason,
        }

    async def _find_resumable_graph_run(  # pragma: no cover - graph import persistence adapter
        self,
        file_hash: str,
        *,
        chunk_limit: int,
        chunk_strategy_version: int,
    ) -> Optional[dict[str, Any]]:
        """查找同文件且分块策略兼容的未完成导入。"""
        row = await fetch_one(
            """
            SELECT *
            FROM graph_import_runs
            WHERE file_hash = :file_hash
              AND chunk_limit = :chunk_limit
              AND chunk_strategy_version = :chunk_strategy_version
              AND status IN ('running', 'paused')
            ORDER BY updated_at DESC
            LIMIT 1
            """,
            {
                "file_hash": file_hash,
                "chunk_limit": chunk_limit,
                "chunk_strategy_version": chunk_strategy_version,
            },
        )
        if not row:
            return None
        row["docling_failed_pages"] = self._json_loads(row.get("docling_failed_pages"), [])
        return row

    async def _find_latest_graph_run_by_file_hash(self, file_hash: str) -> Optional[dict[str, Any]]:  # pragma: no cover - graph import persistence adapter
        """查找同文件最近一次导入记录，不要求策略兼容。"""
        row = await fetch_one(
            """
            SELECT *
            FROM graph_import_runs
            WHERE file_hash = :file_hash
            ORDER BY updated_at DESC
            LIMIT 1
            """,
            {"file_hash": file_hash},
        )
        if not row:
            return None
        row["docling_failed_pages"] = self._json_loads(row.get("docling_failed_pages"), [])
        return row

    async def _find_first_unfinished_chunk_index(self, run_id: str) -> Optional[int]:  # pragma: no cover - graph import persistence adapter
        """返回 run 中第一个未完成块。"""
        row = await fetch_one(
            """
            SELECT chunk_index
            FROM graph_import_chunks
            WHERE run_id = :run_id
              AND status <> 'succeeded'
            ORDER BY chunk_index ASC
            LIMIT 1
            """,
            {"run_id": run_id},
        )
        if not row:
            return None
        try:
            return int(row.get("chunk_index") or 0)
        except (TypeError, ValueError):
            return None

    async def _find_active_job_by_run_id(self, run_id: str) -> Optional[dict[str, Any]]:  # pragma: no cover - graph import persistence adapter
        """查找当前进程内是否已有同 run 的活跃任务。"""
        async with self._job_lock:
            for job in self._jobs.values():
                if job.get("run_id") == run_id and job.get("status") in {
                    self.JOB_STATUS_QUEUED,
                    self.JOB_STATUS_RUNNING,
                }:
                    return self._build_job_response(job)
        return None

    async def _create_graph_import_run(  # pragma: no cover - graph import persistence adapter
        self,
        *,
        run_id: str,
        file_hash: str,
        file_name: str,
        file_size: int,
        prepared: dict[str, Any],
    ) -> None:
        """创建新的图谱导入 run。"""
        await execute(
            """
            INSERT INTO graph_import_runs (
                run_id, file_hash, file_name, file_size, status, parser_used,
                parse_warning, parsed_text_length, total_chars, processed_chars,
                total_chunks, completed_chunks, entity_count, relation_count,
                docling_failed_pages, pdf_fallback_used, chunk_limit, chunk_strategy_version, last_error
            )
            VALUES (
                :run_id, :file_hash, :file_name, :file_size, 'running', :parser_used,
                :parse_warning, :parsed_text_length, :total_chars, 0,
                :total_chunks, 0, 0, 0, :docling_failed_pages, :pdf_fallback_used, :chunk_limit, :chunk_strategy_version, NULL
            )
            """,
            {
                "run_id": run_id,
                "file_hash": file_hash,
                "file_name": file_name,
                "file_size": file_size,
                "parser_used": prepared.get("parser_used") or "",
                "parse_warning": prepared.get("parse_warning") or "",
                "parsed_text_length": prepared.get("parsed_text_length") or 0,
                "total_chars": prepared.get("parsed_text_length") or 0,
                "total_chunks": prepared.get("chunk_count") or 0,
                "docling_failed_pages": self._json_dumps(prepared.get("docling_failed_pages") or []),
                "pdf_fallback_used": 1 if prepared.get("pdf_fallback_used") else 0,
                "chunk_limit": prepared.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT,
                "chunk_strategy_version": prepared.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
            },
        )

    async def _initialize_graph_import_chunks(self, run_id: str, chunks: list[str]) -> None:  # pragma: no cover - graph import persistence adapter
        """初始化块级断点记录。"""
        for index, chunk in enumerate(chunks, start=1):
            await execute(
                """
                INSERT INTO graph_import_chunks (
                    run_id, chunk_index, chunk_chars, status,
                    entity_count, relation_count, node_batches, relation_batches, error_message
                )
                VALUES (
                    :run_id, :chunk_index, :chunk_chars, 'pending',
                    0, 0, 0, 0, NULL
                )
                ON DUPLICATE KEY UPDATE chunk_chars = VALUES(chunk_chars)
                """,
                {
                    "run_id": run_id,
                    "chunk_index": index,
                    "chunk_chars": len(chunk),
                },
            )

    async def _sync_graph_import_run_metadata(  # pragma: no cover - graph import persistence adapter
        self,
        run_id: str,
        *,
        file_name: str,
        file_size: int,
        prepared: dict[str, Any],
        status: str = "running",
        clear_error: bool = False,
    ) -> None:
        """同步 run 的解析元数据。"""
        await execute(
            """
            UPDATE graph_import_runs
            SET file_name = :file_name,
                file_size = :file_size,
                status = :status,
                parser_used = :parser_used,
                parse_warning = :parse_warning,
                parsed_text_length = :parsed_text_length,
                total_chars = :total_chars,
                total_chunks = :total_chunks,
                docling_failed_pages = :docling_failed_pages,
                pdf_fallback_used = :pdf_fallback_used,
                chunk_limit = :chunk_limit,
                chunk_strategy_version = :chunk_strategy_version,
                last_error = CASE WHEN :clear_error = 1 THEN NULL ELSE last_error END
            WHERE run_id = :run_id
            """,
            {
                "run_id": run_id,
                "file_name": file_name,
                "file_size": file_size,
                "status": status,
                "parser_used": prepared.get("parser_used") or "",
                "parse_warning": prepared.get("parse_warning") or "",
                "parsed_text_length": prepared.get("parsed_text_length") or 0,
                "total_chars": prepared.get("parsed_text_length") or 0,
                "total_chunks": prepared.get("chunk_count") or 0,
                "docling_failed_pages": self._json_dumps(prepared.get("docling_failed_pages") or []),
                "pdf_fallback_used": 1 if prepared.get("pdf_fallback_used") else 0,
                "chunk_limit": prepared.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT,
                "chunk_strategy_version": prepared.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
                "clear_error": 1 if clear_error else 0,
            },
        )

    async def _load_graph_import_chunks(self, run_id: str) -> list[dict[str, Any]]:  # pragma: no cover - graph import persistence adapter
        """加载块级断点记录。"""
        rows = await fetch_all(
            """
            SELECT *
            FROM graph_import_chunks
            WHERE run_id = :run_id
            ORDER BY chunk_index ASC
            """,
            {"run_id": run_id},
        )
        for row in rows:
            row["status"] = self._normalize_chunk_status(row.get("status"))
            for key in ("chunk_index", "chunk_chars", "entity_count", "relation_count", "node_batches", "relation_batches"):
                try:
                    row[key] = int(row.get(key) or 0)
                except (TypeError, ValueError):
                    row[key] = 0
        return rows

    async def _recompute_graph_run_progress(self, run_id: str) -> dict[str, int]:  # pragma: no cover - graph import persistence adapter
        """根据块级状态重算 run 进度。"""
        row = await fetch_one(
            """
            SELECT
                COALESCE(SUM(CASE WHEN status = 'succeeded' THEN chunk_chars ELSE 0 END), 0) AS processed_chars,
                COALESCE(SUM(CASE WHEN status = 'succeeded' THEN 1 ELSE 0 END), 0) AS completed_chunks,
                COALESCE(SUM(CASE WHEN status = 'succeeded' THEN entity_count ELSE 0 END), 0) AS entity_count,
                COALESCE(SUM(CASE WHEN status = 'succeeded' THEN relation_count ELSE 0 END), 0) AS relation_count,
                COALESCE(SUM(CASE WHEN status = 'succeeded' THEN node_batches ELSE 0 END), 0) AS node_batches,
                COALESCE(SUM(CASE WHEN status = 'succeeded' THEN relation_batches ELSE 0 END), 0) AS relation_batches
            FROM graph_import_chunks
            WHERE run_id = :run_id
            """,
            {"run_id": run_id},
        ) or {}
        progress = {
            "processed_chars": int(row.get("processed_chars") or 0),
            "completed_chunks": int(row.get("completed_chunks") or 0),
            "entity_count": int(row.get("entity_count") or 0),
            "relation_count": int(row.get("relation_count") or 0),
            "node_batches": int(row.get("node_batches") or 0),
            "relation_batches": int(row.get("relation_batches") or 0),
        }
        await execute(
            """
            UPDATE graph_import_runs
            SET processed_chars = :processed_chars,
                completed_chunks = :completed_chunks,
                entity_count = :entity_count,
                relation_count = :relation_count
            WHERE run_id = :run_id
            """,
            {"run_id": run_id, **progress},
        )
        return progress

    async def _mark_graph_chunk_running(self, run_id: str, chunk_index: int) -> None:  # pragma: no cover - graph import persistence adapter
        """标记块开始执行。"""
        await execute(
            """
            UPDATE graph_import_chunks
            SET status = 'running',
                started_at = CURRENT_TIMESTAMP,
                finished_at = NULL,
                error_message = NULL
            WHERE run_id = :run_id AND chunk_index = :chunk_index
            """,
            {"run_id": run_id, "chunk_index": chunk_index},
        )

    async def _mark_graph_chunk_succeeded(  # pragma: no cover - graph import persistence adapter
        self,
        run_id: str,
        *,
        chunk_index: int,
        entity_count: int,
        relation_count: int,
        node_batches: int,
        relation_batches: int,
    ) -> None:
        """标记块执行成功。"""
        await execute(
            """
            UPDATE graph_import_chunks
            SET status = 'succeeded',
                entity_count = :entity_count,
                relation_count = :relation_count,
                node_batches = :node_batches,
                relation_batches = :relation_batches,
                error_message = NULL,
                finished_at = CURRENT_TIMESTAMP
            WHERE run_id = :run_id AND chunk_index = :chunk_index
            """,
            {
                "run_id": run_id,
                "chunk_index": chunk_index,
                "entity_count": entity_count,
                "relation_count": relation_count,
                "node_batches": node_batches,
                "relation_batches": relation_batches,
            },
        )

    async def _mark_graph_chunk_failed(self, run_id: str, chunk_index: int, error_message: str) -> None:  # pragma: no cover - graph import persistence adapter
        """标记块执行失败。"""
        await execute(
            """
            UPDATE graph_import_chunks
            SET status = 'failed',
                error_message = :error_message,
                finished_at = CURRENT_TIMESTAMP
            WHERE run_id = :run_id AND chunk_index = :chunk_index
            """,
            {
                "run_id": run_id,
                "chunk_index": chunk_index,
                "error_message": error_message[:4000],
            },
        )

    async def _set_graph_run_status(  # pragma: no cover - graph import persistence adapter
        self,
        run_id: str,
        status: str,
        *,
        last_error: Optional[str] = None,
        finished: bool = False,
    ) -> None:
        """更新 run 状态。"""
        await execute(
            """
            UPDATE graph_import_runs
            SET status = :status,
                last_error = :last_error,
                finished_at = CASE WHEN :finished = 1 THEN CURRENT_TIMESTAMP ELSE finished_at END
            WHERE run_id = :run_id
            """,
            {
                "run_id": run_id,
                "status": status,
                "last_error": last_error[:4000] if last_error else None,
                "finished": 1 if finished else 0,
            },
        )

    async def _update_job_resume_state(  # pragma: no cover - graph import persistence adapter
        self,
        job_id: str,
        *,
        run_id: str,
        file_hash: str,
        is_resumed: bool,
        resumed_from_chunk: Optional[int],
        chunk_limit: Optional[int] = None,
        chunk_strategy_version: Optional[int] = None,
        cache_used: Optional[bool] = None,
        cache_invalidated_reason: Optional[str] = None,
        resume_compatible: Optional[bool] = None,
    ) -> None:
        """把恢复元数据写回进程内任务。"""
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                return
            job["run_id"] = run_id
            job["file_hash"] = file_hash
            job["is_resumed"] = is_resumed
            job["resume_run_id"] = run_id if is_resumed else None
            job["resumed_from_chunk"] = resumed_from_chunk
            if chunk_limit is not None:
                job["chunk_limit"] = chunk_limit
            if chunk_strategy_version is not None:
                job["chunk_strategy_version"] = chunk_strategy_version
            if cache_used is not None:
                job["cache_used"] = cache_used
            if cache_invalidated_reason is not None:
                job["cache_invalidated_reason"] = cache_invalidated_reason
            if resume_compatible is not None:
                job["resume_compatible"] = resume_compatible

    async def _update_job_parallel_state(  # pragma: no cover - graph import persistence adapter
        self,
        job_id: str,
        *,
        parallelism: Optional[int] = None,
        active_lanes: Optional[int] = None,
        current_attempt: Optional[int] = None,
        max_attempts: Optional[int] = None,
        lane_statuses: Optional[list[dict[str, Any]]] = None,
    ) -> None:
        """把并发抽取运行态写回进程内任务。"""
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                return
            if parallelism is not None:
                job["parallelism"] = parallelism
            if active_lanes is not None:
                job["active_lanes"] = active_lanes
            if current_attempt is not None:
                job["current_attempt"] = current_attempt
            if max_attempts is not None:
                job["max_attempts"] = max_attempts
            if lane_statuses is not None:
                job["lane_statuses"] = self._clone_lane_statuses(lane_statuses)

    @staticmethod
    def _build_graph_run_progress_text(run_row: dict[str, Any]) -> str:  # pragma: no cover - graph import persistence adapter
        """将持久化 run 转换为最近任务摘要。"""
        status = str(run_row.get("status") or "")
        completed_chunks = int(run_row.get("completed_chunks") or 0)
        total_chunks = int(run_row.get("total_chunks") or 0)
        processed_chars = int(run_row.get("processed_chars") or 0)
        total_chars = int(run_row.get("total_chars") or 0)
        if status == "paused":
            return "已暂停，可重新提交同文件继续"
        if status == "succeeded":
            return "任务执行完成"
        if total_chunks > 0:
            return f"已完成 {completed_chunks}/{total_chunks} 块，已处理 {processed_chars}/{total_chars} 字"
        return "任务执行中"

    def _run_row_to_job_response(self, run_row: dict[str, Any]) -> dict[str, Any]:  # pragma: no cover - graph import persistence adapter
        """把持久化 run 映射为任务响应结构。"""
        docling_failed_pages = self._json_loads(run_row.get("docling_failed_pages"), [])
        result: dict[str, Any] = {
            "file_name": run_row.get("file_name"),
            "file_hash": run_row.get("file_hash"),
            "parsed_text_length": int(run_row.get("parsed_text_length") or 0),
            "processed_text_length": int(run_row.get("processed_chars") or 0),
            "chunk_count": int(run_row.get("total_chunks") or 0),
            "completed_chunks": int(run_row.get("completed_chunks") or 0),
            "chunk_limit": int(run_row.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT),
            "chunk_strategy_version": int(run_row.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION),
            "entity_count": int(run_row.get("entity_count") or 0),
            "relation_count": int(run_row.get("relation_count") or 0),
            "parser_used": run_row.get("parser_used"),
            "parse_warning": run_row.get("parse_warning") or "",
            "docling_failed_pages": docling_failed_pages,
            "pdf_fallback_used": bool(run_row.get("pdf_fallback_used")),
            "cache_used": False,
            "cache_invalidated_reason": "",
            "resume_compatible": True,
            "parallelism": 0,
            "active_lanes": 0,
            "current_attempt": 0,
            "max_attempts": self.GRAPH_LLM_MAX_ATTEMPTS,
            "lane_statuses": [],
            "message": self._build_graph_run_progress_text(run_row),
        }
        if run_row.get("status") == "paused":
            result["pause_reason"] = run_row.get("last_error") or "当前任务已暂停，可重新提交同文件继续"
        return {
            "job_id": f"graph-run-{run_row.get('run_id')}",
            "job_type": self.IMPORT_JOB_TYPE,
            "status": run_row.get("status"),
            "progress_text": self._build_graph_run_progress_text(run_row),
            "created_at": run_row.get("created_at").isoformat(timespec="seconds") + "Z" if isinstance(run_row.get("created_at"), datetime) else run_row.get("created_at"),
            "started_at": run_row.get("created_at").isoformat(timespec="seconds") + "Z" if isinstance(run_row.get("created_at"), datetime) else run_row.get("created_at"),
            "finished_at": run_row.get("finished_at").isoformat(timespec="seconds") + "Z" if isinstance(run_row.get("finished_at"), datetime) else run_row.get("finished_at"),
            "stage": None,
            "stage_detail": "",
            "stage_started_at": None,
            "last_heartbeat_at": run_row.get("updated_at").isoformat(timespec="seconds") + "Z" if isinstance(run_row.get("updated_at"), datetime) else run_row.get("updated_at"),
            "total_chars": int(run_row.get("total_chars") or 0),
            "processed_chars": int(run_row.get("processed_chars") or 0),
            "total_chunks": int(run_row.get("total_chunks") or 0),
            "completed_chunks": int(run_row.get("completed_chunks") or 0),
            "current_chunk_index": None,
            "current_chunk_chars": None,
            "file_hash": run_row.get("file_hash"),
            "run_id": run_row.get("run_id"),
            "is_resumed": False,
            "resume_run_id": None,
            "resumed_from_chunk": None,
            "chunk_limit": int(run_row.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT),
            "chunk_strategy_version": int(run_row.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION),
            "cache_used": False,
            "cache_invalidated_reason": "",
            "resume_compatible": True,
            "parallelism": 0,
            "active_lanes": 0,
            "current_attempt": 0,
            "max_attempts": self.GRAPH_LLM_MAX_ATTEMPTS,
            "lane_statuses": [],
            "warnings": [],
            "trace": [],
            "result": result,
            "error": run_row.get("last_error"),
        }

    async def _register_job(
        self,
        job_type: str,
        *,
        seed: Optional[dict[str, Any]] = None,
        dedupe_active: bool = False,
    ) -> dict[str, Any]:
        """注册进程内后台任务，并按文件哈希去重活跃任务。

        Args:
            job_type: 后台任务类型。
            seed: 可选初始任务字段。
            dedupe_active: 是否复用同文件、同策略的活跃任务。

        Returns:
            dict[str, Any]: 新建任务快照，或去重命中的已有任务快照。
        """
        async with self._job_lock:
            if dedupe_active and seed:
                file_hash = str(seed.get("file_hash") or "")
                chunk_limit = int(seed.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT)
                chunk_strategy_version = int(seed.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION)
                if file_hash:
                    for existing in self._jobs.values():
                        if (
                            existing.get("job_type") == job_type
                            and existing.get("status") in {self.JOB_STATUS_QUEUED, self.JOB_STATUS_RUNNING}
                            and str(existing.get("file_hash") or "") == file_hash
                            and int(existing.get("chunk_limit") or 0) == chunk_limit
                            and int(existing.get("chunk_strategy_version") or 0) == chunk_strategy_version
                        ):
                            row = dict(existing)
                            row["deduplicated"] = True
                            row["existing_job_id"] = str(existing.get("job_id") or "")
                            row["existing_run_id"] = str(existing.get("run_id") or existing.get("resume_run_id") or "")
                            row["_deduplicated"] = True
                            return row
            job_id = f"graph-job-{uuid4().hex}"
            job = {
                "job_id": job_id,
                "job_type": job_type,
                "status": self.JOB_STATUS_QUEUED,
                "progress_text": "任务已提交，等待执行",
                "created_at": self._utc_now_iso(),
                "started_at": None,
                "finished_at": None,
                "stage": None,
                "stage_detail": "",
                "stage_started_at": None,
                "last_heartbeat_at": None,
                "total_chars": 0,
                "processed_chars": 0,
                "total_chunks": 0,
                "completed_chunks": 0,
                "current_chunk_index": 0,
                "current_chunk_chars": 0,
                "file_hash": None,
                "run_id": None,
                "is_resumed": False,
                "resume_run_id": None,
                "resumed_from_chunk": None,
                "chunk_limit": self.GRAPH_IMPORT_CHUNK_LIMIT,
                "chunk_strategy_version": self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
                "cache_used": False,
                "cache_invalidated_reason": "",
                "resume_compatible": None,
                "parallelism": 0,
                "active_lanes": 0,
                "current_attempt": 0,
                "max_attempts": self.GRAPH_LLM_MAX_ATTEMPTS,
                "lane_statuses": [],
                "labels": None,
                "batch_size": None,
                "warnings": [],
                "trace": [],
                "result": None,
                "error": None,
                "_task": None,
                "_watchdog_task": None,
                "_last_soft_timeout_notice_at": None,
            }
            if seed:
                job.update({key: value for key, value in seed.items() if value is not None})
            self._jobs[job_id] = job
            self._prune_jobs_locked()
            return dict(job)

    def _prune_jobs_locked(self) -> None:
        """裁剪旧任务，避免进程内状态无限增长。"""
        if len(self._jobs) <= self._max_retained_jobs:
            return
        removable = sorted(
            self._jobs.values(),
            key=lambda item: item.get("created_at") or "",
        )
        while len(removable) > 0 and len(self._jobs) > self._max_retained_jobs:
            candidate = removable.pop(0)
            if candidate.get("status") in {
                self.JOB_STATUS_PAUSED,
                self.JOB_STATUS_SUCCEEDED,
                self.JOB_STATUS_FAILED,
                self.JOB_STATUS_CANCELLED,
            }:
                self._jobs.pop(candidate["job_id"], None)

    def _make_stage_event_callback(self, loop: asyncio.AbstractEventLoop, job_id: str) -> Callable[..., None]:  # pragma: no cover - background thread callback
        """生成线程安全的任务阶段事件回调，供阻塞线程回传进度。"""
        def _progress(
            action: str,
            stage: Optional[str] = None,
            detail: Optional[str] = None,
            progress_text: Optional[str] = None,
            metrics: Optional[dict[str, Any]] = None,
        ) -> None:
            try:
                asyncio.run_coroutine_threadsafe(
                    self._handle_job_stage_event(
                        job_id,
                        action=action,
                        stage=stage,
                        detail=detail,
                        progress_text=progress_text,
                        metrics=metrics,
                    ),
                    loop,
                )
            except RuntimeError:
                # 进程关闭或 loop 不可用时静默跳过，避免后台线程崩溃。
                return
        return _progress

    async def _handle_job_stage_event(  # pragma: no cover - background thread callback
        self,
        job_id: str,
        *,
        action: str,
        stage: Optional[str] = None,
        detail: Optional[str] = None,
        progress_text: Optional[str] = None,
        metrics: Optional[dict[str, Any]] = None,
    ) -> None:
        """接收阻塞线程回传的阶段事件并更新进程内任务状态。

        Args:
            job_id: 任务标识。
            action: 阶段事件动作，支持 enter、heartbeat、complete、warning。
            stage: 可选阶段名称。
            detail: 可选阶段明细。
            progress_text: 可选对外进度文案。
            metrics: 可选进度指标。

        Returns:
            None: 无返回值。
        """
        if action == "enter" and stage:
            await self._enter_job_stage(job_id, stage, detail or progress_text or "", metrics=metrics)
            return
        if action == "heartbeat":
            await self._heartbeat_job_stage(job_id, detail=detail or progress_text, progress_text=progress_text, metrics=metrics)
            return
        if action == "complete":
            await self._complete_job_stage(job_id, detail=detail or progress_text, metrics=metrics)
            return
        if action == "warning":
            message = detail or progress_text or ""
            if message:
                await self._add_job_warning(job_id, message)

    @staticmethod
    def _apply_job_metrics(job: dict[str, Any], metrics: Optional[dict[str, Any]] = None) -> None:
        """将线程回传的进度指标合并到任务对象。"""
        if not metrics:
            return
        for key in (
            "total_chars",
            "processed_chars",
            "total_chunks",
            "completed_chunks",
            "current_chunk_index",
            "current_chunk_chars",
        ):
            if key in metrics and metrics[key] is not None:
                job[key] = metrics[key]

    async def _set_job_state(
        self,
        job_id: str,
        *,
        status: Optional[str] = None,
        progress_text: Optional[str] = None,
        result: Optional[dict[str, Any]] = None,
        error: Optional[str] = None,
        started: bool = False,
        finished: bool = False,
    ) -> None:
        """更新进程内任务状态、结果、错误和生命周期时间。

        Args:
            job_id: 任务标识。
            status: 可选任务状态。
            progress_text: 可选对外进度文案。
            result: 可选任务结果。
            error: 可选错误文本。
            started: 是否写入开始时间。
            finished: 是否写入结束时间。

        Returns:
            None: 无返回值。
        """
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                return
            if status:
                job["status"] = status
            if progress_text is not None:
                job["progress_text"] = progress_text
            if result is not None:
                job["result"] = result
            if error is not None:
                job["error"] = error
                self._mark_current_trace_finished(job, status or job.get("status") or self.JOB_STATUS_FAILED, error)
            if started and not job.get("started_at"):
                job["started_at"] = self._utc_now_iso()
            if finished:
                job["finished_at"] = self._utc_now_iso()

    def _mark_current_trace_finished(self, job: dict[str, Any], status: str, error: Optional[str] = None) -> None:
        """把任务当前运行中的 trace 阶段标记为结束。"""
        trace = job.get("trace") or []
        for item in reversed(trace):
            if not item.get("finished_at"):
                finished_at = self._utc_now_iso()
                item["finished_at"] = finished_at
                item["status"] = status
                item["error"] = error
                started_at = self._parse_utc_iso(item.get("started_at"))
                finished_dt = self._parse_utc_iso(finished_at)
                if started_at and finished_dt:
                    item["duration_ms"] = int((finished_dt - started_at).total_seconds() * 1000)
                break

    async def _enter_job_stage(
        self,
        job_id: str,
        stage: str,
        detail: str = "",
        *,
        metrics: Optional[dict[str, Any]] = None,
    ) -> None:
        """进入任务新阶段并追加 trace 记录。

        Args:
            job_id: 任务标识。
            stage: 阶段名称。
            detail: 阶段说明文本。
            metrics: 可选进度指标。

        Returns:
            None: 无返回值。
        """
        now_iso = self._utc_now_iso()
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                return
            self._apply_job_metrics(job, metrics)
            if job.get("stage") and job.get("stage") != stage:
                self._mark_current_trace_finished(job, "succeeded")
            job["stage"] = stage
            job["stage_detail"] = detail or ""
            job["stage_started_at"] = now_iso
            job["last_heartbeat_at"] = now_iso
            job["_last_soft_timeout_notice_at"] = None
            job["progress_text"] = detail or stage
            trace = job.setdefault("trace", [])
            trace.append(
                {
                    "stage": stage,
                    "detail": detail or "",
                    "status": "running",
                    "started_at": now_iso,
                    "finished_at": None,
                    "duration_ms": None,
                    "error": None,
                }
            )

    async def _heartbeat_job_stage(
        self,
        job_id: str,
        *,
        detail: Optional[str] = None,
        progress_text: Optional[str] = None,
        metrics: Optional[dict[str, Any]] = None,
    ) -> None:
        """刷新当前阶段心跳、进度文案和进度指标。

        Args:
            job_id: 任务标识。
            detail: 可选阶段明细。
            progress_text: 可选对外进度文案。
            metrics: 可选进度指标。

        Returns:
            None: 无返回值。
        """
        now_iso = self._utc_now_iso()
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                return
            self._apply_job_metrics(job, metrics)
            job["last_heartbeat_at"] = now_iso
            if detail is not None:
                job["stage_detail"] = detail
            if progress_text is not None:
                job["progress_text"] = progress_text
            elif detail is not None:
                job["progress_text"] = detail
            trace = job.get("trace") or []
            for item in reversed(trace):
                if not item.get("finished_at"):
                    if detail is not None:
                        item["detail"] = detail
                    break

    async def _complete_job_stage(
        self,
        job_id: str,
        detail: Optional[str] = None,
        *,
        metrics: Optional[dict[str, Any]] = None,
    ) -> None:
        """完成任务当前阶段并收口 trace。

        Args:
            job_id: 任务标识。
            detail: 可选完成明细。
            metrics: 可选进度指标。

        Returns:
            None: 无返回值。
        """
        now_iso = self._utc_now_iso()
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job or not job.get("stage"):
                return
            self._apply_job_metrics(job, metrics)
            job["last_heartbeat_at"] = now_iso
            if detail is not None:
                job["stage_detail"] = detail
                job["progress_text"] = detail
                trace = job.get("trace") or []
                for item in reversed(trace):
                    if not item.get("finished_at"):
                        item["detail"] = detail
                        break
            self._mark_current_trace_finished(job, "succeeded")

    async def _add_job_warning(
        self,
        job_id: str,
        message: str,
        *,
        elapsed_seconds: Optional[int] = None,
    ) -> None:
        """为任务追加可展示 warning。

        Args:
            job_id: 任务标识。
            message: warning 文本。
            elapsed_seconds: 可选阶段已耗时秒数。

        Returns:
            None: 无返回值。
        """
        warning = {
            "message": message,
            "stage": None,
            "timestamp": self._utc_now_iso(),
            "elapsed_seconds": elapsed_seconds,
        }
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                return
            warning["stage"] = job.get("stage")
            job.setdefault("warnings", []).append(warning)

    async def _watch_job_soft_timeout(self, job_id: str) -> None:  # pragma: no cover - wall-clock watchdog
        """监控单阶段运行时长，超过阈值只提示不失败。"""
        try:
            while True:
                await asyncio.sleep(5)
                async with self._job_lock:
                    job = self._jobs.get(job_id)
                    if not job or job.get("status") != self.JOB_STATUS_RUNNING:
                        return
                    stage = job.get("stage")
                    stage_started_at = self._parse_utc_iso(job.get("stage_started_at"))
                    if not stage or stage_started_at is None:
                        continue
                    now_dt = datetime.now(stage_started_at.tzinfo)
                    elapsed_seconds = int((now_dt - stage_started_at).total_seconds())
                    if elapsed_seconds < self.SOFT_TIMEOUT_SECONDS:
                        continue
                    last_notice_at = self._parse_utc_iso(job.get("_last_soft_timeout_notice_at"))
                    if last_notice_at and (now_dt - last_notice_at) < timedelta(seconds=self.SOFT_TIMEOUT_REPEAT_SECONDS):
                        continue
                    message = f"当前任务已在阶段 {stage} 持续 {elapsed_seconds} 秒"
                    job["_last_soft_timeout_notice_at"] = self._utc_now_iso()
                    job.setdefault("warnings", []).append(
                        {
                            "message": message,
                            "stage": stage,
                            "timestamp": job["_last_soft_timeout_notice_at"],
                            "elapsed_seconds": elapsed_seconds,
                        }
                    )
                    logger.warning(message)
        except asyncio.CancelledError:
            raise

    async def get_graph_job(self, job_id: str) -> dict[str, Any]:
        """获取图谱后台任务状态，支持进程内任务和持久化 run。

        Args:
            job_id: 后台任务标识或 graph-run 前缀的持久化 run 标识。

        Returns:
            dict[str, Any]: 任务状态快照或任务不存在错误。
        """
        async with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                pass
            else:
                return {"ok": True, "data": self._build_job_response(job)}
        if job_id.startswith("graph-run-"):
            run_id = job_id.removeprefix("graph-run-")
            row = await fetch_one("SELECT * FROM graph_import_runs WHERE run_id = :run_id", {"run_id": run_id})
            if row:
                return {"ok": True, "data": self._run_row_to_job_response(row)}
            return {"ok": False, "error": "任务不存在或已过期"}
        if job_id.startswith("graph-job-"):
            return {"ok": False, "error": "任务不存在或已过期"}
        return {"ok": False, "error": "任务不存在或已过期"}

    async def list_graph_jobs(self, limit: int = 10) -> dict[str, Any]:
        """返回最近图谱后台任务列表，合并进程内任务和持久化 run。

        Args:
            limit: 返回数量上限。

        Returns:
            dict[str, Any]: 最近任务快照列表。
        """
        safe_limit = max(1, min(int(limit or 10), 50))
        async with self._job_lock:
            jobs = [self._build_job_response(job) for job in self._jobs.values()]
        active_run_ids = {str(job.get("run_id") or "") for job in jobs if job.get("run_id")}
        persisted_rows = await fetch_all(
            """
            SELECT *
            FROM graph_import_runs
            ORDER BY updated_at DESC
            LIMIT :limit
            """,
            {"limit": safe_limit * 3},
        )
        persisted_jobs = [
            self._run_row_to_job_response(row)
            for row in persisted_rows
            if str(row.get("run_id") or "") not in active_run_ids
        ]
        jobs = sorted(
            [*jobs, *persisted_jobs],
            key=lambda item: item.get("created_at") or "",
            reverse=True,
        )[:safe_limit]
        return {"ok": True, "data": jobs}

    async def _start_background_job(  # pragma: no cover - background task orchestration
        self,
        *,
        job_type: str,
        runner: Callable[[str], Awaitable[dict[str, Any]]],
        seed: Optional[dict[str, Any]] = None,
        dedupe_active: bool = False,
    ) -> dict[str, Any]:
        """创建图谱后台任务、启动 watchdog，并立即返回任务信息。

        Args:
            job_type: 后台任务类型。
            runner: 接收 job_id 并执行实际后台工作的异步函数。
            seed: 可选初始任务字段。
            dedupe_active: 是否复用同文件、同策略的活跃任务。

        Returns:
            dict[str, Any]: 任务提交响应。
        """
        job = await self._register_job(job_type, seed=seed, dedupe_active=dedupe_active)
        if job.get("_deduplicated"):
            response_data = self._build_job_response(job)
            response_data["message"] = "检测到同资源任务正在执行，已返回现有任务"
            return {"ok": True, "data": response_data}
        job_id = job["job_id"]

        async def _run() -> None:
            await self._set_job_state(
                job_id,
                status=self.JOB_STATUS_RUNNING,
                progress_text="任务开始执行",
                started=True,
            )
            watchdog_task = asyncio.create_task(self._watch_job_soft_timeout(job_id), name=f"watchdog:{job_id}")
            async with self._job_lock:
                if job_id in self._jobs:
                    self._jobs[job_id]["_watchdog_task"] = watchdog_task
            try:
                result = await runner(job_id)
                job_status = str(result.pop("_job_status", self.JOB_STATUS_SUCCEEDED)) if isinstance(result, dict) else self.JOB_STATUS_SUCCEEDED
                if job_status == self.JOB_STATUS_PAUSED:
                    pause_message = ""
                    pause_error = None
                    if isinstance(result, dict):
                        pause_message = str(result.get("message") or "任务已暂停，可重新提交同文件继续")
                        pause_error = result.get("pause_reason") or result.get("last_error")
                    await self._set_job_state(
                        job_id,
                        status=self.JOB_STATUS_PAUSED,
                        progress_text=pause_message or "任务已暂停",
                        result=result if isinstance(result, dict) else None,
                        error=str(pause_error) if pause_error else None,
                        finished=True,
                    )
                else:
                    await self._set_job_state(
                        job_id,
                        status=self.JOB_STATUS_SUCCEEDED,
                        progress_text="任务执行完成",
                        result=result,
                        finished=True,
                    )
            except asyncio.CancelledError:
                await self._set_job_state(
                    job_id,
                    status=self.JOB_STATUS_CANCELLED,
                    progress_text="任务已取消",
                    error="任务已取消",
                    finished=True,
                )
                raise
            except Exception as exc:
                await self._set_job_state(
                    job_id,
                    status=self.JOB_STATUS_FAILED,
                    progress_text="任务执行失败",
                    error=f"{exc}\n{traceback.format_exc()}",
                    finished=True,
                )
            finally:
                watchdog = None
                async with self._job_lock:
                    if job_id in self._jobs:
                        watchdog = self._jobs[job_id].get("_watchdog_task")
                        self._jobs[job_id]["_watchdog_task"] = None
                if watchdog:
                    watchdog.cancel()
                    try:
                        await watchdog
                    except asyncio.CancelledError:
                        pass

        task = asyncio.create_task(_run(), name=f"{job_type}:{job_id}")
        async with self._job_lock:
            if job_id in self._jobs:
                self._jobs[job_id]["_task"] = task
        response_data = self._build_job_response(job)
        response_data["message"] = "任务已提交到后台处理"
        return {
            "ok": True,
            "data": response_data,
        }

    async def submit_graph_extract_import_job(self, file: UploadFile) -> dict[str, Any]:  # pragma: no cover - background task orchestration
        """读取上传文件并提交图谱抽取导入后台任务。

        Args:
            file: FastAPI 上传文件对象。

        Returns:
            dict[str, Any]: 任务提交结果、去重命中结果或校验错误。
        """
        filename = file.filename or "uploaded_file"
        try:
            file_bytes = await file.read()
            if not file_bytes:
                return {"ok": False, "error": "上传文件为空"}
            if len(file_bytes) > 50 * 1024 * 1024:
                return {"ok": False, "error": "文件大小不能超过 50MB"}
            file_hash = self._compute_file_hash(file_bytes)
            resumable_run = await self._find_resumable_graph_run(
                file_hash,
                chunk_limit=self.GRAPH_IMPORT_CHUNK_LIMIT,
                chunk_strategy_version=self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
            )
            if resumable_run:
                active_job = await self._find_active_job_by_run_id(str(resumable_run.get("run_id") or ""))
                if active_job and resumable_run.get("status") == "running":
                    active_job["message"] = "检测到同文件导入任务仍在执行，已返回当前任务状态"
                    active_job["deduplicated"] = True
                    active_job["existing_job_id"] = active_job.get("job_id")
                    active_job["existing_run_id"] = resumable_run.get("run_id")
                    return {"ok": True, "data": active_job}
            resumed_from_chunk = None
            if resumable_run:
                resumed_from_chunk = await self._find_first_unfinished_chunk_index(str(resumable_run.get("run_id") or ""))
            return await self._start_background_job(
                job_type=self.IMPORT_JOB_TYPE,
                runner=lambda job_id: self._extract_and_import_graph_sync(file_bytes, filename, job_id),
                dedupe_active=True,
                seed={
                    "file_hash": file_hash,
                    "is_resumed": bool(resumable_run),
                    "resume_run_id": resumable_run.get("run_id") if resumable_run else None,
                    "resumed_from_chunk": resumed_from_chunk,
                    "chunk_limit": self.GRAPH_IMPORT_CHUNK_LIMIT,
                    "chunk_strategy_version": self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
                    "resume_compatible": bool(resumable_run) if resumable_run else None,
                    "parallelism": min(self.GRAPH_LLM_MAX_PARALLELISM, len(self._graph_llm_lane_configs)),
                    "active_lanes": min(self.GRAPH_LLM_MAX_PARALLELISM, len(self._graph_llm_lane_configs)),
                    "current_attempt": 0,
                    "max_attempts": self.GRAPH_LLM_MAX_ATTEMPTS,
                    "progress_text": self._format_resume_message(resumed_from_chunk) if resumable_run else "任务已提交，等待执行",
                },
            )
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/jobs/extract-import", "extract_import_submit", exc, filename=filename)
            return {"ok": False, "error": f"提交抽取导入任务失败: {exc}"}

    async def submit_graph_embed_milvus_job(self, data: Optional[dict[str, Any]] = None) -> dict[str, Any]:  # pragma: no cover - background task orchestration
        """提交图谱嵌入 Milvus 后台任务。

        Args:
            data: 可选标签列表和批大小配置。

        Returns:
            dict[str, Any]: 任务提交结果或错误响应。
        """
        try:
            payload = data or {}
            labels_obj = payload.get("labels")
            labels = [str(item).strip() for item in labels_obj if str(item).strip()] if isinstance(labels_obj, list) else None
            batch_size_raw = payload.get("batch_size", 100)
            try:
                batch_size = int(batch_size_raw)
            except (TypeError, ValueError):
                batch_size = 100
            return await self._start_background_job(
                job_type=self.EMBED_JOB_TYPE,
                runner=lambda job_id: self._embed_graph_to_milvus_sync(payload, job_id),
                seed={
                    "labels": labels or None,
                    "batch_size": batch_size,
                },
            )
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/jobs/embed-milvus", "embed_graph_submit", exc)
            return {"ok": False, "error": f"提交嵌入任务失败: {exc}"}

    async def _extract_and_import_graph_sync(self, file_bytes: bytes, filename: str, job_id: str) -> dict[str, Any]:  # pragma: no cover - external LLM/Neo4j background runtime
        """执行图谱抽取导入主体，包含解析、续跑、抽取和写入。

        Args:
            file_bytes: 上传文件二进制内容。
            filename: 上传文件名。
            job_id: 进程内后台任务标识。

        Returns:
            dict[str, Any]: 图谱导入结果，可能返回暂停态结果。
        """
        loop = asyncio.get_running_loop()
        stage_callback = self._make_stage_event_callback(loop, job_id)
        file_hash = self._compute_file_hash(file_bytes)
        prepared = await asyncio.to_thread(
            self._parse_and_split_graph_blocking,
            file_bytes,
            filename,
            stage_callback,
        )
        chunks = list(prepared.pop("_chunks", []))
        total_chunks = len(chunks)
        file_size = len(file_bytes)
        chunk_limit = int(prepared.get("chunk_limit") or self.GRAPH_IMPORT_CHUNK_LIMIT)
        chunk_strategy_version = int(
            prepared.get("chunk_strategy_version") or self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION
        )
        cache_used = bool(prepared.get("cache_used"))
        cache_invalidated_reason = str(prepared.get("cache_invalidated_reason") or "")

        resumable_run = await self._find_resumable_graph_run(
            file_hash,
            chunk_limit=chunk_limit,
            chunk_strategy_version=chunk_strategy_version,
        )
        latest_run = await self._find_latest_graph_run_by_file_hash(file_hash)
        existing_chunks: list[dict[str, Any]] = []
        is_resumed = False
        run_id = ""
        resumed_from_chunk: Optional[int] = None
        resume_compatible: Optional[bool] = None

        if latest_run and (
            int(latest_run.get("chunk_limit") or 0) != chunk_limit
            or int(latest_run.get("chunk_strategy_version") or 0) != chunk_strategy_version
        ):
            resume_compatible = False
            compatibility_reason = (
                f"旧断点分块策略已过期（旧策略: limit={int(latest_run.get('chunk_limit') or 0)}, "
                f"version={int(latest_run.get('chunk_strategy_version') or 0)}；"
                f"当前策略: limit={chunk_limit}, version={chunk_strategy_version}），已重新建任务"
            )
            if str(latest_run.get("status") or "") in {"running", "paused"}:
                await self._set_graph_run_status(
                    str(latest_run.get("run_id") or ""),
                    "failed",
                    last_error=compatibility_reason,
                    finished=True,
                )
            cache_invalidated_reason = compatibility_reason
            logger.info(
                "graph_import.resume_incompatible file_hash=%s latest_run_id=%s reason=%s",
                file_hash,
                latest_run.get("run_id"),
                compatibility_reason,
            )
        elif resumable_run:
            resume_compatible = True

        if resumable_run:
            run_id = str(resumable_run.get("run_id") or "")
            existing_chunks = await self._load_graph_import_chunks(run_id)
            if existing_chunks and len(existing_chunks) != total_chunks:
                resume_compatible = False
                incompatibility_reason = "旧导入断点与当前分块结果不一致，已终止旧断点并改为新任务导入"
                await self._set_graph_run_status(
                    run_id,
                    "failed",
                    last_error=incompatibility_reason,
                    finished=True,
                )
                cache_invalidated_reason = incompatibility_reason
                resumable_run = None
                existing_chunks = []
                logger.info(
                    "graph_import.resume_chunk_mismatch file_hash=%s run_id=%s chunk_limit=%s total_chunks=%s",
                    file_hash,
                    run_id,
                    chunk_limit,
                    total_chunks,
                )

        if not resumable_run:
            run_id = f"graph-run-{uuid4().hex}"
            await self._create_graph_import_run(
                run_id=run_id,
                file_hash=file_hash,
                file_name=filename,
                file_size=file_size,
                prepared=prepared,
            )
            await self._initialize_graph_import_chunks(run_id, chunks)
            existing_chunks = await self._load_graph_import_chunks(run_id)
        else:
            is_resumed = True
            await self._sync_graph_import_run_metadata(
                run_id,
                file_name=filename,
                file_size=file_size,
                prepared=prepared,
                status="running",
                clear_error=True,
            )
            if not existing_chunks:
                await self._initialize_graph_import_chunks(run_id, chunks)
                existing_chunks = await self._load_graph_import_chunks(run_id)

        logger.info(
            "graph_import.run_selected file_hash=%s run_id=%s resumed=%s chunk_limit=%s chunk_strategy_version=%s total_chunks=%s first_chunk_chars=%s cache_used=%s resume_compatible=%s",
            file_hash,
            run_id,
            is_resumed,
            chunk_limit,
            chunk_strategy_version,
            total_chunks,
            len(chunks[0]) if chunks else 0,
            cache_used,
            resume_compatible,
        )

        progress = await self._recompute_graph_run_progress(run_id)
        chunk_map = {int(item.get("chunk_index") or 0): item for item in existing_chunks}
        for index in range(1, total_chunks + 1):
            if chunk_map.get(index, {}).get("status") != "succeeded":
                resumed_from_chunk = index if is_resumed else None
                break

        await self._update_job_resume_state(
            job_id,
            run_id=run_id,
            file_hash=file_hash,
            is_resumed=is_resumed,
            resumed_from_chunk=resumed_from_chunk,
            chunk_limit=chunk_limit,
            chunk_strategy_version=chunk_strategy_version,
            cache_used=cache_used,
            cache_invalidated_reason=cache_invalidated_reason,
            resume_compatible=resume_compatible,
        )
        prepared.update(
            {
                "file_hash": file_hash,
                "run_id": run_id,
                "is_resumed": is_resumed,
                "resume_run_id": run_id if is_resumed else None,
                "resumed_from_chunk": resumed_from_chunk,
                "chunk_limit": chunk_limit,
                "chunk_strategy_version": chunk_strategy_version,
                "cache_used": cache_used,
                "cache_invalidated_reason": cache_invalidated_reason,
                "resume_compatible": resume_compatible,
            }
        )

        if is_resumed:
            await self._heartbeat_job_stage(
                job_id,
                detail=self._format_resume_message(resumed_from_chunk),
                progress_text=self._format_resume_message(resumed_from_chunk),
                metrics={
                    "total_chars": prepared.get("parsed_text_length") or 0,
                    "processed_chars": progress["processed_chars"],
                    "total_chunks": total_chunks,
                    "completed_chunks": progress["completed_chunks"],
                    "current_chunk_index": resumed_from_chunk or 0,
                    "current_chunk_chars": len(chunks[(resumed_from_chunk or 1) - 1]) if resumed_from_chunk and resumed_from_chunk <= total_chunks else 0,
                },
            )

        total_entities = int(progress["entity_count"])
        total_relations = int(progress["relation_count"])
        total_node_batches = int(progress.get("node_batches") or 0)
        total_relation_batches = int(progress.get("relation_batches") or 0)
        processed_chars = int(progress["processed_chars"])
        completed_chunks = int(progress["completed_chunks"])
        from app.knowledge_graph.extract_graph_data_utils import create_graph_llm_bundle

        lane_configs = list(self._graph_llm_lane_configs[: self.GRAPH_LLM_MAX_PARALLELISM])
        if not lane_configs:
            raise RuntimeError("未配置可用的图谱抽取模型，请检查 MODEL_* 或 GRAPH_LLM_* 环境变量")

        lane_statuses = [
            {
                "lane_id": int(item["lane_id"]),
                "label": self._format_lane_label(item),
                "model_name": str(item["model_name"]),
                "status": "ready",
                "current_chunk_index": None,
                "current_attempt": 0,
                "last_error": "",
                "backoff_until": None,
            }
            for item in lane_configs
        ]
        parallelism = len(lane_statuses)
        await self._update_job_parallel_state(
            job_id,
            parallelism=parallelism,
            active_lanes=self._count_active_lanes(lane_statuses),
            current_attempt=0,
            max_attempts=self.GRAPH_LLM_MAX_ATTEMPTS,
            lane_statuses=lane_statuses,
        )
        prepared.update(
            {
                "parallelism": parallelism,
                "active_lanes": self._count_active_lanes(lane_statuses),
                "current_attempt": 0,
                "max_attempts": self.GRAPH_LLM_MAX_ATTEMPTS,
                "lane_statuses": self._clone_lane_statuses(lane_statuses),
            }
        )
        logger.info(
            "graph_import.parallel_extract_start run_id=%s parallelism=%s lanes=%s",
            run_id,
            parallelism,
            [self._format_lane_label(item) for item in lane_configs],
        )

        pending_chunks = [
            (index, chunk)
            for index, chunk in enumerate(chunks, start=1)
            if self._normalize_chunk_status(chunk_map.get(index, {}).get("status")) != "succeeded"
        ]
        pending_chunks_map = {index: chunk for index, chunk in pending_chunks}

        if pending_chunks:
            work_queue: asyncio.Queue[tuple[int, str]] = asyncio.Queue()
            for item in pending_chunks[:parallelism]:
                await work_queue.put(item)
            next_schedule_cursor = min(parallelism, len(pending_chunks))
            write_cursor = 0
            active_extractions = 0
            attempt_counts: dict[int, int] = {}
            result_buffer: dict[int, dict[str, Any]] = {}
            stop_event = asyncio.Event()
            state_condition = asyncio.Condition()
            terminal_error: Optional[str] = None

            lane_bundles = {
                int(item["lane_id"]): create_graph_llm_bundle(
                    api_key=str(item["api_key"]),
                    base_url=str(item["base_url"]),
                    model_name=str(item["model_name"]),
                    timeout=self.GRAPH_LLM_REQUEST_TIMEOUT_SECONDS,
                    max_retries=0,
                    lane_id=int(item["lane_id"]),
                )
                for item in lane_configs
            }

            async def sync_lane_status(
                lane_id: int,
                *,
                status: Optional[str] = None,
                current_chunk_index: Optional[int] = None,
                current_attempt: Optional[int] = None,
                last_error: Optional[str] = None,
                backoff_until: Optional[str] = None,
            ) -> None:
                """同步单个 LLM lane 的当前运行状态到任务快照。"""
                for item in lane_statuses:
                    if int(item.get("lane_id") or 0) != lane_id:
                        continue
                    if status is not None:
                        item["status"] = status
                    item["current_chunk_index"] = current_chunk_index
                    if current_attempt is not None:
                        item["current_attempt"] = current_attempt
                    if last_error is not None:
                        item["last_error"] = last_error
                    item["backoff_until"] = backoff_until
                    break
                await self._update_job_parallel_state(
                    job_id,
                    parallelism=parallelism,
                    active_lanes=self._count_active_lanes(lane_statuses),
                    current_attempt=current_attempt,
                    max_attempts=self.GRAPH_LLM_MAX_ATTEMPTS,
                    lane_statuses=lane_statuses,
                )

            async def fail_chunk_and_task(chunk_index: int, error_message: str) -> None:
                """标记当前块和整个导入任务失败，并通知并发协程停止。"""
                nonlocal terminal_error
                if terminal_error:
                    return
                terminal_error = error_message
                await self._mark_graph_chunk_failed(run_id, chunk_index, error_message)
                await self._set_graph_run_status(run_id, "failed", last_error=error_message)
                async with state_condition:
                    state_condition.notify_all()
                stop_event.set()

            async def lane_worker(lane_config: dict[str, Any]) -> None:
                """从队列领取待抽取文本块，并在对应 LLM lane 上执行抽取。"""
                nonlocal active_extractions, terminal_error, write_cursor
                lane_id = int(lane_config["lane_id"])
                llm_bundle = lane_bundles[lane_id]
                lane_label = self._format_lane_label(lane_config)
                while not stop_event.is_set():
                    try:
                        chunk_index, chunk = await asyncio.wait_for(work_queue.get(), timeout=1.0)
                    except asyncio.TimeoutError:
                        if stop_event.is_set() or write_cursor >= len(pending_chunks):
                            break
                        continue

                    if stop_event.is_set():
                        work_queue.task_done()
                        break

                    attempt = int(attempt_counts.get(chunk_index) or 0) + 1
                    attempt_counts[chunk_index] = attempt
                    chunk_chars = len(chunk)
                    lane_stage_callback = self._make_lane_stage_callback(
                        stage_callback,
                        lane_id=lane_id,
                        chunk_index=chunk_index,
                        total_chunks=total_chunks,
                        attempt=attempt,
                        max_attempts=self.GRAPH_LLM_MAX_ATTEMPTS,
                    )

                    await self._mark_graph_chunk_running(run_id, chunk_index)
                    await sync_lane_status(
                        lane_id,
                        status="running",
                        current_chunk_index=chunk_index,
                        current_attempt=attempt,
                        last_error="",
                        backoff_until=None,
                    )
                    await self._heartbeat_job_stage(
                        job_id,
                        detail=(
                            f"正在并发抽取第 {chunk_index}/{total_chunks} 段文本，"
                            f"{lane_label} 正在执行第 {attempt}/{self.GRAPH_LLM_MAX_ATTEMPTS} 次尝试"
                        ),
                        progress_text=(
                            f"并发抽取中：第 {chunk_index}/{total_chunks} 段文本，"
                            f"{lane_label} 第 {attempt}/{self.GRAPH_LLM_MAX_ATTEMPTS} 次尝试"
                        ),
                        metrics={
                            "total_chars": prepared.get("parsed_text_length") or 0,
                            "processed_chars": processed_chars,
                            "total_chunks": total_chunks,
                            "completed_chunks": completed_chunks,
                            "current_chunk_index": chunk_index,
                            "current_chunk_chars": chunk_chars,
                        },
                    )

                    async with state_condition:
                        active_extractions += 1

                    try:
                        chunk_result = await asyncio.to_thread(
                            self._extract_single_graph_chunk_blocking,
                            chunk,
                            chunk_index,
                            total_chunks,
                            prepared.get("parsed_text_length") or 0,
                            processed_chars,
                            lane_stage_callback,
                            llm_bundle,
                        )
                    except Exception as exc:
                        error_text = str(exc).strip() or exc.__class__.__name__
                        failure_text = (
                            f"第 {chunk_index}/{total_chunks} 块在 {lane_label} 上第 "
                            f"{attempt}/{self.GRAPH_LLM_MAX_ATTEMPTS} 次尝试失败: {error_text}"
                        )
                        await self._mark_graph_chunk_failed(run_id, chunk_index, failure_text)
                        error_type = self._classify_lane_error(exc)
                        first_error = self._first_error_line(error_text) or error_text

                        async with state_condition:
                            active_extractions = max(0, active_extractions - 1)
                            state_condition.notify_all()

                        if attempt >= self.GRAPH_LLM_MAX_ATTEMPTS:
                            await fail_chunk_and_task(
                                chunk_index,
                                (
                                    f"第 {chunk_index}/{total_chunks} 块累计失败 "
                                    f"{self.GRAPH_LLM_MAX_ATTEMPTS} 次，任务失败。最近错误: {first_error}"
                                ),
                            )
                            work_queue.task_done()
                            break

                        await work_queue.put((chunk_index, chunk))
                        await self._add_job_warning(
                            job_id,
                            (
                                f"{lane_label} 处理第 {chunk_index}/{total_chunks} 块失败，"
                                f"已回队重试（{attempt}/{self.GRAPH_LLM_MAX_ATTEMPTS}）：{first_error}"
                            ),
                        )

                        if error_type == "permanent":
                            await sync_lane_status(
                                lane_id,
                                status="disabled",
                                current_chunk_index=chunk_index,
                                current_attempt=attempt,
                                last_error=first_error,
                                backoff_until=None,
                            )
                            await self._add_job_warning(job_id, f"{lane_label} 因鉴权或模型配置错误已摘除")
                            if self._count_viable_lanes(lane_statuses) <= 0:
                                await fail_chunk_and_task(
                                    chunk_index,
                                    "所有图谱抽取模型 lane 均不可用，任务失败",
                                )
                                work_queue.task_done()
                                break
                            work_queue.task_done()
                            break

                        backoff_until = (
                            datetime.utcnow() + timedelta(seconds=self.GRAPH_LLM_TRANSIENT_BACKOFF_SECONDS)
                        ).isoformat(timespec="seconds") + "Z"
                        await sync_lane_status(
                            lane_id,
                            status="backoff",
                            current_chunk_index=chunk_index,
                            current_attempt=attempt,
                            last_error=first_error,
                            backoff_until=backoff_until,
                        )
                        work_queue.task_done()
                        if not stop_event.is_set():
                            await asyncio.sleep(self.GRAPH_LLM_TRANSIENT_BACKOFF_SECONDS)
                            if not stop_event.is_set():
                                await sync_lane_status(
                                    lane_id,
                                    status="ready",
                                    current_chunk_index=None,
                                    current_attempt=attempt,
                                    last_error="",
                                    backoff_until=None,
                                )
                        continue

                    async with state_condition:
                        active_extractions = max(0, active_extractions - 1)
                        result_buffer[chunk_index] = {
                            **chunk_result,
                            "attempt": attempt,
                            "lane_id": lane_id,
                            "lane_label": lane_label,
                        }
                        state_condition.notify_all()
                    await sync_lane_status(
                        lane_id,
                        status="ready",
                        current_chunk_index=None,
                        current_attempt=attempt,
                        last_error="",
                        backoff_until=None,
                    )
                    work_queue.task_done()

            async def writer_loop() -> None:
                """按块序等待抽取结果并串行写入 Neo4j，保证续跑进度有序推进。"""
                nonlocal next_schedule_cursor, write_cursor, processed_chars, completed_chunks
                nonlocal total_entities, total_relations, total_node_batches, total_relation_batches, terminal_error
                while write_cursor < len(pending_chunks):
                    index, chunk = pending_chunks[write_cursor]
                    chunk_chars = len(chunk)
                    async with state_condition:
                        await state_condition.wait_for(
                            lambda: terminal_error is not None or index in result_buffer
                        )
                        if terminal_error:
                            raise RuntimeError(terminal_error)
                        chunk_result = result_buffer.pop(index)

                    node_queries = chunk_result["node_queries"]
                    relation_queries = chunk_result["relation_queries"]
                    stats = chunk_result["stats"]
                    node_batches = len(node_queries)
                    relation_batches = len(relation_queries)
                    attempt = int(chunk_result.get("attempt") or 0)
                    lane_label = str(chunk_result.get("lane_label") or "")

                    await self._update_job_parallel_state(
                        job_id,
                        parallelism=parallelism,
                        active_lanes=self._count_active_lanes(lane_statuses),
                        current_attempt=attempt,
                        max_attempts=self.GRAPH_LLM_MAX_ATTEMPTS,
                        lane_statuses=lane_statuses,
                    )
                    await self._enter_job_stage(
                        job_id,
                        "write_neo4j",
                        (
                            f"正在写入第 {index}/{total_chunks} 段文本到 Neo4j"
                            f"{f'（来源 {lane_label}）' if lane_label else ''}"
                        ),
                        metrics={
                            "total_chars": prepared.get("parsed_text_length") or 0,
                            "processed_chars": processed_chars,
                            "total_chunks": total_chunks,
                            "completed_chunks": completed_chunks,
                            "current_chunk_index": index,
                            "current_chunk_chars": chunk_chars,
                        },
                    )
                    try:
                        if node_queries:
                            await self._heartbeat_job_stage(
                                job_id,
                                detail=f"正在写入第 {index}/{total_chunks} 段文本对应的节点",
                                progress_text=f"正在写入第 {index}/{total_chunks} 段文本对应的节点",
                            )
                            await self._neo4j_client.run_multiple_cypher(node_queries)
                        if relation_queries:
                            await self._heartbeat_job_stage(
                                job_id,
                                detail=f"正在写入第 {index}/{total_chunks} 段文本对应的关系",
                                progress_text=f"正在写入第 {index}/{total_chunks} 段文本对应的关系",
                            )
                            await self._neo4j_client.run_multiple_cypher(relation_queries)
                    except Exception as exc:
                        write_error = f"第 {index}/{total_chunks} 块写入 Neo4j 失败: {exc}"
                        await self._mark_graph_chunk_failed(run_id, index, write_error)
                        await self._set_graph_run_status(run_id, "failed", last_error=write_error)
                        terminal_error = write_error
                        stop_event.set()
                        raise RuntimeError(write_error) from exc

                    await self._mark_graph_chunk_succeeded(
                        run_id,
                        chunk_index=index,
                        entity_count=int(stats.get("entities") or 0),
                        relation_count=int(stats.get("relations") or 0),
                        node_batches=node_batches,
                        relation_batches=relation_batches,
                    )
                    progress = await self._recompute_graph_run_progress(run_id)
                    processed_chars = progress["processed_chars"]
                    completed_chunks = progress["completed_chunks"]
                    total_entities = progress["entity_count"]
                    total_relations = progress["relation_count"]
                    total_node_batches += node_batches
                    total_relation_batches += relation_batches
                    await self._complete_job_stage(
                        job_id,
                        detail=f"已完成第 {index}/{total_chunks} 段文本写入",
                        metrics={
                            "total_chars": prepared.get("parsed_text_length") or 0,
                            "processed_chars": processed_chars,
                            "total_chunks": total_chunks,
                            "completed_chunks": completed_chunks,
                            "current_chunk_index": index,
                            "current_chunk_chars": chunk_chars,
                        },
                    )

                    write_cursor += 1
                    if next_schedule_cursor < len(pending_chunks):
                        await work_queue.put(pending_chunks[next_schedule_cursor])
                        next_schedule_cursor += 1

            worker_tasks = [
                asyncio.create_task(lane_worker(item), name=f"graph-lane-{int(item['lane_id'])}")
                for item in lane_configs
            ]
            try:
                await writer_loop()
            finally:
                stop_event.set()
                for task in worker_tasks:
                    task.cancel()
                if worker_tasks:
                    await asyncio.gather(*worker_tasks, return_exceptions=True)

        await self._set_graph_run_status(run_id, "succeeded", finished=True)
        await self._enter_job_stage(job_id, "finalize", "图谱抽取完成，正在整理结果")
        await self._complete_job_stage(job_id, detail="任务执行完成")

        prepared.update(
            {
                "file_hash": file_hash,
                "run_id": run_id,
                "is_resumed": is_resumed,
                "resume_run_id": run_id if is_resumed else None,
                "resumed_from_chunk": resumed_from_chunk,
                "chunk_limit": chunk_limit,
                "chunk_strategy_version": chunk_strategy_version,
                "cache_used": cache_used,
                "cache_invalidated_reason": cache_invalidated_reason,
                "resume_compatible": resume_compatible,
                "parallelism": parallelism,
                "active_lanes": self._count_active_lanes(lane_statuses),
                "current_attempt": int(max(attempt_counts.values()) if pending_chunks and attempt_counts else 0),
                "max_attempts": self.GRAPH_LLM_MAX_ATTEMPTS,
                "lane_statuses": self._clone_lane_statuses(lane_statuses),
                "processed_text_length": processed_chars,
                "entity_count": total_entities,
                "relation_count": total_relations,
                "node_batches": total_node_batches,
                "relation_batches": total_relation_batches,
                "message": f"抽取并导入完成，实体 {total_entities} 个，关系 {total_relations} 条",
            }
        )
        return prepared

    def _parse_and_split_graph_blocking(  # pragma: no cover - external document parsing runtime
        self,
        file_bytes: bytes,
        filename: str,
        stage_callback: Callable[..., None],
    ) -> dict[str, Any]:
        """在线程中执行文档解析、缓存恢复和文本分块。

        Args:
            file_bytes: 上传文件二进制内容。
            filename: 上传文件名。
            stage_callback: 线程安全阶段回调，用于回传解析和分块进度。

        Returns:
            dict[str, Any]: 解析元数据和内部使用的 `_chunks` 文本块列表。
        """
        suffix = Path(filename).suffix.lower()
        supported_suffixes = {".txt", ".md", ".pdf", ".docx", ".doc", ".pptx", ".ppt", ".html", ".htm", ".jpg", ".jpeg", ".png"}
        if suffix not in supported_suffixes:
            raise RuntimeError(f"不支持的文件格式: {suffix}，仅支持 {', '.join(sorted(supported_suffixes))}")
        chunk_limit = self.GRAPH_IMPORT_CHUNK_LIMIT

        def split_text_by_sentence(text: str, limit: int = self.GRAPH_IMPORT_CHUNK_LIMIT) -> list[str]:
            """按句号类标点尽量切分长文本，超过限制时按长度硬切。"""
            chunks: list[str] = []
            remaining = (text or "").strip()
            sentence_ends = ["。", ".", "!", "！", "?", "？"]
            while remaining:
                if len(remaining) <= limit:
                    chunks.append(remaining)
                    break
                window = remaining[:limit]
                split_at = max(window.rfind(mark) for mark in sentence_ends)
                if split_at < 0:
                    split_at = limit
                else:
                    split_at += 1
                chunk = remaining[:split_at].strip()
                if not chunk:
                    chunk = remaining[:limit]
                    split_at = limit
                chunks.append(chunk)
                remaining = remaining[split_at:].strip()
            return chunks

        def extract_text_fallback(raw_bytes: bytes, raw_filename: str, raw_suffix: str) -> str:
            """Docling 不可用时按文件类型执行内置文本解析。"""
            _ = raw_filename
            if raw_suffix == ".pdf":
                import fitz

                doc = fitz.open(stream=raw_bytes, filetype="pdf")
                try:
                    return "\n".join(page.get_text("text") for page in doc)
                finally:
                    doc.close()
            if raw_suffix == ".docx":
                from io import BytesIO
                from docx import Document

                doc = Document(BytesIO(raw_bytes))
                return "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text and p.text.strip())
            if raw_suffix in {".html", ".htm"}:
                from bs4 import BeautifulSoup

                html = raw_bytes.decode("utf-8", errors="replace")
                return BeautifulSoup(html, "html.parser").get_text(separator="\n", strip=True)
            raise RuntimeError(f"当前环境无法使用 Docling，且该格式无内置回退解析器: {raw_suffix}")

        if not file_bytes:
            raise RuntimeError("上传文件为空")
        if len(file_bytes) > 50 * 1024 * 1024:
            raise RuntimeError("文件大小不能超过 50MB")

        stage_callback("enter", "validate_file", "正在校验文件")
        stage_callback("complete", "validate_file", "文件校验完成")
        file_hash = self._compute_file_hash(file_bytes)
        logger.info(
            "graph_import.parse_prepare file_hash=%s chunk_limit=%s chunk_strategy_version=%s file_name=%s",
            file_hash,
            chunk_limit,
            self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
            filename,
        )

        stage_callback("enter", "parse_document", "正在检查本地缓存")
        cached_prepared, cache_invalidated_reason = self._load_graph_import_cache(file_hash)
        if cached_prepared:
            stage_callback(
                "complete",
                "parse_document",
                f"已从本地缓存恢复解析结果，共 {cached_prepared.get('chunk_count') or 0} 段文本",
            )
            logger.info(
                "graph_import.cache_hit file_hash=%s chunk_limit=%s chunk_count=%s first_chunk_chars=%s",
                file_hash,
                cached_prepared.get("chunk_limit"),
                cached_prepared.get("chunk_count"),
                len((cached_prepared.get("_chunks") or [""])[0]) if cached_prepared.get("_chunks") else 0,
            )
            return cached_prepared
        if cache_invalidated_reason:
            logger.info(
                "graph_import.cache_invalidated file_hash=%s chunk_limit=%s reason=%s",
                file_hash,
                chunk_limit,
                cache_invalidated_reason,
            )

        parser_used = "plain_text"
        parse_warning = ""
        parsed = None
        text = ""

        if suffix in {".txt", ".md"}:
            stage_callback("heartbeat", "parse_document", "正在解析纯文本内容")
            text = file_bytes.decode("utf-8", errors="replace")
            stage_callback("complete", "parse_document", "文档正文解析完成")
        else:
            try:
                from app.document_parser.docling_parser import DoclingParser

                parser = DoclingParser(enable_image_ocr=True, strict_image_ocr=True, progress_callback=stage_callback)
                parsed = parser.parse_bytes(file_bytes, filename)
                text = parsed.text
                parser_used = "docling"
            except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as docling_error:
                if "图片 OCR 失败" in str(docling_error):
                    raise
                self._log_api_error("/admin/graph/jobs/extract-import", "docling_parse_fallback", docling_error, filename=filename)
                stage_callback("heartbeat", "parse_document", "Docling 不可用，正在执行回退解析")
                text = extract_text_fallback(file_bytes, filename, suffix)
                stage_callback("complete", "parse_document", "回退解析完成")
                parser_used = "fallback"
                parse_warning = f"Docling 不可用，已自动回退解析: {docling_error}"

        if not text or not text.strip():
            raise RuntimeError("文件解析后文本为空，无法抽取图谱")

        if parse_warning:
            stage_callback("warning", detail=parse_warning)

        chunks = split_text_by_sentence(text, chunk_limit)
        total_chars = len(text)
        total_chunk_chars = sum(len(chunk) for chunk in chunks)

        ocr_engine = None
        docling_failed_pages: list[int] = []
        pdf_fallback_used = False
        if parsed and parser_used == "docling":
            metadata = parsed.metadata if isinstance(parsed.metadata, dict) else {}
            ocr_engine = metadata.get("ocr_engine")
            docling_failed_pages = list(metadata.get("docling_failed_pages") or [])
            pdf_fallback_used = bool(metadata.get("pdf_fallback_used"))
            if metadata.get("pdf_fallback_used"):
                parser_used = "docling_with_pymupdf_fallback"
                if not parse_warning:
                    parse_warning = str(metadata.get("pdf_fallback_reason") or "")

        response_data: dict[str, Any] = {
            "file_name": filename,
            "parsed_text_length": len(text),
            "total_chunk_chars": total_chunk_chars,
            "processed_text_length": 0,
            "chunk_limit": chunk_limit,
            "chunk_strategy_version": self.GRAPH_IMPORT_CHUNK_STRATEGY_VERSION,
            "parser_used": parser_used,
            "ocr_engine": ocr_engine,
            "parse_warning": parse_warning,
            "docling_failed_pages": docling_failed_pages,
            "pdf_fallback_used": pdf_fallback_used,
            "chunk_count": len(chunks),
            "entity_count": 0,
            "relation_count": 0,
            "node_batches": 0,
            "relation_batches": 0,
            "message": f"文档解析完成，共切分为 {len(chunks)} 段文本",
            "cache_invalidated_reason": cache_invalidated_reason or "",
        }
        if parsed:
            response_data["parse_details"] = {
                "format": parsed.format,
                "pages": parsed.pages,
                "tables_count": len(parsed.tables),
                "images_count": len(parsed.images),
                "metadata": parsed.metadata,
                "text_preview": parsed.text[:2000] if parsed.text else "",
            }
        response_data["_chunks"] = chunks
        response_data["cache_used"] = False
        logger.info(
            "graph_import.chunked file_hash=%s chunk_limit=%s chunk_count=%s first_chunk_chars=%s current_strategy=%s",
            file_hash,
            chunk_limit,
            len(chunks),
            len(chunks[0]) if chunks else 0,
            self._get_chunk_strategy_signature(),
        )
        try:
            self._write_graph_import_cache(file_hash, response_data, chunks)
        except OSError as exc:
            logger.warning("写入图谱导入缓存失败: file_hash=%s error=%s", file_hash, exc)
        return response_data

    def _extract_single_graph_chunk_blocking(  # pragma: no cover - external LLM runtime
        self,
        chunk: str,
        chunk_index: int,
        total_chunks: int,
        total_chars: int,
        processed_chars: int,
        stage_callback: Callable[..., None],
        llm_bundle: dict[str, Any],
    ) -> dict[str, Any]:
        """在线程中执行单块图谱抽取（单次尝试）。

        Args:
            chunk: 当前文本块。
            chunk_index: 当前块序号。
            total_chunks: 总块数。
            total_chars: 文档总字符数。
            processed_chars: 已处理字符数。
            stage_callback: 阶段进度回调。
            llm_bundle: 当前 lane 使用的 LLM 配置和客户端上下文。

        Returns:
            dict[str, Any]: 当前块抽取出的节点查询、关系查询和统计数据。
        """
        from app.knowledge_graph.extract_graph_data_utils import extract_text_to_graph_queries

        chunk_chars = len(chunk)
        stage_callback(
            "heartbeat",
            detail=f"文档解析完成，准备抽取第 {chunk_index}/{total_chunks} 段文本",
            progress_text=f"正在抽取第 {chunk_index}/{total_chunks} 段文本",
            metrics={
                "total_chars": total_chars,
                "processed_chars": processed_chars,
                "total_chunks": total_chunks,
                "completed_chunks": chunk_index - 1,
                "current_chunk_index": chunk_index,
                "current_chunk_chars": chunk_chars,
            },
        )
        graph_data = extract_text_to_graph_queries(
            chunk,
            chunk_chars,
            stage_callback=stage_callback,
            llm_bundle=llm_bundle,
        )
        return {
            "chunk_index": chunk_index,
            "chunk_chars": chunk_chars,
            "node_queries": graph_data["node_queries"],
            "relation_queries": graph_data["relation_queries"],
            "stats": graph_data["stats"],
        }

    @staticmethod
    def _make_lane_stage_callback(  # pragma: no cover - external LLM lane callback
        base_callback: Callable[..., None],
        *,
        lane_id: int,
        chunk_index: int,
        total_chunks: int,
        attempt: int,
        max_attempts: int,
    ) -> Callable[..., None]:
        """为并发 lane 包装带上下文的阶段回调。

        Args:
            base_callback: 原始阶段回调。
            lane_id: 当前 lane 标识。
            chunk_index: 当前块序号。
            total_chunks: 总块数。
            attempt: 当前尝试次数。
            max_attempts: 最大尝试次数。

        Returns:
            Callable[..., None]: 会自动补充 lane、块和尝试次数前缀的回调。
        """

        prefix = f"[lane{lane_id} 第 {chunk_index}/{total_chunks} 块 第 {attempt}/{max_attempts} 次]"

        def _callback(
            action: str,
            stage: Optional[str] = None,
            detail: Optional[str] = None,
            progress_text: Optional[str] = None,
            metrics: Optional[dict[str, Any]] = None,
        ) -> None:
            next_detail = f"{prefix} {detail}" if detail else prefix
            next_progress = f"{prefix} {progress_text}" if progress_text else next_detail
            base_callback(
                action,
                stage=stage,
                detail=next_detail,
                progress_text=next_progress,
                metrics=metrics,
            )

        return _callback

    async def _embed_graph_to_milvus_sync(self, data: Optional[dict[str, Any]], job_id: str) -> dict[str, Any]:  # pragma: no cover - external Milvus runtime
        """执行图谱嵌入主体，从 Neo4j 读取数据并重建 Milvus 索引。

        Args:
            data: 可选标签过滤和批大小配置。
            job_id: 进程内后台任务标识。

        Returns:
            dict[str, Any]: Milvus 重建统计和完成时间。
        """
        await self._enter_job_stage(job_id, "read_graph", "准备读取图谱数据")
        data = data or {}
        labels = data.get("labels")
        batch_size = data.get("batch_size", 100)
        kg_store = self._get_shared_kg_milvus_store()
        loop = asyncio.get_running_loop()

        def progress_callback(text: str) -> None:
            """将 Milvus 重建进度回传到后台任务心跳。"""
            try:
                asyncio.run_coroutine_threadsafe(
                    self._heartbeat_job_stage(job_id, detail=str(text or ""), progress_text=str(text or "")),
                    loop,
                )
            except RuntimeError:
                return

        stats = await kg_store.rebuild_from_neo4j(
            self._neo4j_client,
            labels,
            batch_size,
            progress_callback=progress_callback,
        )
        await self._complete_job_stage(job_id, detail="Milvus 嵌入完成")
        await self._enter_job_stage(job_id, "finalize", "Milvus 嵌入已完成，正在整理结果")
        await self._complete_job_stage(job_id, detail="任务执行完成")
        return {
            "message": f"知识图谱嵌入成功，共 {stats['total_documents']} 条记录",
            "stats": stats,
            "timestamp": stats.get("timestamp"),
        }

    async def get_graph_labels(self) -> dict[str, Any]:
        """获取 Neo4j 中所有节点标签。

        Returns:
            dict[str, Any]: 标签列表或查询错误响应。
        """
        try:
            query = """
            MATCH (n)
            UNWIND labels(n) AS label
            RETURN DISTINCT label
            ORDER BY label
            """
            result = await self._await_call(self._neo4j_client.run_cypher, query)
            labels = [record["label"] for record in result]
            return {"ok": True, "data": labels}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/labels", "neo4j_query", exc)
            return {"ok": False, "error": str(exc)}

    async def get_graph_relationship_types(self) -> dict[str, Any]:
        """获取 Neo4j 中所有关系类型。

        Returns:
            dict[str, Any]: 关系类型列表或查询错误响应。
        """
        try:
            query = """
            MATCH ()-[r]->()
            RETURN DISTINCT type(r) AS rel_type
            ORDER BY rel_type
            """
            result = await self._await_call(self._neo4j_client.run_cypher, query)
            rel_types = [record["rel_type"] for record in result]
            return {"ok": True, "data": rel_types}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/relationship-types", "neo4j_query", exc)
            return {"ok": False, "error": str(exc)}

    async def get_graph_nodes(self, label: str = "", keyword: str = "", limit: int = 100, offset: int = 0) -> dict[str, Any]:
        """按标签、关键词和分页参数获取图谱节点列表。

        Args:
            label: 可选节点标签过滤条件。
            keyword: 可选名称/标题关键词。
            limit: 返回数量上限。
            offset: 分页偏移量。

        Returns:
            dict[str, Any]: 节点列表、总数或查询错误响应。
        """
        try:
            if label:
                if keyword:
                    query = f"""
                    MATCH (n:{label})
                    WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                    RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                    ORDER BY n.name
                    SKIP $offset LIMIT $limit
                    """
                    count_query = f"""
                    MATCH (n:{label})
                    WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                    RETURN count(n) AS total
                    """
                else:
                    query = f"""
                    MATCH (n:{label})
                    RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                    ORDER BY n.name
                    SKIP $offset LIMIT $limit
                    """
                    count_query = f"""
                    MATCH (n:{label})
                    RETURN count(n) AS total
                    """
            else:
                if keyword:
                    query = """
                    MATCH (n)
                    WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                    RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                    ORDER BY n.name
                    SKIP $offset LIMIT $limit
                    """
                    count_query = """
                    MATCH (n)
                    WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                    RETURN count(n) AS total
                    """
                else:
                    query = """
                    MATCH (n)
                    RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                    ORDER BY n.name
                    SKIP $offset LIMIT $limit
                    """
                    count_query = """
                    MATCH (n)
                    RETURN count(n) AS total
                    """

            params = {"keyword": keyword, "offset": offset, "limit": limit}
            result = await self._await_call(self._neo4j_client.run_cypher, query, params)
            count_result = await self._await_call(self._neo4j_client.run_cypher, count_query, {"keyword": keyword} if keyword else {})
            total = count_result[0]["total"] if count_result else 0
            return {"ok": True, "data": result, "total": total}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/nodes", "neo4j_query", exc, label=label, keyword=keyword)
            return {"ok": False, "error": str(exc)}

    async def get_graph_node(self, node_id: int) -> dict[str, Any]:
        """获取单个节点详情及其入边、出边。

        Args:
            node_id: Neo4j 内部节点 ID。

        Returns:
            dict[str, Any]: 节点详情、关系详情或节点不存在错误。
        """
        try:
            node_query = """
            MATCH (n)
            WHERE id(n) = $node_id
            RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
            """
            node_result = await self._await_call(self._neo4j_client.run_cypher, node_query, {"node_id": node_id})
            if not node_result:
                return {"ok": False, "error": "节点不存在"}

            incoming_query = """
            MATCH (n)<-[r]-(m)
            WHERE id(n) = $node_id
            RETURN id(r) AS rel_id, type(r) AS rel_type, properties(r) AS rel_props,
                   id(m) AS source_id, labels(m) AS source_labels, m.name AS source_name
            """
            incoming = await self._await_call(self._neo4j_client.run_cypher, incoming_query, {"node_id": node_id})

            outgoing_query = """
            MATCH (n)-[r]->(m)
            WHERE id(n) = $node_id
            RETURN id(r) AS rel_id, type(r) AS rel_type, properties(r) AS rel_props,
                   id(m) AS target_id, labels(m) AS target_labels, m.name AS target_name
            """
            outgoing = await self._await_call(self._neo4j_client.run_cypher, outgoing_query, {"node_id": node_id})

            return {
                "ok": True,
                "data": {
                    "node": node_result[0],
                    "relationships": {
                        "incoming": incoming,
                        "outgoing": outgoing,
                    },
                },
            }
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/nodes/{node_id}", "neo4j_query", exc, node_id=node_id)
            return {"ok": False, "error": str(exc)}

    async def create_graph_node(self, data: dict[str, Any]) -> dict[str, Any]:
        """创建新的图谱节点。

        Args:
            data: 包含 label 和 properties 的节点创建 payload。

        Returns:
            dict[str, Any]: 创建后的节点 ID 或校验/写入错误。
        """
        try:
            raw_label = str(data.get("label", "")).strip()
            label = self._validate_cypher_identifier(raw_label)
            properties = self._sanitize_graph_properties(data.get("properties", {}))

            if not raw_label:
                return {"ok": False, "error": "节点标签不能为空"}
            if not label:
                return {"ok": False, "error": "节点标签非法，仅允许字母数字下划线且不能以数字开头"}

            if properties:
                props_str = ", ".join([f"`{k}`: ${k}" for k in properties.keys()])
                query = f"CREATE (n:`{label}` {{{props_str}}}) RETURN id(n) AS id"
            else:
                query = f"CREATE (n:`{label}`) RETURN id(n) AS id"

            result = await self._await_call(self._neo4j_client.run_cypher, query, properties)
            if result:
                return {"ok": True, "data": {"id": result[0]["id"], "message": "节点创建成功"}}
            return {"ok": False, "error": "创建失败"}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/nodes", "create_node", exc)
            return {"ok": False, "error": str(exc)}

    async def update_graph_node(self, node_id: int, data: dict[str, Any]) -> dict[str, Any]:
        """更新指定图谱节点属性。

        Args:
            node_id: Neo4j 内部节点 ID。
            data: 包含 properties 的节点更新 payload。

        Returns:
            dict[str, Any]: 更新结果或节点不存在错误。
        """
        try:
            properties = self._sanitize_graph_properties(data.get("properties", {}))
            if not properties:
                return {"ok": False, "error": "属性不能为空"}

            set_clauses = [f"n.`{k}` = ${k}" for k in properties.keys()]
            set_str = ", ".join(set_clauses)
            query = f"""
            MATCH (n)
            WHERE id(n) = $node_id
            SET {set_str}
            RETURN id(n) AS id, properties(n) AS properties
            """
            params = {"node_id": node_id, **properties}
            result = await self._await_call(self._neo4j_client.run_cypher, query, params)
            if result:
                return {"ok": True, "data": {"id": result[0]["id"], "message": "节点更新成功"}}
            return {"ok": False, "error": "节点不存在或更新失败"}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/nodes/{node_id}", "update_node", exc, node_id=node_id)
            return {"ok": False, "error": str(exc)}

    async def delete_graph_node(self, node_id: int) -> dict[str, Any]:
        """删除指定图谱节点及其关联关系。

        Args:
            node_id: Neo4j 内部节点 ID。

        Returns:
            dict[str, Any]: 删除结果或节点不存在错误。
        """
        try:
            query = """
            MATCH (n)
            WHERE id(n) = $node_id
            DETACH DELETE n
            RETURN count(n) AS deleted_count
            """
            result = await self._await_call(self._neo4j_client.run_cypher, query, {"node_id": node_id})
            if result and result[0]["deleted_count"] > 0:
                return {"ok": True, "message": "节点删除成功"}
            return {"ok": False, "error": "节点不存在"}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/nodes/{node_id}", "delete_node", exc, node_id=node_id)
            return {"ok": False, "error": str(exc)}

    async def create_graph_relationship(self, data: dict[str, Any]) -> dict[str, Any]:
        """创建两个节点之间的图谱关系。

        Args:
            data: 包含 source_id、target_id、rel_type 和 properties 的关系创建 payload。

        Returns:
            dict[str, Any]: 创建后的关系 ID 或校验/写入错误。
        """
        try:
            source_id = data.get("source_id")
            target_id = data.get("target_id")
            raw_rel_type = str(data.get("rel_type", "")).strip()
            rel_type = self._validate_cypher_identifier(raw_rel_type)
            properties = self._sanitize_graph_properties(data.get("properties", {}))

            if source_id is None or target_id is None or not raw_rel_type:
                return {"ok": False, "error": "源节点、目标节点和关系类型不能为空"}
            if not rel_type:
                return {"ok": False, "error": "关系类型非法，仅允许字母数字下划线且不能以数字开头"}

            try:
                source_id_int = int(source_id)
                target_id_int = int(target_id)
            except (TypeError, ValueError) as exc:
                self._log_api_error(
                    "/admin/graph/relationships",
                    "create_relationship_id_parse",
                    exc,
                    source_id=source_id,
                    target_id=target_id,
                )
                return {"ok": False, "error": "源节点和目标节点ID必须是整数"}

            if properties:
                props_str = "{ " + ", ".join([f"`{k}`: ${k}" for k in properties.keys()]) + " }"
                query = f"""
                MATCH (source), (target)
                WHERE id(source) = $source_id AND id(target) = $target_id
                CREATE (source)-[r:`{rel_type}` {props_str}]->(target)
                RETURN id(r) AS rel_id
                """
                params = {"source_id": source_id_int, "target_id": target_id_int, **properties}
            else:
                query = f"""
                MATCH (source), (target)
                WHERE id(source) = $source_id AND id(target) = $target_id
                CREATE (source)-[r:`{rel_type}`]->(target)
                RETURN id(r) AS rel_id
                """
                params = {"source_id": source_id_int, "target_id": target_id_int}

            result = await self._await_call(self._neo4j_client.run_cypher, query, params)
            if result:
                return {"ok": True, "data": {"rel_id": result[0]["rel_id"], "message": "关系创建成功"}}
            return {"ok": False, "error": "创建失败，请检查节点是否存在"}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/relationships", "create_relationship", exc)
            return {"ok": False, "error": str(exc)}

    async def delete_graph_relationship(self, rel_id: int) -> dict[str, Any]:
        """删除指定图谱关系。

        Args:
            rel_id: Neo4j 内部关系 ID。

        Returns:
            dict[str, Any]: 删除结果或关系不存在错误。
        """
        try:
            query = """
            MATCH ()-[r]->()
            WHERE id(r) = $rel_id
            DELETE r
            RETURN count(r) AS deleted_count
            """
            result = await self._await_call(self._neo4j_client.run_cypher, query, {"rel_id": rel_id})
            if result and result[0]["deleted_count"] > 0:
                return {"ok": True, "message": "关系删除成功"}
            return {"ok": False, "error": "关系不存在"}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/relationships/{rel_id}", "delete_relationship", exc, rel_id=rel_id)
            return {"ok": False, "error": str(exc)}

    async def execute_cypher(self, data: dict[str, Any]) -> dict[str, Any]:
        """执行管理员输入的只读 Cypher 查询。

        Args:
            data: 包含 query 字段的查询 payload。

        Returns:
            dict[str, Any]: 查询结果或安全校验错误。
        """
        try:
            query = str(data.get("query") or "").strip()
            if not query:
                return {"ok": False, "error": "查询语句不能为空"}

            forbidden_keywords = ["create", "merge", "delete", "remove", "set", "drop", "alter"]
            query_lower = query.lower()
            for keyword in forbidden_keywords:
                if keyword in query_lower:
                    return {"ok": False, "error": f"出于安全考虑，不允许执行包含 '{keyword}' 的操作"}

            result = await self._await_call(self._neo4j_client.run_cypher, query)
            return {"ok": True, "data": result}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/execute-cypher", "execute_cypher", exc)
            return {"ok": False, "error": str(exc)}

    async def extract_and_import_graph(self, file: UploadFile) -> dict[str, Any]:
        """兼容旧接口，提交文件抽取导入后台任务。

        Args:
            file: FastAPI 上传文件对象。

        Returns:
            dict[str, Any]: 图谱抽取导入任务提交结果。
        """
        return await self.submit_graph_extract_import_job(file=file)

    async def embed_graph_to_milvus(self, data: Optional[dict[str, Any]] = None) -> dict[str, Any]:
        """兼容旧接口，提交图谱嵌入 Milvus 后台任务。

        Args:
            data: 可选标签和批大小等嵌入参数。

        Returns:
            dict[str, Any]: 图谱嵌入任务提交结果。
        """
        return await self.submit_graph_embed_milvus_job(data=data)

    async def get_graph_embed_status(self) -> dict[str, Any]:
        """获取知识图谱向量库嵌入状态。

        Returns:
            dict[str, Any]: 图谱向量库统计信息或错误响应。
        """
        try:
            kg_store = await self._await_call(self._get_shared_kg_milvus_store)
            stats = await self._await_call(kg_store.get_stats)
            return {"ok": True, "data": stats}
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            self._log_api_error("/admin/graph/embed-status", "embed_status", exc)
            return {"ok": False, "error": str(exc)}
