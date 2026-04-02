"""main 模块。

提供 main 相关能力与对外接口。
"""


# ruff: noqa: E402

import asyncio
from datetime import datetime
import inspect
import json
import os
import re
import time
import uuid
from contextlib import asynccontextmanager
from pathlib import Path
import logging
from collections import defaultdict
import threading
from typing import Any, AsyncIterator, Callable, Dict, List, Optional

# 首先加载环境变量，确保后续模块能读取到配置
from dotenv import load_dotenv
project_root = Path(__file__).resolve().parent.parent.parent
env_path = project_root / ".env"
load_dotenv(dotenv_path=env_path, override=True)

from fastapi import FastAPI, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from app.a2a import A2AService, all_agent_cards
from app.backend.repositories.database import SessionManager, MessageManager, SessionMemoryManager, fetch_one, initialize_database
from app.backend.services import SessionMemoryService
from app.common.auth_simple import (
    get_current_admin_user,
    is_logged_in,
    logout,
    reset_current_admin_token,
    set_current_admin_token,
    verify_login,
)
from fastapi import UploadFile, File, Form
from app.a2a.agent.gating_mechanism.evidence_evaluator import evaluate_and_gate
from app.a2a.agent import GraphSubAgent
from app.knowledge_graph.neo4j_manager import neo4j_client
from app.a2a.agent import process_query
from app.a2a.agent import RagSubAgent
from app.common.db_manager import check_databases
from app.mcp.tools import get_mcp_runtime_status, get_mcp_server
from app.common.perf_logger import log_duration
from app.common.baseline_report import generate_api_baseline_report
from app.common.feature_flags import get_feature_flag, get_runtime_flag_snapshot
from app.common.langsmith_tracing import set_trace_metadata
from app.common.request_rate_limiter import InMemorySlidingWindowRateLimiter
from app.common.session_logger import (
    create_request_id,
    ensure_session_meta,
    log_api_call,
    reset_request_context,
    reset_session_context,
    set_request_context,
    set_session_context,
)
from app.rag_core.runtime_resources import (
    get_shared_kg_milvus_store,
    get_shared_vector_store,
)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def _log_api_error(endpoint: str, stage: str, exc: Exception, **kwargs: Any) -> None:
    """记录 API 异常并输出结构化日志。

    Args:
        endpoint: 接口路径标识。
        stage: 当前处理阶段标识。
        exc: 捕获到的异常对象。
        **kwargs: 补充上下文字段。

    Returns:
        None: 无返回值。
    """
    payload = {"endpoint": endpoint, "stage": stage, "error_type": type(exc).__name__, "error": str(exc)}
    payload.update(kwargs)
    logger.exception("[api_error] %s", json.dumps(payload, ensure_ascii=False))

app: FastAPI

# MCP 服务器实例（延迟初始化）
_mcp_server = None
_a2a_service: Optional[A2AService] = None
_a2a_graph_service: Optional[A2AService] = None
_a2a_retrieval_service: Optional[A2AService] = None
_a2a_evaluator_service: Optional[A2AService] = None
_graph_subagent: Optional[GraphSubAgent] = None
_rag_subagent: Optional[RagSubAgent] = None
_mysql_initialized: Optional[bool] = None
_CYPHER_IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,63}$")
_rate_limiter: Optional[InMemorySlidingWindowRateLimiter] = None
_session_memory_service: Optional[SessionMemoryService] = None


def _env_bool(name: str, default: bool) -> bool:
    raw = str(os.getenv(name, "1" if default else "0")).strip().lower()
    if raw in {"1", "true", "yes", "on"}:
        return True
    if raw in {"0", "false", "no", "off"}:
        return False
    return default


@asynccontextmanager
async def _app_lifespan(_: FastAPI):
    """应用生命周期（替代 on_event）。"""
    global _mysql_initialized, _a2a_service, _a2a_graph_service, _a2a_retrieval_service, _a2a_evaluator_service

    print("\n" + "=" * 50)
    print("正在初始化 MySQL 数据库...")
    print("=" * 50)
    _mysql_initialized = bool(await initialize_database())
    if not _mysql_initialized:
        print("⚠ MySQL 数据库初始化失败，部分功能可能不可用")

    print("\n" + "=" * 50)
    print("正在检查数据库连接...")
    print("=" * 50)
    status = await check_databases()
    if bool(status.get("all_running")):
        print("✓ Neo4j 运行正常")
        print("✓ Milvus 运行正常")
    else:
        print("⚠ 检测到部分数据库未启动，请手动启动后再使用相关功能")
    print("=" * 50 + "\n")

    print("正在预热共享检索资源...")
    try:
        with log_duration("startup.prewarm_vector_store"):
            get_shared_vector_store()
        kg_store = get_shared_kg_milvus_store()
        if _env_bool("KG_ENTITY_RELATION_PRELOAD_EMBEDDER", True):
            with log_duration("startup.prewarm_kg_embedder"):
                kg_store.ensure_embedder_loaded()
        print("✓ 共享检索资源预热完成")
    except Exception as exc:
        _log_api_error("startup", "retrieval_prewarm", exc)
        print(f"⚠ 共享检索资源预热失败: {exc}")

    print("正在初始化 MCP 服务...")
    try:
        mcp = get_mcp()
        tools = mcp.list_tools()
        print(f"✓ MCP 服务初始化成功，已注册 {len(tools)} 个工具:")
        for tool in tools:
            print(f"  - {tool['name']}: {tool['description']}")
    except Exception as exc:
        _log_api_error("startup", "mcp_init", exc)
        print(f"⚠ MCP 服务初始化失败: {exc}")

    print("正在初始化 A2A 服务...")
    try:
        a2a_service = get_a2a()
        card = a2a_service.get_agent_card()
        print(
            "✓ A2A 服务初始化成功: "
            f"name={card.get('name', '')}, protocol={card.get('protocolVersion', '')}"
        )
        graph_card = get_a2a_graph().get_agent_card()
        retrieval_card = get_a2a_retrieval().get_agent_card()
        evaluator_card = get_a2a_evaluator().get_agent_card()
        print(
            "✓ A2A 子代理已初始化: "
            f"{graph_card.get('name', '')}, {retrieval_card.get('name', '')}, {evaluator_card.get('name', '')}"
        )
    except Exception as exc:
        _log_api_error("startup", "a2a_init", exc)
        print(f"⚠ A2A 服务初始化失败: {exc}")

    flags = get_runtime_flag_snapshot()
    print("当前运行时开关:")
    for flag_name, flag_value in flags.items():
        print(f"  - {flag_name}={flag_value}")
    print("=" * 50 + "\n")

    try:
        yield
    finally:
        if _a2a_service is not None:
            await _a2a_service.ashutdown()
            _a2a_service = None
        if _a2a_graph_service is not None:
            await _a2a_graph_service.ashutdown()
            _a2a_graph_service = None
        if _a2a_retrieval_service is not None:
            await _a2a_retrieval_service.ashutdown()
            _a2a_retrieval_service = None
        if _a2a_evaluator_service is not None:
            await _a2a_evaluator_service.ashutdown()
            _a2a_evaluator_service = None


app = FastAPI(lifespan=_app_lifespan)


def _read_int_env(name: str, default: int, min_value: int, max_value: int) -> int:
    """读取整型环境变量并限制取值范围。

    Args:
        name: 环境变量名称。
        default: 默认值。
        min_value: 允许的最小值。
        max_value: 允许的最大值。

    Returns:
        int: 限幅后的整型值。
    """
    raw = str(os.getenv(name, str(default))).strip()
    try:
        value = int(raw)
    except ValueError:
        value = default
    if value < min_value:
        return min_value
    if value > max_value:
        return max_value
    return value


def _read_float_env(name: str, default: float, min_value: float, max_value: float) -> float:
    """读取浮点环境变量并限制取值范围。

    Args:
        name: 环境变量名称。
        default: 默认值。
        min_value: 允许的最小值。
        max_value: 允许的最大值。

    Returns:
        float: 限幅后的浮点值。
    """
    raw = str(os.getenv(name, str(default))).strip()
    try:
        value = float(raw)
    except ValueError:
        value = default
    if value < min_value:
        return min_value
    if value > max_value:
        return max_value
    return value


_RATE_LIMIT_WINDOW_SECONDS = _read_float_env(
    "HEAVY_API_RATE_LIMIT_WINDOW_SECONDS",
    60.0,
    1.0,
    3600.0,
)
_rate_limiter = InMemorySlidingWindowRateLimiter(
    max_keys=_read_int_env("HEAVY_API_RATE_LIMIT_MAX_KEYS", 20000, 1000, 200000),
)

_HEAVY_ENDPOINT_LIMITS: dict[str, dict[str, int]] = {
    "/agent/run": {
        "user_limit": _read_int_env("RATE_LIMIT_AGENT_RUN_USER", 12, 1, 5000),
        "session_limit": _read_int_env("RATE_LIMIT_AGENT_RUN_SESSION", 18, 1, 5000),
    },
    "/process_with_session": {
        "user_limit": _read_int_env("RATE_LIMIT_PROCESS_SESSION_USER", 20, 1, 5000),
        "session_limit": _read_int_env("RATE_LIMIT_PROCESS_SESSION_SESSION", 30, 1, 5000),
    },
    "/mcp/invoke": {
        "user_limit": _read_int_env("RATE_LIMIT_MCP_INVOKE_USER", 60, 1, 20000),
        "session_limit": _read_int_env("RATE_LIMIT_MCP_INVOKE_SESSION", 80, 1, 20000),
    },
}


def _extract_bearer_token(authorization: str) -> Optional[str]:
    """从 Authorization 头提取 Bearer Token。

    Args:
        authorization: HTTP Authorization 头原始值。

    Returns:
        Optional[str]: 解析出的 token；无效时返回 None。
    """
    auth_text = str(authorization or "").strip()
    if not auth_text:
        return None
    parts = auth_text.split(" ", 1)
    if len(parts) != 2 or parts[0].lower() != "bearer":
        return None
    token = parts[1].strip()
    if not token:
        return None
    return token


async def _is_session_owner(session_id: str, user_id: str) -> bool:
    """校验指定会话是否属于目标用户。

    Args:
        session_id: 会话 ID。
        user_id: 用户 ID。

    Returns:
        bool: 是否为会话所有者。
    """
    session = await _await_call(SessionManager.get_session, session_id)
    if not session:
        return False
    return str(session.get("user_id") or "") == user_id


def _validate_cypher_identifier(raw_name: Any) -> Optional[str]:
    """校验 Cypher 标识符是否合法。

    Args:
        raw_name: 待校验的原始标识符。

    Returns:
        Optional[str]: 合法标识符或 None。
    """
    name = str(raw_name).strip()
    if not name:
        return None
    if not _CYPHER_IDENTIFIER_PATTERN.match(name):
        return None
    return name


def _sanitize_graph_properties(raw_properties: Any) -> Dict[str, Any]:
    """清洗并校验图谱属性键。

    Args:
        raw_properties: 待清洗的属性对象。

    Returns:
        Dict[str, Any]: 清洗后的属性字典。
    """
    if raw_properties is None:
        return {}
    if not isinstance(raw_properties, dict):
        raise ValueError("属性必须是对象")

    sanitized: Dict[str, Any] = {}
    for key, value in raw_properties.items():
        safe_key = _validate_cypher_identifier(key)
        if not safe_key:
            raise ValueError(f"非法属性键: {key}")
        sanitized[safe_key] = value
    return sanitized


def _resolve_trace_id_from_request(request: Optional[Request]) -> str:
    """从请求上下文解析 trace_id。

    Args:
        request: FastAPI 请求对象。

    Returns:
        str: trace_id 字符串。
    """
    if request is None:
        return ""
    return str(getattr(request.state, "request_id", "") or "")


def _enforce_heavy_api_rate_limit(
    endpoint: str,
    user_id: str = "",
    session_id: str = "",
    request: Optional[Request] = None,
) -> Optional[JSONResponse]:
    """执行重接口限流检查。

    Args:
        endpoint: 接口路径标识。
        user_id: 用户 ID。
        session_id: 会话 ID。
        request: FastAPI 请求对象。

    Returns:
        Optional[JSONResponse]: 触发限流时返回响应，否则为 None。
    """
    if not get_feature_flag("ENABLE_HEAVY_API_RATE_LIMIT", True):
        return None
    if _rate_limiter is None:
        return None
    config = _HEAVY_ENDPOINT_LIMITS.get(endpoint)
    if not config:
        return None

    trace_id = _resolve_trace_id_from_request(request)
    safe_user_id = str(user_id or "").strip()
    safe_session_id = str(session_id or "").strip()
    checks: list[tuple[str, str, int]] = []

    user_limit = int(config.get("user_limit") or 0)
    session_limit = int(config.get("session_limit") or 0)
    if safe_user_id and user_limit > 0:
        checks.append(("user_id", safe_user_id, user_limit))
    if safe_session_id and session_limit > 0:
        checks.append(("session_id", safe_session_id, session_limit))

    for dimension, value, limit in checks:
        key = f"{endpoint}:{dimension}:{value}"
        result = _rate_limiter.check(
            key=key,
            limit=limit,
            window_seconds=_RATE_LIMIT_WINDOW_SECONDS,
        )
        if bool(result.get("allowed")):
            continue
        retry_after_ms = int(result.get("retry_after_ms") or 0)
        return JSONResponse(
            status_code=429,
            content={
                "ok": False,
                "error": "rate_limited",
                "endpoint": endpoint,
                "dimension": dimension,
                "retry_after_ms": retry_after_ms,
                "trace_id": trace_id,
                "limit": int(result.get("limit") or limit),
                "window_seconds": float(result.get("window_seconds") or _RATE_LIMIT_WINDOW_SECONDS),
            },
        )
    return None

def get_mcp():
    """获取或初始化 MCP 服务器"""
    global _mcp_server
    if _mcp_server is None:
        _mcp_server = get_mcp_server()
    return _mcp_server


def get_session_memory_service() -> SessionMemoryService:
    """获取或初始化会话记忆服务。"""
    global _session_memory_service
    if _session_memory_service is None:
        _session_memory_service = SessionMemoryService()
    return _session_memory_service


async def _run_a2a_query(
    input_text: str,
    context_id: str = "",
    user_id: str = "",
    metadata: Optional[Dict[str, Any]] = None,
    cancel_check: Optional[Callable[[], bool]] = None,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
) -> Dict[str, Any]:
    """A2A 任务执行适配：复用现有 process_query 主链路。"""
    meta = metadata if isinstance(metadata, dict) else {}
    session_id = str(context_id or meta.get("session_id") or "").strip()
    safe_user_id = str(user_id or meta.get("user_id") or "").strip()
    strategy = str(meta.get("strategy") or "iterative").strip().lower() or "iterative"
    raw_max_turns = meta.get("max_turns")
    max_turns: Optional[int] = (
        raw_max_turns
        if isinstance(raw_max_turns, int) and raw_max_turns > 0
        else None
    )
    original_input_text = str(meta.get("original_input_text") or input_text or "").strip()
    resolved_query_obj = meta.get("resolved_query")
    resolved_query = resolved_query_obj if isinstance(resolved_query_obj, dict) else {}
    memory_context_obj = meta.get("memory_context")
    memory_context = memory_context_obj if isinstance(memory_context_obj, dict) else {}

    return await process_query(
        input_text=str(input_text or ""),
        session_id=session_id,
        user_id=safe_user_id,
        strategy=strategy,
        max_turns=max_turns,
        original_input_text=original_input_text,
        resolved_query=resolved_query,
        memory_context=memory_context,
        cancel_check=cancel_check,
        progress_callback=progress_callback,
    )


def get_a2a() -> A2AService:
    """获取或初始化 A2A 服务。"""
    global _a2a_service
    if _a2a_service is None:
        _a2a_service = A2AService(
            run_query=_run_a2a_query,
            agent_name="Hypertension Orchestrator Agent",
            agent_description="高血压问答多 Agent 编排代理（A2A 适配）",
            protocol_version="0.2.5",
            agent_url="/a2a/rpc",
            events_base_path="/a2a/tasks",
        )
    return _a2a_service


def _safe_float(value: Any, default: float = 0.0) -> float:
    """安全转换为浮点数。

    Args:
        value: 待规范化值。
        default: 默认值。

    Returns:
        float: 转换后的浮点值；失败时返回 default。
    """
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _extract_text_from_parts(parts_obj: Any) -> str:
    """从消息 parts 中提取文本。

    Args:
        parts_obj: parts 字段原始对象（期望为列表）。

    Returns:
        str: 拼接后的文本内容。
    """
    parts = parts_obj if isinstance(parts_obj, list) else []
    chunks: List[str] = []
    for part in parts:
        if not isinstance(part, dict):
            continue
        if str(part.get("kind") or "").strip().lower() != "text":
            continue
        text = str(part.get("text") or "").strip()
        if text:
            chunks.append(text)
    return "\n".join(chunks).strip()


def _extract_a2a_artifact(task: Dict[str, Any], artifact_name: str) -> Any:
    """从 A2A 任务中提取指定产物。

    Args:
        task: 任务对象字典。
        artifact_name: 目标产物名称。

    Returns:
        Any: 命中时返回 data/text 对应内容，未命中返回 None。
    """
    artifacts_obj = task.get("artifacts")
    artifacts = artifacts_obj if isinstance(artifacts_obj, list) else []
    for artifact in artifacts:
        if not isinstance(artifact, dict):
            continue
        if str(artifact.get("name") or "").strip() != artifact_name:
            continue
        parts_obj = artifact.get("parts")
        parts = parts_obj if isinstance(parts_obj, list) else []
        for part in parts:
            if not isinstance(part, dict):
                continue
            kind = str(part.get("kind") or "").strip().lower()
            if kind == "data":
                return part.get("data")
            if kind == "text":
                return str(part.get("text") or "")
    return None


def _extract_a2a_status_message(task: Dict[str, Any]) -> Dict[str, Any]:
    """提取任务状态消息对象。

    Args:
        task: 任务对象字典。

    Returns:
        Dict[str, Any]: 标准化后的 status.message 字典。
    """
    status_obj = task.get("status")
    status = status_obj if isinstance(status_obj, dict) else {}
    message_obj = status.get("message")
    message = message_obj if isinstance(message_obj, dict) else {}
    return message


def _build_traces_from_processing_steps(steps_obj: Any) -> List[Dict[str, Any]]:
    """将处理步骤转换为 trace 列表。

    Args:
        steps_obj: 处理步骤对象（通常为列表）。

    Returns:
        List[Dict[str, Any]]: 结构化 trace 列表。
    """
    steps = steps_obj if isinstance(steps_obj, list) else []
    rows: List[Dict[str, Any]] = []
    for idx, step in enumerate(steps, start=1):
        rows.append(
            {
                "step": idx,
                "action": str(step or ""),
            }
        )
    return rows


def _build_processing_steps_from_progress_events(events_obj: Any) -> List[str]:
    """从进度事件生成处理步骤文本。

    Args:
        events_obj: 进度事件对象（通常为列表）。

    Returns:
        List[str]: 提取出的步骤文案列表。
    """
    events = events_obj if isinstance(events_obj, list) else []
    rows: List[str] = []
    for event in events:
        if not isinstance(event, dict):
            continue
        detail = str(event.get("detail") or event.get("event") or "").strip()
        if detail:
            rows.append(detail)
    return rows


def _resolve_agent_run_mode(payload: Dict[str, Any]) -> str:
    """解析 agent 运行模式。

    Args:
        payload: 请求负载字典。

    Returns:
        str: 归一化后的运行模式。
    """
    mode = str(payload.get("a2a_mode") or "").strip().lower()
    if not mode:
        return "auto"
    if mode in {"auto", "force", "prefer"}:
        return mode
    if mode in {"off", "legacy"}:
        raise ValueError("a2a_mode_not_supported")
    raise ValueError("a2a_mode_invalid")


async def _run_agent_via_a2a(
    input_text: str,
    session_id: str,
    user_id: str,
    strategy: str,
    max_turns: Optional[int],
    trace_id: str,
) -> Dict[str, Any]:
    """通过 A2A 服务执行一次 Agent 调用。

    Args:
        input_text: 用户输入文本。
        session_id: 会话 ID。
        user_id: 用户 ID。
        strategy: 策略名称。
        max_turns: 最大迭代轮数。
        trace_id: 请求追踪 ID。

    Returns:
        Dict[str, Any]: A2A 执行结果摘要。
    """
    service = get_a2a()
    metadata: Dict[str, Any] = {
        "session_id": str(session_id or ""),
        "user_id": str(user_id or ""),
        "strategy": str(strategy or "iterative"),
        "trace_id": str(trace_id or ""),
        "source": "agent_run",
    }
    if isinstance(max_turns, int) and max_turns > 0:
        metadata["max_turns"] = max_turns

    message = {
        "role": "user",
        "parts": [{"kind": "text", "text": str(input_text or "")}],
        "metadata": metadata,
        "contextId": str(session_id or ""),
        "userId": str(user_id or ""),
    }
    task = await service.asend_message(message, trace_id)
    task_id = str((task or {}).get("id") or "").strip()
    if not task_id:
        raise RuntimeError("a2a_task_id_missing")

    timeout_seconds = max(
        3.0,
        min(
            _safe_float(os.getenv("AGENT_RUN_A2A_TIMEOUT_SECONDS", "180"), default=180.0),
            900.0,
        ),
    )
    poll_ms = max(
        100,
        min(
            int(_safe_float(os.getenv("AGENT_RUN_A2A_POLL_MS", "250"), default=250.0)),
            3000,
        ),
    )
    started = time.time()

    while True:
        snapshot = await service.aget_task(task_id)
        if not snapshot:
            raise RuntimeError("a2a_task_not_found")
        status_obj = snapshot.get("status")
        status = status_obj if isinstance(status_obj, dict) else {}
        state = str(status.get("state") or "").strip().lower()
        if state in {"completed", "failed", "canceled"}:
            break
        if time.time() - started >= timeout_seconds:
            raise TimeoutError(f"a2a_task_timeout:{task_id}")
        await asyncio.sleep(float(poll_ms) / 1000.0)

    if state != "completed":
        message_obj = _extract_a2a_status_message(snapshot)
        status_text = _extract_text_from_parts(message_obj.get("parts"))
        metadata_obj = message_obj.get("metadata")
        status_meta = metadata_obj if isinstance(metadata_obj, dict) else {}
        if not status_text:
            status_text = str(status_meta.get("error") or status_meta.get("error_type") or state)
        raise RuntimeError(f"a2a_task_{state}:{status_text}")

    final_answer_obj = _extract_a2a_artifact(snapshot, "final_answer")
    final_answer = str(final_answer_obj or "").strip()
    message_obj = _extract_a2a_status_message(snapshot)
    if not final_answer:
        final_answer = _extract_text_from_parts(message_obj.get("parts"))

    message_meta_obj = message_obj.get("metadata")
    message_meta = message_meta_obj if isinstance(message_meta_obj, dict) else {}
    output_type = str(message_meta.get("output_type") or "").strip()
    confidence = _safe_float(message_meta.get("confidence"), default=0.0)

    tool_observations_obj = _extract_a2a_artifact(snapshot, "tool_observations")
    tool_observations = tool_observations_obj if isinstance(tool_observations_obj, list) else []
    telemetry_obj = _extract_a2a_artifact(snapshot, "telemetry")
    telemetry = telemetry_obj if isinstance(telemetry_obj, dict) else {}
    progress_events_obj = _extract_a2a_artifact(snapshot, "progress_events")
    progress_events = progress_events_obj if isinstance(progress_events_obj, list) else []
    processing_steps = _build_processing_steps_from_progress_events(progress_events)

    return {
        "output": final_answer,
        "output_type": output_type,
        "confidence": confidence,
        "tool_observations": tool_observations,
        "reflection_events": progress_events,
        "processing_steps": processing_steps,
        "telemetry": telemetry,
        "a2a_task_id": task_id,
    }


def _map_a2a_state_to_done_status(state: str) -> str:
    """将 A2A 状态映射为前端完成状态。

    Args:
        state: 任务状态字符串。

    Returns:
        str: 映射后的终态标识。
    """
    safe_state = str(state or "").strip().lower()
    if safe_state == "completed":
        return "done"
    if safe_state == "canceled":
        return "cancelled"
    if safe_state == "failed":
        return "failed"
    return safe_state or "failed"


def _extract_progress_detail(event: Dict[str, Any]) -> str:
    """提取进度事件展示文案。

    Args:
        event: 事件字典。

    Returns:
        str: 可展示的进度文本。
    """
    if not isinstance(event, dict):
        return ""
    return str(event.get("detail") or event.get("event") or "").strip()


def _extract_task_final_text(task: Dict[str, Any], state: str) -> str:
    """提取任务终态展示文本。

    Args:
        task: 任务对象字典。
        state: 任务状态字符串。

    Returns:
        str: 任务终态文案。
    """
    if not isinstance(task, dict):
        return str(state or "")
    if state == "completed":
        final_answer_obj = _extract_a2a_artifact(task, "final_answer")
        final_answer = str(final_answer_obj or "").strip()
        if final_answer:
            return final_answer

    status_message = _extract_a2a_status_message(task)
    status_text = _extract_text_from_parts(status_message.get("parts"))
    if status_text and status_text not in {"task_submitted", "task_running", "task_completed"}:
        return status_text

    metadata_obj = status_message.get("metadata")
    metadata = metadata_obj if isinstance(metadata_obj, dict) else {}
    fallback = str(metadata.get("error") or metadata.get("error_type") or state or "").strip()
    return fallback or "任务结束"


async def _build_a2a_runtime_snapshot(limit_per_service: int = 200) -> Dict[str, Any]:
    """汇总 A2A 运行时任务快照。

    Args:
        limit_per_service: 每个服务统计的任务上限。

    Returns:
        Dict[str, Any]: 各 A2A 服务任务统计。
    """
    services = {
        "orchestrator": get_a2a(),
        "graph": get_a2a_graph(),
        "retrieval": get_a2a_retrieval(),
        "evaluator": get_a2a_evaluator(),
    }
    result: Dict[str, Any] = {}
    total = 0
    for name, service in services.items():
        rows = await service.alist_tasks(limit=limit_per_service)
        counter: Dict[str, int] = {}
        for row in rows:
            status_obj = row.get("status")
            status = status_obj if isinstance(status_obj, dict) else {}
            state = str(status.get("state") or "unknown")
            counter[state] = int(counter.get(state, 0)) + 1
        result[name] = {
            "recent_tasks": len(rows),
            "states": counter,
        }
        total += len(rows)
    return {
        "recent_total_tasks": total,
        "services": result,
    }


async def _generate_a2a_stream_events(
    input_text: str,
    user_id: str,
    session_id: str = "",
    request_id: str = "",
    strategy: str = "iterative",
    max_turns: Optional[int] = None,
    original_input_text: str = "",
    resolved_query: Optional[Dict[str, Any]] = None,
    memory_context: Optional[Dict[str, Any]] = None,
) -> Any:
    """生成 Agent 流式输出事件。

    Args:
        input_text: 用户输入文本。
        user_id: 用户 ID。
        session_id: 会话 ID。
        request_id: 请求 ID。
        strategy: 策略名称。
        max_turns: 最大迭代轮数。

    Returns:
        Any: 可异步迭代的事件流。
    """
    req_id = str(request_id or uuid.uuid4())
    service = get_a2a()

    safe_original_input_text = str(original_input_text or input_text or "")
    safe_resolved_query = resolved_query if isinstance(resolved_query, dict) else {}
    safe_memory_context = memory_context if isinstance(memory_context, dict) else {}

    metadata: Dict[str, Any] = {
        "session_id": str(session_id or ""),
        "user_id": str(user_id or ""),
        "strategy": str(strategy or "iterative"),
        "original_input_text": safe_original_input_text,
        "resolved_query": safe_resolved_query,
        "memory_context": safe_memory_context,
        "source": "stream_api",
    }
    if isinstance(max_turns, int) and max_turns > 0:
        metadata["max_turns"] = max_turns

    message = {
        "role": "user",
        "parts": [{"kind": "text", "text": str(input_text or "")}],
        "metadata": metadata,
        "contextId": str(session_id or ""),
        "userId": str(user_id or ""),
    }
    task = await service.asend_message(message, req_id)
    task_id = str((task or {}).get("id") or "").strip()
    if not task_id:
        raise RuntimeError("a2a_task_id_missing")

    yield {"type": "job", "job_id": task_id, "request_id": req_id, "task_id": task_id}

    safe_poll_ms = max(
        100,
        min(int(_safe_float(os.getenv("AGENT_RUN_A2A_POLL_MS", "250"), default=250.0)), 3000),
    )
    traces: List[Dict[str, Any]] = []
    cursor = 0

    while True:
        snapshot = await service.aget_task_updates(task_id, cursor)
        if snapshot is None:
            yield {
                "type": "done",
                "msg": "任务不存在或已被清理",
                "status": "failed",
                "traces": traces,
                "task_id": task_id,
            }
            break

        events_obj = snapshot.get("events")
        events = events_obj if isinstance(events_obj, list) else []
        cursor = int(snapshot.get("next_cursor") or cursor)

        for event in events:
            detail = _extract_progress_detail(event if isinstance(event, dict) else {})
            if not detail:
                continue
            traces.append({"step": len(traces) + 1, "action": detail})
            yield {"type": "msg", "msg": detail, "task_id": task_id}

        state = str(snapshot.get("state") or "").strip().lower()
        if bool(snapshot.get("terminal")):
            task_obj = snapshot.get("task")
            task_payload = task_obj if isinstance(task_obj, dict) else {}
            final_text = _extract_task_final_text(task_payload, state)
            processing_steps_obj = _extract_a2a_artifact(task_payload, "processing_steps")
            processing_steps = processing_steps_obj if isinstance(processing_steps_obj, list) else []
            tool_observations_obj = _extract_a2a_artifact(task_payload, "tool_observations")
            tool_observations = tool_observations_obj if isinstance(tool_observations_obj, list) else []
            evidence_summary_obj = _extract_a2a_artifact(task_payload, "evidence_summary")
            evidence_summary = evidence_summary_obj if isinstance(evidence_summary_obj, dict) else {}
            if state == "completed" and final_text:
                yield {"type": "reply", "msg": final_text, "task_id": task_id}
            yield {
                "type": "done",
                "msg": final_text,
                "status": _map_a2a_state_to_done_status(state),
                "traces": traces,
                "task_id": task_id,
                "processing_steps": processing_steps,
                "tool_observations": tool_observations,
                "evidence_summary": evidence_summary,
            }
            break

        await asyncio.sleep(float(safe_poll_ms) / 1000.0)


def get_graph_subagent() -> GraphSubAgent:
    """获取图谱子 Agent 单例。"""
    global _graph_subagent
    if _graph_subagent is None:
        _graph_subagent = GraphSubAgent(get_mcp())
    return _graph_subagent


def get_rag_subagent() -> RagSubAgent:
    """获取 RAG 子 Agent 单例。"""
    global _rag_subagent
    if _rag_subagent is None:
        _rag_subagent = RagSubAgent(get_mcp())
    return _rag_subagent


def _build_subagent_requests(
    input_text: str,
    metadata: Dict[str, Any],
    default_tool_name: str,
) -> List[Dict[str, Any]]:
    """从 A2A metadata 组装子 Agent 请求列表。"""
    requests_obj = metadata.get("requests")
    if isinstance(requests_obj, list):
        rows: List[Dict[str, Any]] = []
        for idx, item in enumerate(requests_obj, start=1):
            if not isinstance(item, dict):
                continue
            args_obj = item.get("arguments")
            arguments = args_obj if isinstance(args_obj, dict) else {}
            subquery = str(item.get("subquery") or input_text or "").strip()
            tool_name = str(item.get("tool_name") or default_tool_name).strip() or default_tool_name
            rows.append(
                {
                    "subquery_id": int(item.get("subquery_id") or idx),
                    "subquery": subquery,
                    "tool_name": tool_name,
                    "arguments": arguments,
                    "reason": str(item.get("reason") or "").strip(),
                }
            )
        if rows:
            return rows

    args_obj = metadata.get("arguments")
    arguments = args_obj if isinstance(args_obj, dict) else {}
    subquery = str(metadata.get("subquery") or input_text or "").strip()
    tool_name = str(metadata.get("tool_name") or default_tool_name).strip() or default_tool_name
    return [
        {
            "subquery_id": 1,
            "subquery": subquery,
            "tool_name": tool_name,
            "arguments": arguments,
            "reason": str(metadata.get("reason") or "").strip(),
        }
    ]


async def _run_graph_a2a_query(
    input_text: str,
    context_id: str = "",
    user_id: str = "",
    metadata: Optional[Dict[str, Any]] = None,
    cancel_check: Optional[Callable[[], bool]] = None,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
) -> Dict[str, Any]:
    """Graph 子代理 A2A 执行入口。"""
    _ = cancel_check
    meta = metadata if isinstance(metadata, dict) else {}
    session_id = str(context_id or meta.get("session_id") or "").strip()
    safe_user_id = str(user_id or meta.get("user_id") or "").strip()
    reason = str(meta.get("reason") or "a2a_graph_rpc").strip()
    requests = _build_subagent_requests(
        input_text=str(input_text or ""),
        metadata=meta,
        default_tool_name="kg_entity_relation_search",
    )
    if progress_callback is not None:
        progress_callback(
            {
                "event": "graph_batch_start",
                "detail": f"graph requests={len(requests)}",
                "tool_name": "graph_subagent",
            }
        )

    result, result_meta = await get_graph_subagent().invoke_batch(
        requests=requests,
        user_query=str(input_text or ""),
        session_id=session_id,
        user_id=safe_user_id,
        reason=reason,
    )
    payload_obj = result.get("result") if isinstance(result, dict) else {}
    payload = payload_obj if isinstance(payload_obj, dict) else {}
    summary_obj = payload.get("summary")
    summary = summary_obj if isinstance(summary_obj, dict) else {}
    items_obj = payload.get("items")
    items = items_obj if isinstance(items_obj, list) else []
    signal_count = int(summary.get("signal_count") or 0)
    total_requests = int(summary.get("total_requests") or len(requests))
    note = str(summary.get("note") or "").strip()
    if not note:
        note = f"Graph 子代理完成 {total_requests} 个检索任务"
    confidence = 0.75 if signal_count > 0 else 0.45
    if progress_callback is not None:
        progress_callback(
            {
                "event": "graph_batch_done",
                "detail": note,
                "tool_name": "graph_subagent",
            }
        )
    return {
        "output": note,
        "output_type": "graph_subagent_batch",
        "confidence": confidence,
        "tool_observations": items,
        "graph_result": result,
        "graph_meta": result_meta,
        "telemetry": {
            "tool_calls": total_requests,
            "graph_batch_size": total_requests,
            "graph_signal_count": signal_count,
            "graph_ok_count": int(summary.get("ok_count") or 0),
        },
        "evidence_summary": {
            "evidence_count": signal_count,
            "conflict_count": 0,
            "hop_coverage": 1.0 if signal_count > 0 else 0.0,
            "uncovered_points": [],
            "high_risk": False,
            "high_risk_tokens": [],
            "disambiguation_ambiguous_count": 0,
            "community_count": 0,
            "top_community_score": 0.0,
            "force_supplement_required": False,
        },
    }


async def _run_retrieval_a2a_query(
    input_text: str,
    context_id: str = "",
    user_id: str = "",
    metadata: Optional[Dict[str, Any]] = None,
    cancel_check: Optional[Callable[[], bool]] = None,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
) -> Dict[str, Any]:
    """Retrieval 子代理 A2A 执行入口。"""
    _ = cancel_check
    meta = metadata if isinstance(metadata, dict) else {}
    session_id = str(context_id or meta.get("session_id") or "").strip()
    safe_user_id = str(user_id or meta.get("user_id") or "").strip()
    reason = str(meta.get("reason") or "a2a_retrieval_rpc").strip()
    requests = _build_subagent_requests(
        input_text=str(input_text or ""),
        metadata=meta,
        default_tool_name="rag_search",
    )
    if progress_callback is not None:
        progress_callback(
            {
                "event": "retrieval_batch_start",
                "detail": f"retrieval requests={len(requests)}",
                "tool_name": "rag_subagent",
            }
        )

    result, result_meta = await get_rag_subagent().invoke_batch(
        requests=requests,
        user_query=str(input_text or ""),
        session_id=session_id,
        user_id=safe_user_id,
        reason=reason,
    )
    payload_obj = result.get("result") if isinstance(result, dict) else {}
    payload = payload_obj if isinstance(payload_obj, dict) else {}
    summary_obj = payload.get("summary")
    summary = summary_obj if isinstance(summary_obj, dict) else {}
    items_obj = payload.get("items")
    items = items_obj if isinstance(items_obj, list) else []
    signal_count = int(summary.get("signal_count") or 0)
    total_requests = int(summary.get("total_requests") or len(requests))
    note = str(summary.get("note") or "").strip()
    if not note:
        note = f"Retrieval 子代理完成 {total_requests} 个检索任务"
    confidence = 0.72 if signal_count > 0 else 0.42
    if progress_callback is not None:
        progress_callback(
            {
                "event": "retrieval_batch_done",
                "detail": note,
                "tool_name": "rag_subagent",
            }
        )
    return {
        "output": note,
        "output_type": "retrieval_subagent_batch",
        "confidence": confidence,
        "tool_observations": items,
        "retrieval_result": result,
        "retrieval_meta": result_meta,
        "telemetry": {
            "tool_calls": total_requests,
            "retrieval_batch_size": total_requests,
            "retrieval_signal_count": signal_count,
            "retrieval_ok_count": int(summary.get("ok_count") or 0),
        },
        "evidence_summary": {
            "evidence_count": signal_count,
            "conflict_count": 0,
            "hop_coverage": 1.0 if signal_count > 0 else 0.0,
            "uncovered_points": [],
            "high_risk": False,
            "high_risk_tokens": [],
            "disambiguation_ambiguous_count": 0,
            "community_count": 0,
            "top_community_score": 0.0,
            "force_supplement_required": False,
        },
    }


async def _run_evaluator_a2a_query(
    input_text: str,
    context_id: str = "",
    user_id: str = "",
    metadata: Optional[Dict[str, Any]] = None,
    cancel_check: Optional[Callable[[], bool]] = None,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
) -> Dict[str, Any]:
    """Evaluator 子代理 A2A 执行入口。"""
    _ = input_text
    _ = context_id
    _ = user_id
    _ = cancel_check
    meta = metadata if isinstance(metadata, dict) else {}
    traces_obj = meta.get("traces") or meta.get("tool_observations")
    traces = traces_obj if isinstance(traces_obj, list) else []
    classification_obj = meta.get("classification")
    classification = classification_obj if isinstance(classification_obj, dict) else {}
    backtrack_count = int(meta.get("backtrack_count") or 0)
    max_backtrack = int(meta.get("max_backtrack") or 2)
    is_degraded = bool(meta.get("is_degraded"))
    if progress_callback is not None:
        progress_callback(
            {
                "event": "evaluator_start",
                "detail": "evidence evaluator running",
                "tool_name": "evaluator_agent",
            }
        )
    gate_result = evaluate_and_gate(
        traces=traces,
        classification=classification,
        backtrack_count=backtrack_count,
        max_backtrack=max_backtrack,
        is_degraded=is_degraded,
    )
    gate_passed = bool(gate_result.get("gate_passed"))
    note = "证据门控通过，可终答" if gate_passed else "证据仍不足或需补检索"
    if progress_callback is not None:
        progress_callback(
            {
                "event": "evaluator_done",
                "detail": note,
                "tool_name": "evaluator_agent",
            }
        )
    return {
        "output": note,
        "output_type": "evaluator_gate",
        "confidence": float(gate_result.get("confidence") or 0.0),
        "evaluator_result": gate_result,
        "telemetry": {
            "gate_passed": gate_passed,
            "sufficient": bool(gate_result.get("sufficient")),
            "should_backtrack": bool(gate_result.get("should_backtrack")),
            "evidence_count": int(gate_result.get("evidence_count") or 0),
            "conflict_count": int(gate_result.get("conflict_count") or 0),
        },
        "evidence_summary": {
            "evidence_count": int(gate_result.get("evidence_count") or 0),
            "conflict_count": int(gate_result.get("conflict_count") or 0),
            "hop_coverage": float(gate_result.get("hop_coverage") or 0.0),
            "uncovered_points": gate_result.get("uncovered_points") or [],
            "high_risk": bool(gate_result.get("high_risk")),
            "high_risk_tokens": gate_result.get("high_risk_tokens") or [],
            "disambiguation_ambiguous_count": int(
                gate_result.get("disambiguation_ambiguous_count") or 0
            ),
            "community_count": int(gate_result.get("community_count") or 0),
            "top_community_score": float(gate_result.get("top_community_score") or 0.0),
            "force_supplement_required": bool(gate_result.get("force_supplement_required")),
        },
    }


def get_a2a_graph() -> A2AService:
    """获取图谱子代理 A2A 服务。"""
    global _a2a_graph_service
    if _a2a_graph_service is None:
        _a2a_graph_service = A2AService(
            run_query=_run_graph_a2a_query,
            agent_name="Hypertension Graph Agent",
            agent_description="图谱检索与关系证据代理（A2A 适配）",
            protocol_version="0.2.5",
            agent_url="/a2a/graph/rpc",
            events_base_path="/a2a/graph/tasks",
        )
    return _a2a_graph_service


def get_a2a_retrieval() -> A2AService:
    """获取检索子代理 A2A 服务。"""
    global _a2a_retrieval_service
    if _a2a_retrieval_service is None:
        _a2a_retrieval_service = A2AService(
            run_query=_run_retrieval_a2a_query,
            agent_name="Hypertension Retrieval Agent",
            agent_description="语义检索与文本证据代理（A2A 适配）",
            protocol_version="0.2.5",
            agent_url="/a2a/retrieval/rpc",
            events_base_path="/a2a/retrieval/tasks",
        )
    return _a2a_retrieval_service


def get_a2a_evaluator() -> A2AService:
    """获取评估子代理 A2A 服务。"""
    global _a2a_evaluator_service
    if _a2a_evaluator_service is None:
        _a2a_evaluator_service = A2AService(
            run_query=_run_evaluator_a2a_query,
            agent_name="Hypertension Evaluator Agent",
            agent_description="证据门控与风险评估代理（A2A 适配）",
            protocol_version="0.2.5",
            agent_url="/a2a/evaluator/rpc",
            events_base_path="/a2a/evaluator/tasks",
        )
    return _a2a_evaluator_service


async def _measure_event_loop_lag(samples: int = 20, interval_ms: int = 20) -> Dict[str, Any]:
    """
    采样估算当前事件循环延迟。

    Args:
        samples: 采样次数
        interval_ms: 采样间隔毫秒

    Returns:
        Dict[str, Any]: 延迟统计信息（毫秒）
    """
    safe_samples = max(5, min(int(samples), 200))
    safe_interval_ms = max(1, min(int(interval_ms), 200))
    interval_s = float(safe_interval_ms) / 1000.0
    lags_ms: list[float] = []

    for _ in range(safe_samples):
        started = time.perf_counter()
        await asyncio.sleep(interval_s)
        elapsed = time.perf_counter() - started
        lag_ms = max(0.0, (elapsed - interval_s) * 1000.0)
        lags_ms.append(lag_ms)

    sorted_lags = sorted(lags_ms)

    def _pct(p: float) -> float:
        """计算事件循环延迟分位值。

        Args:
            p: 分位比例，范围通常为 0 到 1。

        Returns:
            float: 对应分位值（毫秒）。
        """
        if not sorted_lags:
            return 0.0
        if len(sorted_lags) == 1:
            return sorted_lags[0]
        rank = max(0.0, min(1.0, p)) * (len(sorted_lags) - 1)
        low = int(rank)
        high = min(low + 1, len(sorted_lags) - 1)
        weight = rank - low
        return sorted_lags[low] * (1.0 - weight) + sorted_lags[high] * weight

    return {
        "samples": safe_samples,
        "interval_ms": safe_interval_ms,
        "avg_ms": round(sum(sorted_lags) / max(1, len(sorted_lags)), 2),
        "p50_ms": round(_pct(0.50), 2),
        "p95_ms": round(_pct(0.95), 2),
        "p99_ms": round(_pct(0.99), 2),
        "max_ms": round(sorted_lags[-1] if sorted_lags else 0.0, 2),
    }


async def _await_call(func: Callable[..., Any], *args: Any, **kwargs: Any) -> Any:
    """统一等待可等待对象；仅保留异步主链路。"""
    result = func(*args, **kwargs)
    if inspect.isawaitable(result):
        return await result
    return result


async def _build_readiness_report() -> Dict[str, Any]:
    """构建服务就绪性报告。

    Returns:
        Dict[str, Any]: 当前服务 readiness 报告。
    """
    db_status = await check_databases()

    neo4j_ready = bool((db_status.get("neo4j") or {}).get("is_running"))
    milvus_ready = bool((db_status.get("milvus") or {}).get("is_running"))
    mysql_ready = bool(_mysql_initialized)

    dependencies = [
        {
            "name": "mysql",
            "critical": True,
            "ready": mysql_ready,
            "detail": "数据库初始化成功" if mysql_ready else "数据库初始化失败或尚未完成",
        },
        {
            "name": "neo4j",
            "critical": False,
            "ready": neo4j_ready,
            "detail": f"host={(db_status.get('neo4j') or {}).get('host', '')}",
        },
        {
            "name": "milvus",
            "critical": False,
            "ready": milvus_ready,
            "detail": f"host={(db_status.get('milvus') or {}).get('host', '')}",
        },
    ]

    mcp_detail: Dict[str, Any] = {
        "core_ready": False,
        "missing_required": [],
        "runtime_unavailable_required": [],
        "missing_optional": [],
        "runtime_unavailable_optional": [],
        "runtime_errors": {},
        "tool_count": 0,
        "tool_names": [],
        "optional_issues": [],
    }
    mcp_error = ""

    try:
        mcp_detail = await get_mcp_runtime_status(get_mcp())
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        mcp_error = str(exc)

    mcp_core_ready = bool(mcp_detail.get("core_ready")) and not mcp_error
    dependencies.append(
        {
            "name": "mcp_core_tools",
            "critical": False,
            "ready": mcp_core_ready,
            "detail": mcp_error or json.dumps(
                {
                    "missing_required": mcp_detail.get("missing_required", []),
                    "runtime_unavailable_required": mcp_detail.get("runtime_unavailable_required", []),
                },
                ensure_ascii=False,
            ),
        }
    )

    optional_issues = list(mcp_detail.get("optional_issues") or [])
    if not neo4j_ready:
        optional_issues.append("neo4j_unavailable")
    if not milvus_ready:
        optional_issues.append("milvus_unavailable")
    if not mcp_core_ready:
        optional_issues.append("mcp_core_degraded")
    if mcp_error:
        optional_issues.append("mcp_server_unavailable")

    critical_not_ready = [item["name"] for item in dependencies if item["critical"] and not item["ready"]]

    level = "ready"
    if critical_not_ready:
        level = "unavailable"
    elif optional_issues:
        level = "degraded"

    return {
        "level": level,
        "ready": level != "unavailable",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "critical_not_ready": critical_not_ready,
        "optional_issues": optional_issues,
        "dependencies": dependencies,
        "db": db_status,
        "mcp": mcp_detail,
        "runtime_flags": get_runtime_flag_snapshot(),
    }


async def _ensure_core_ready(endpoint: str) -> Optional[JSONResponse]:
    """在接口执行前校验核心依赖就绪状态。

    Args:
        endpoint: 接口路径标识。

    Returns:
        Optional[JSONResponse]: 不可用时返回 503 响应，否则返回 None。
    """
    report = await _await_call(_build_readiness_report)
    if report["level"] != "unavailable":
        return None

    logger.warning(
        "[readiness_block] %s",
        json.dumps(
            {
                "endpoint": endpoint,
                "critical_not_ready": report.get("critical_not_ready", []),
            },
            ensure_ascii=False,
        ),
    )
    return JSONResponse(
        status_code=503,
        content={
            "ok": False,
            "error": "service_unavailable",
            "message": "核心依赖未就绪，请稍后重试",
            "readiness": report,
        },
    )


async def _wrap_stream_body_iterator(
    body_iterator: Any,
    on_close: Callable[[str], Any],
) -> AsyncIterator[bytes]:
    """包装 StreamingResponse body_iterator，确保结束时回调统一收口。"""
    stream_error_type = ""
    try:
        if hasattr(body_iterator, "__aiter__"):
            async for chunk in body_iterator:
                yield chunk
        else:
            for chunk in body_iterator:
                yield chunk
    except BaseException as exc:
        stream_error_type = type(exc).__name__
        raise
    finally:
        callback_ret = on_close(stream_error_type)
        if inspect.isawaitable(callback_ret):
            await callback_ret


def _resolve_api_session_user(request: Request) -> Dict[str, str]:
    """解析接口日志身份信息，优先使用 request.state。"""
    state_session_id = str(getattr(request.state, "session_id", "") or "")
    state_user_id = str(getattr(request.state, "user_id", "") or "")
    query_session_id = str(request.query_params.get("session_id") or "")
    query_user_id = str(request.query_params.get("user_id") or "")
    return {
        "session_id": state_session_id or query_session_id,
        "user_id": state_user_id or query_user_id,
    }


@app.middleware("http")
async def request_timing_middleware(request: Request, call_next):
    """记录请求耗时与链路上下文。

    Args:
        request: FastAPI 请求对象。
        call_next: 下一个中间件或路由处理函数。

    Returns:
        Any: 下游中间件响应对象。
    """
    request_id = create_request_id()
    request.state.request_id = request_id
    request.state.trace_id = request_id
    request_token = set_request_context(request_id)
    bearer_token = _extract_bearer_token(request.headers.get("Authorization", ""))
    token_handle = set_current_admin_token(bearer_token)
    started = time.perf_counter()
    status_code = 500
    error_type = ""
    ttfb_ms = 0.0
    logged = False
    is_streaming_response = False
    client_ip = request.client.host if request.client else ""
    request_timer_ctx = log_duration("http_request", method=request.method, path=request.url.path)
    request_timer_ctx.__enter__()
    request_timer_closed = False

    def _finish_request_timer(exc: Optional[BaseException] = None) -> None:
        nonlocal request_timer_closed
        if request_timer_closed:
            return
        request_timer_closed = True
        if exc is None:
            request_timer_ctx.__exit__(None, None, None)
            return
        request_timer_ctx.__exit__(type(exc), exc, exc.__traceback__)

    async def _log_stream_finished(stream_error_type: str) -> None:
        nonlocal logged
        if logged:
            _finish_request_timer()
            return
        stream_total_ms = (time.perf_counter() - started) * 1000
        identity = _resolve_api_session_user(request)
        log_api_call(
            request_id=request_id,
            method=request.method,
            path=request.url.path,
            status_code=status_code,
            is_streaming=True,
            ttfb_ms=ttfb_ms,
            stream_total_ms=stream_total_ms,
            client_ip=client_ip,
            session_id=identity["session_id"],
            user_id=identity["user_id"],
            error_type=stream_error_type or error_type,
        )
        logged = True
        _finish_request_timer()

    try:
        set_trace_metadata(
            {
                "trace_id": request_id,
                "request_id": request_id,
                "method": request.method,
                "path": request.url.path,
            },
            tags=["api"],
        )
        response = await call_next(request)
        response.headers["X-Trace-Id"] = request_id
        status_code = int(response.status_code)

        ttfb_ms = (time.perf_counter() - started) * 1000
        if isinstance(response, StreamingResponse):
            response.body_iterator = _wrap_stream_body_iterator(
                response.body_iterator,
                _log_stream_finished,
            )
            is_streaming_response = True
            return response

        identity = _resolve_api_session_user(request)
        log_api_call(
            request_id=request_id,
            method=request.method,
            path=request.url.path,
            status_code=status_code,
            is_streaming=False,
            ttfb_ms=ttfb_ms,
            stream_total_ms=ttfb_ms,
            client_ip=client_ip,
            session_id=identity["session_id"],
            user_id=identity["user_id"],
            error_type=error_type,
        )
        logged = True
        _finish_request_timer()
        return response
    except Exception as exc:
        error_type = type(exc).__name__
        if not logged:
            stream_total_ms = (time.perf_counter() - started) * 1000
            fallback_ttfb = ttfb_ms if ttfb_ms > 0.0 else stream_total_ms
            identity = _resolve_api_session_user(request)
            log_api_call(
                request_id=request_id,
                method=request.method,
                path=request.url.path,
                status_code=status_code,
                is_streaming=False,
                ttfb_ms=fallback_ttfb,
                stream_total_ms=stream_total_ms,
                client_ip=client_ip,
                session_id=identity["session_id"],
                user_id=identity["user_id"],
                error_type=error_type,
            )
            logged = True
        _finish_request_timer(exc)
        raise
    finally:
        if not is_streaming_response:
            _finish_request_timer()
        reset_current_admin_token(token_handle)
        reset_request_context(request_token)


# 数据库状态检查端点
@app.get("/db/status")
async def db_status():
    """获取数据库连接状态。

    Returns:
        Dict[str, Any]: MySQL、Milvus、Neo4j 的连接状态汇总。
    """
    status = await _await_call(check_databases)
    return status

# 从环境变量读取允许的跨域来源，默认开发环境前端地址
_allowed_origins = os.getenv("ALLOWED_ORIGINS", "http://localhost:5173,http://localhost:3000")
ALLOWED_ORIGINS = [origin.strip() for origin in _allowed_origins.split(",") if origin.strip()]

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
    allow_headers=["Authorization", "Content-Type", "X-Requested-With"],
)

UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

app.mount("/uploads", StaticFiles(directory=str(UPLOAD_DIR)), name="uploads")


@app.post("/process")
async def process(data: dict, request: Request):
    """处理流式问答请求入口。

    Args:
        data: 请求体数据字典。
        request: FastAPI 请求对象。

    Returns:
        StreamingResponse: 流式响应对象。
    """
    readiness_error = await _ensure_core_ready("/process")
    if readiness_error is not None:
        return readiness_error

    input_text = data.get("input", "")
    user_id = data.get("user_id") or ""
    session_id = data.get("session_id") or ""
    if not user_id:
        user_id = str(uuid.uuid4())
    return StreamingResponse(
        generate_data(
            input_text,
            user_id,
            session_id,
            request_id=str(getattr(request.state, "request_id", "") or ""),
        ),
        media_type="text/plain",
    )


@app.post("/agent/run")
async def agent_run(data: dict, request: Request):
    """运行 Agent（仅保留 A2A 架构）。"""
    trace_id = _resolve_trace_id_from_request(request)
    readiness_error = await _ensure_core_ready("/agent/run")
    if readiness_error is not None:
        return readiness_error

    input_text = (data.get("input") or "").strip()
    session_id = data.get("session_id", "")
    user_id = data.get("user_id", "")
    strategy = str(data.get("strategy") or "iterative").strip().lower()
    try:
        a2a_mode = _resolve_agent_run_mode(data)
    except ValueError as exc:
        return {"ok": False, "error": f"invalid_request:{str(exc)}", "trace_id": trace_id}
    raw_max_turns = data.get("max_turns")
    max_turns: Optional[int] = raw_max_turns if isinstance(raw_max_turns, int) and raw_max_turns > 0 else None

    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "session_id": str(session_id or ""),
            "user_id": str(user_id or ""),
            "endpoint": "/agent/run",
            "strategy": strategy,
            "a2a_mode": a2a_mode,
        },
        tags=["api", "agent_run"],
    )
    limited = _enforce_heavy_api_rate_limit(
        endpoint="/agent/run",
        user_id=str(user_id or ""),
        session_id=str(session_id or ""),
        request=request,
    )
    if limited is not None:
        return limited
    
    if not input_text:
        return {"ok": False, "error": "input_required", "trace_id": trace_id}
    
    if session_id:
        ensure_session_meta(str(session_id), str(user_id or ""))
    session_tokens = set_session_context(str(session_id or ""), str(user_id or ""))

    try:
        result = await _run_agent_via_a2a(
            input_text=input_text,
            session_id=session_id,
            user_id=user_id,
            strategy=strategy,
            max_turns=max_turns,
            trace_id=trace_id,
        )

        traces = _build_traces_from_processing_steps(result.get("processing_steps", []))
        response = {
            "ok": True,
            "answer": result.get("output", ""),
            "output_type": result.get("output_type", ""),
            "strategy": strategy,
            "traces": traces,
            "tool_observations": result.get("tool_observations", []),
            "reflection_events": result.get("reflection_events", []),
            "telemetry": result.get("telemetry", {}),
            "confidence": result.get("confidence", 0.0),
            "trace_id": trace_id,
            "run_path": "a2a",
            "a2a_mode": a2a_mode,
        }
        if str(result.get("a2a_task_id") or "").strip():
            response["a2a_task_id"] = str(result.get("a2a_task_id"))
        return response
    except ValueError as exc:
        _log_api_error("/agent/run", "validation", exc, strategy=strategy)
        return {"ok": False, "error": f"invalid_request:{str(exc)}", "trace_id": trace_id}
    except RuntimeError as exc:
        _log_api_error("/agent/run", "runtime", exc, strategy=strategy)
        return {"ok": False, "error": f"agent_runtime_error:{str(exc)}", "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/agent/run", "unknown", exc, strategy=strategy)
        return {"ok": False, "error": f"agent_unknown_error:{str(exc)}", "trace_id": trace_id}
    finally:
        reset_session_context(session_tokens)


@app.get("/health/readiness")
async def health_readiness():
    """返回服务 readiness 分级状态。"""
    report = await _await_call(_build_readiness_report)
    status_code = 200 if report["level"] != "unavailable" else 503
    return JSONResponse(status_code=status_code, content={"ok": report["ready"], "data": report})


@app.get("/agent/tools")
async def agent_tools():
    """返回 Agent 策略与可用工具信息"""
    mcp = get_mcp()
    mcp_tools = await _await_call(mcp.list_tools)
    return {
        "default_strategy": "iterative",
        "supported_strategies": ["iterative", "graph"],
        "iterative_actions": [
            {"name": "call_tool", "description": "调用 MCP 工具继续获取证据"},
            {"name": "final_answer", "description": "证据充分后直接输出最终答案"},
        ],
        "mcp_tools": mcp_tools,
    }


@app.get("/mcp/tools")
async def mcp_tools():
    """返回 MCP 服务中注册的工具列表"""
    try:
        mcp = get_mcp()
        tools = await _await_call(mcp.list_tools)
        return {
            "ok": True,
            "tools": tools
        }
    except RuntimeError as exc:
        _log_api_error("/mcp/tools", "runtime", exc)
        return {"ok": False, "error": f"mcp_runtime_error:{str(exc)}"}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/mcp/tools", "unknown", exc)
        return {"ok": False, "error": f"mcp_unknown_error:{str(exc)}"}


@app.post("/mcp/invoke")
async def mcp_invoke(data: dict, request: Request):
    """调用 MCP 工具
    
    Args:
        data: {
            "tool_name": "工具名称",
            "payload": {工具参数}
        }
    """
    trace_id = _resolve_trace_id_from_request(request)
    tool_name = data.get("tool_name")
    server_name = data.get("server_name")
    payload_obj = data.get("payload", {})
    payload = payload_obj if isinstance(payload_obj, dict) else {}
    user_id = str(data.get("user_id") or payload.get("user_id") or "").strip()
    session_id = str(data.get("session_id") or payload.get("session_id") or "").strip()

    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": "/mcp/invoke",
            "tool_name": str(tool_name or ""),
            "server_name": str(server_name or ""),
            "user_id": user_id,
            "session_id": session_id,
        },
        tags=["api", "mcp_invoke"],
    )
    limited = _enforce_heavy_api_rate_limit(
        endpoint="/mcp/invoke",
        user_id=user_id,
        session_id=session_id,
        request=request,
    )
    if limited is not None:
        return limited
    
    if not tool_name:
        return {"ok": False, "error": "tool_name_required", "trace_id": trace_id}
    
    try:
        mcp = get_mcp()
        resolved_tool_name = (
            f"{str(server_name).strip()}::{str(tool_name).strip()}"
            if str(server_name or "").strip()
            else str(tool_name).strip()
        )
        single_call = [
            {
                "tool_name": resolved_tool_name,
                "payload": payload,
            }
        ]
        rows = await mcp.ainvoke_batch(single_call)
        row_obj = rows[0] if isinstance(rows, list) and rows else {}
        result = (
            row_obj
            if isinstance(row_obj, dict)
            else {
                "ok": False,
                "error": "mcp_batch_empty_result",
                "error_type": "runtime_error",
                "meta": {},
            }
        )
        if isinstance(result, dict):
            meta_obj = result.get("meta")
            meta = meta_obj if isinstance(meta_obj, dict) else {}
            meta["trace_id"] = trace_id
            result["meta"] = meta
            if "trace_id" not in result:
                result["trace_id"] = trace_id
        return result
    except ValueError as exc:
        _log_api_error("/mcp/invoke", "payload_validation", exc, tool_name=tool_name, server_name=server_name)
        return {"ok": False, "error": f"invalid_payload:{str(exc)}", "trace_id": trace_id}
    except RuntimeError as exc:
        _log_api_error("/mcp/invoke", "runtime", exc, tool_name=tool_name, server_name=server_name)
        return {"ok": False, "error": f"mcp_runtime_error:{str(exc)}", "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/mcp/invoke", "unknown", exc, tool_name=tool_name, server_name=server_name)
        return {"ok": False, "error": f"mcp_unknown_error:{str(exc)}", "trace_id": trace_id}


@app.get("/a2a/agent-card")
async def a2a_agent_card(request: Request):
    """返回 A2A Agent Card。"""
    trace_id = _resolve_trace_id_from_request(request)
    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": "/a2a/agent-card",
        },
        tags=["api", "a2a"],
    )
    try:
        service = get_a2a()
        return {"ok": True, "agent_card": service.get_agent_card(), "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/a2a/agent-card", "runtime", exc)
        return {"ok": False, "error": f"a2a_card_error:{str(exc)}", "trace_id": trace_id}


@app.get("/a2a/agents/cards")
async def a2a_agent_cards(request: Request):
    """返回当前可发现的 Agent Card 列表。"""
    trace_id = _resolve_trace_id_from_request(request)
    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": "/a2a/agents/cards",
        },
        tags=["api", "a2a"],
    )
    try:
        return {"ok": True, "cards": all_agent_cards(), "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/a2a/agents/cards", "runtime", exc)
        return {"ok": False, "error": f"a2a_cards_error:{str(exc)}", "trace_id": trace_id}


@app.post("/a2a/rpc")
async def a2a_rpc(data: dict, request: Request):
    """A2A JSON-RPC 入口。"""
    trace_id = _resolve_trace_id_from_request(request)
    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": "/a2a/rpc",
            "rpc_method": str((data or {}).get("method") or ""),
        },
        tags=["api", "a2a"],
    )
    try:
        service = get_a2a()
        response = await service.ahandle_rpc(data, trace_id)
        return JSONResponse(status_code=200, content=response)
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/a2a/rpc", "runtime", exc)
        fallback = {
            "jsonrpc": "2.0",
            "id": (data or {}).get("id"),
            "error": {
                "code": -32000,
                "message": "Server error",
                "data": {"detail": str(exc), "trace_id": trace_id},
            },
        }
        return JSONResponse(status_code=200, content=fallback)


@app.get("/a2a/tasks")
async def a2a_list_tasks(
    request: Request,
    context_id: str = "",
    limit: int = 20,
):
    """列出 A2A 任务。

    Args:
        request: FastAPI 请求对象。
        context_id: 可选上下文 ID，用于筛选某个上下文下的任务。
        limit: 结果数量上限。

    Returns:
        Dict[str, Any]: 包含任务列表与 trace_id 的查询结果。
    """
    trace_id = _resolve_trace_id_from_request(request)
    try:
        service = get_a2a()
        rows = await service.alist_tasks(context_id, limit)
        return {"ok": True, "tasks": rows, "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/a2a/tasks", "runtime", exc, context_id=context_id)
        return {"ok": False, "error": f"a2a_list_tasks_error:{str(exc)}", "trace_id": trace_id}


@app.get("/a2a/tasks/{task_id}")
async def a2a_get_task(task_id: str, request: Request):
    """查询 A2A 任务状态。"""
    trace_id = _resolve_trace_id_from_request(request)
    try:
        service = get_a2a()
        task = await service.aget_task(task_id)
        if not task:
            return JSONResponse(
                status_code=404,
                content={"ok": False, "error": "task_not_found", "task_id": task_id, "trace_id": trace_id},
            )
        return {"ok": True, "task": task, "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/a2a/tasks/{task_id}", "runtime", exc, task_id=task_id)
        return {"ok": False, "error": f"a2a_get_task_error:{str(exc)}", "trace_id": trace_id}


@app.get("/a2a/tasks/{task_id}/events")
async def a2a_task_events(
    task_id: str,
    request: Request,
    cursor: int = 0,
    poll_ms: int = 500,
    timeout_seconds: int = 300,
):
    """A2A 任务事件流（SSE）。支持基于 cursor 断线续读。"""
    trace_id = _resolve_trace_id_from_request(request)
    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": "/a2a/tasks/{task_id}/events",
            "task_id": str(task_id or ""),
        },
        tags=["api", "a2a", "stream"],
    )
    try:
        service = get_a2a()
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/a2a/tasks/{task_id}/events", "runtime", exc, task_id=task_id)
        return JSONResponse(
            status_code=500,
            content={"ok": False, "error": f"a2a_stream_init_error:{str(exc)}", "trace_id": trace_id},
        )

    safe_cursor = max(0, int(cursor or 0))
    safe_poll_ms = max(100, min(int(poll_ms or 500), 5000))
    safe_timeout_seconds = max(5, min(int(timeout_seconds or 300), 3600))

    # 优先使用显式 cursor，未提供时尝试 Last-Event-ID
    if safe_cursor <= 0:
        last_event_id = str(request.headers.get("Last-Event-ID") or "").strip()
        try:
            if last_event_id:
                safe_cursor = max(0, int(last_event_id))
        except ValueError:
            safe_cursor = 0

    async def _event_stream():
        """生成 SSE 事件流。

        Returns:
            Any: SSE 事件字符串异步迭代器。
        """
        nonlocal safe_cursor
        started_at = time.time()
        while True:
            snapshot = await service.aget_task_updates(task_id, safe_cursor)
            if snapshot is None:
                payload = {"type": "error", "error": "task_not_found", "task_id": task_id}
                yield f"event: error\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"
                break

            events_obj = snapshot.get("events")
            events = events_obj if isinstance(events_obj, list) else []
            next_cursor = int(snapshot.get("next_cursor") or safe_cursor)
            base_cursor = next_cursor - len(events)

            for idx, event in enumerate(events, start=1):
                event_cursor = base_cursor + idx
                payload = {
                    "type": "progress",
                    "task_id": str(task_id or ""),
                    "cursor": event_cursor,
                    "event": event,
                }
                yield f"id: {event_cursor}\nevent: progress\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"

            safe_cursor = next_cursor
            state = str(snapshot.get("state") or "")
            terminal = bool(snapshot.get("terminal"))
            if terminal:
                task_payload = {
                    "type": "task",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                    "state": state,
                    "task": snapshot.get("task"),
                }
                yield f"id: {safe_cursor}\nevent: task\ndata: {json.dumps(task_payload, ensure_ascii=False)}\n\n"
                done_payload = {
                    "type": "done",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                    "state": state,
                }
                yield f"id: {safe_cursor}\nevent: done\ndata: {json.dumps(done_payload, ensure_ascii=False)}\n\n"
                break

            if time.time() - started_at >= safe_timeout_seconds:
                timeout_payload = {
                    "type": "stream_timeout",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                }
                yield f"id: {safe_cursor}\nevent: timeout\ndata: {json.dumps(timeout_payload, ensure_ascii=False)}\n\n"
                break

            if not events:
                heartbeat_payload = {
                    "type": "heartbeat",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                }
                yield f"id: {safe_cursor}\nevent: heartbeat\ndata: {json.dumps(heartbeat_payload, ensure_ascii=False)}\n\n"

            await asyncio.sleep(float(safe_poll_ms) / 1000.0)

    return StreamingResponse(
        _event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@app.post("/a2a/tasks/{task_id}/cancel")
async def a2a_cancel_task(task_id: str, request: Request, data: Optional[Dict[str, Any]] = None):
    """取消 A2A 任务。

    Args:
        task_id: 任务 ID。
        request: FastAPI 请求对象。
        data: 请求体数据字典。

    Returns:
        Dict[str, Any]: 取消请求处理结果，包含 ok/error 与 trace_id 信息。
    """
    trace_id = _resolve_trace_id_from_request(request)
    payload = data if isinstance(data, dict) else {}
    reason = str(payload.get("reason") or "manual_cancel")
    try:
        service = get_a2a()
        result = await service.acancel_task(task_id, reason)
        if not bool(result.get("ok")) and str(result.get("error") or "") == "task_not_found":
            return JSONResponse(
                status_code=404,
                content={"ok": False, "error": "task_not_found", "task_id": task_id, "trace_id": trace_id},
            )
        result["trace_id"] = trace_id
        return result
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/a2a/tasks/{task_id}/cancel", "runtime", exc, task_id=task_id)
        return {"ok": False, "error": f"a2a_cancel_task_error:{str(exc)}", "trace_id": trace_id}


def _build_a2a_rpc_fallback(data: Dict[str, Any], trace_id: str, detail: str) -> Dict[str, Any]:
    """构建 A2A RPC 兜底错误响应。

    Args:
        data: 请求体数据字典。
        trace_id: 请求追踪 ID。
        detail: 错误详情文本。

    Returns:
        Dict[str, Any]: JSON-RPC 标准错误响应对象。
    """
    return {
        "jsonrpc": "2.0",
        "id": (data or {}).get("id"),
        "error": {
            "code": -32000,
            "message": "Server error",
            "data": {"detail": str(detail or ""), "trace_id": trace_id},
        },
    }


async def _a2a_rpc_dispatch(
    data: Dict[str, Any],
    request: Request,
    service_getter: Callable[[], A2AService],
    endpoint: str,
) -> JSONResponse:
    """分发通用 A2A JSON-RPC 请求。

    Args:
        data: 请求体数据字典。
        request: FastAPI 请求对象。
        service_getter: A2A 服务获取函数。
        endpoint: 接口路径标识。

    Returns:
        JSONResponse: RPC 分发响应。
    """
    trace_id = _resolve_trace_id_from_request(request)
    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": endpoint,
            "rpc_method": str((data or {}).get("method") or ""),
        },
        tags=["api", "a2a"],
    )
    try:
        service = service_getter()
        response = await service.ahandle_rpc(data, trace_id)
        return JSONResponse(status_code=200, content=response)
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error(endpoint, "runtime", exc)
        fallback = _build_a2a_rpc_fallback(data=data, trace_id=trace_id, detail=str(exc))
        return JSONResponse(status_code=200, content=fallback)


async def _a2a_list_tasks_dispatch(
    request: Request,
    service_getter: Callable[[], A2AService],
    endpoint: str,
    context_id: str = "",
    limit: int = 20,
) -> Dict[str, Any]:
    """分发 A2A 任务列表查询。

    Args:
        request: FastAPI 请求对象。
        service_getter: A2A 服务获取函数。
        endpoint: 接口路径标识。
        context_id: 上下文 ID。
        limit: 结果数量上限。

    Returns:
        Dict[str, Any]: 任务列表查询结果。
    """
    trace_id = _resolve_trace_id_from_request(request)
    try:
        service = service_getter()
        rows = await service.alist_tasks(context_id, limit)
        return {"ok": True, "tasks": rows, "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error(endpoint, "runtime", exc, context_id=context_id)
        return {"ok": False, "error": f"a2a_list_tasks_error:{str(exc)}", "trace_id": trace_id}


async def _a2a_get_task_dispatch(
    task_id: str,
    request: Request,
    service_getter: Callable[[], A2AService],
    endpoint: str,
) -> Dict[str, Any] | JSONResponse:
    """分发 A2A 单任务查询。

    Args:
        task_id: 任务 ID。
        request: FastAPI 请求对象。
        service_getter: A2A 服务获取函数。
        endpoint: 接口路径标识。

    Returns:
        Dict[str, Any] | JSONResponse: 查询结果或 404 响应。
    """
    trace_id = _resolve_trace_id_from_request(request)
    try:
        service = service_getter()
        task = await service.aget_task(task_id)
        if not task:
            return JSONResponse(
                status_code=404,
                content={"ok": False, "error": "task_not_found", "task_id": task_id, "trace_id": trace_id},
            )
        return {"ok": True, "task": task, "trace_id": trace_id}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error(endpoint, "runtime", exc, task_id=task_id)
        return {"ok": False, "error": f"a2a_get_task_error:{str(exc)}", "trace_id": trace_id}


async def _a2a_task_events_dispatch(
    task_id: str,
    request: Request,
    service_getter: Callable[[], A2AService],
    endpoint: str,
    cursor: int = 0,
    poll_ms: int = 500,
    timeout_seconds: int = 300,
) -> StreamingResponse | JSONResponse:
    """分发 A2A 任务事件流请求。

    Args:
        task_id: 任务 ID。
        request: FastAPI 请求对象。
        service_getter: A2A 服务获取函数。
        endpoint: 接口路径标识。
        cursor: 事件游标位置。
        poll_ms: 轮询间隔（毫秒）。
        timeout_seconds: 轮询超时（秒）。

    Returns:
        StreamingResponse | JSONResponse: 事件流或错误响应。
    """
    trace_id = _resolve_trace_id_from_request(request)
    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": endpoint,
            "task_id": str(task_id or ""),
        },
        tags=["api", "a2a", "stream"],
    )
    try:
        service = service_getter()
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error(endpoint, "runtime", exc, task_id=task_id)
        return JSONResponse(
            status_code=500,
            content={"ok": False, "error": f"a2a_stream_init_error:{str(exc)}", "trace_id": trace_id},
        )

    safe_cursor = max(0, int(cursor or 0))
    safe_poll_ms = max(100, min(int(poll_ms or 500), 5000))
    safe_timeout_seconds = max(5, min(int(timeout_seconds or 300), 3600))

    if safe_cursor <= 0:
        last_event_id = str(request.headers.get("Last-Event-ID") or "").strip()
        try:
            if last_event_id:
                safe_cursor = max(0, int(last_event_id))
        except ValueError:
            safe_cursor = 0

    async def _event_stream():
        """生成 SSE 事件流。

        Returns:
            Any: SSE 事件字符串异步迭代器。
        """
        nonlocal safe_cursor
        started_at = time.time()
        while True:
            snapshot = await service.aget_task_updates(task_id, safe_cursor)
            if snapshot is None:
                payload = {"type": "error", "error": "task_not_found", "task_id": task_id}
                yield f"event: error\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"
                break

            events_obj = snapshot.get("events")
            events = events_obj if isinstance(events_obj, list) else []
            next_cursor = int(snapshot.get("next_cursor") or safe_cursor)
            base_cursor = next_cursor - len(events)

            for idx, event in enumerate(events, start=1):
                event_cursor = base_cursor + idx
                payload = {
                    "type": "progress",
                    "task_id": str(task_id or ""),
                    "cursor": event_cursor,
                    "event": event,
                }
                yield f"id: {event_cursor}\nevent: progress\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"

            safe_cursor = next_cursor
            state = str(snapshot.get("state") or "")
            terminal = bool(snapshot.get("terminal"))
            if terminal:
                task_payload = {
                    "type": "task",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                    "state": state,
                    "task": snapshot.get("task"),
                }
                yield f"id: {safe_cursor}\nevent: task\ndata: {json.dumps(task_payload, ensure_ascii=False)}\n\n"
                done_payload = {
                    "type": "done",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                    "state": state,
                }
                yield f"id: {safe_cursor}\nevent: done\ndata: {json.dumps(done_payload, ensure_ascii=False)}\n\n"
                break

            if time.time() - started_at >= safe_timeout_seconds:
                timeout_payload = {
                    "type": "stream_timeout",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                }
                yield f"id: {safe_cursor}\nevent: timeout\ndata: {json.dumps(timeout_payload, ensure_ascii=False)}\n\n"
                break

            if not events:
                heartbeat_payload = {
                    "type": "heartbeat",
                    "task_id": str(task_id or ""),
                    "cursor": safe_cursor,
                }
                yield f"id: {safe_cursor}\nevent: heartbeat\ndata: {json.dumps(heartbeat_payload, ensure_ascii=False)}\n\n"

            await asyncio.sleep(float(safe_poll_ms) / 1000.0)

    return StreamingResponse(
        _event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


async def _a2a_cancel_task_dispatch(
    task_id: str,
    request: Request,
    service_getter: Callable[[], A2AService],
    endpoint: str,
    data: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any] | JSONResponse:
    """分发 A2A 任务取消请求。

    Args:
        task_id: 任务 ID。
        request: FastAPI 请求对象。
        service_getter: A2A 服务获取函数。
        endpoint: 接口路径标识。
        data: 请求体数据字典。

    Returns:
        Dict[str, Any] | JSONResponse: 取消结果或错误响应。
    """
    trace_id = _resolve_trace_id_from_request(request)
    payload = data if isinstance(data, dict) else {}
    reason = str(payload.get("reason") or "manual_cancel")
    try:
        service = service_getter()
        result = await service.acancel_task(task_id, reason)
        if not bool(result.get("ok")) and str(result.get("error") or "") == "task_not_found":
            return JSONResponse(
                status_code=404,
                content={"ok": False, "error": "task_not_found", "task_id": task_id, "trace_id": trace_id},
            )
        result["trace_id"] = trace_id
        return result
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error(endpoint, "runtime", exc, task_id=task_id)
        return {"ok": False, "error": f"a2a_cancel_task_error:{str(exc)}", "trace_id": trace_id}


@app.post("/a2a/graph/rpc")
async def a2a_graph_rpc(data: dict, request: Request):
    """Graph 子代理 A2A JSON-RPC 入口。"""
    return await _a2a_rpc_dispatch(data, request, get_a2a_graph, "/a2a/graph/rpc")


@app.get("/a2a/graph/tasks")
async def a2a_graph_list_tasks(
    request: Request,
    context_id: str = "",
    limit: int = 20,
):
    """列出 Graph 子代理 A2A 任务。"""
    return await _a2a_list_tasks_dispatch(
        request=request,
        service_getter=get_a2a_graph,
        endpoint="/a2a/graph/tasks",
        context_id=context_id,
        limit=limit,
    )


@app.get("/a2a/graph/tasks/{task_id}")
async def a2a_graph_get_task(task_id: str, request: Request):
    """查询 Graph 子代理 A2A 任务状态。"""
    return await _a2a_get_task_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_graph,
        endpoint="/a2a/graph/tasks/{task_id}",
    )


@app.get("/a2a/graph/tasks/{task_id}/events")
async def a2a_graph_task_events(
    task_id: str,
    request: Request,
    cursor: int = 0,
    poll_ms: int = 500,
    timeout_seconds: int = 300,
):
    """Graph 子代理 A2A 任务事件流（SSE）。"""
    return await _a2a_task_events_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_graph,
        endpoint="/a2a/graph/tasks/{task_id}/events",
        cursor=cursor,
        poll_ms=poll_ms,
        timeout_seconds=timeout_seconds,
    )


@app.post("/a2a/graph/tasks/{task_id}/cancel")
async def a2a_graph_cancel_task(
    task_id: str,
    request: Request,
    data: Optional[Dict[str, Any]] = None,
):
    """取消 Graph 子代理 A2A 任务。"""
    return await _a2a_cancel_task_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_graph,
        endpoint="/a2a/graph/tasks/{task_id}/cancel",
        data=data,
    )


@app.post("/a2a/retrieval/rpc")
async def a2a_retrieval_rpc(data: dict, request: Request):
    """Retrieval 子代理 A2A JSON-RPC 入口。"""
    return await _a2a_rpc_dispatch(data, request, get_a2a_retrieval, "/a2a/retrieval/rpc")


@app.get("/a2a/retrieval/tasks")
async def a2a_retrieval_list_tasks(
    request: Request,
    context_id: str = "",
    limit: int = 20,
):
    """列出 Retrieval 子代理 A2A 任务。"""
    return await _a2a_list_tasks_dispatch(
        request=request,
        service_getter=get_a2a_retrieval,
        endpoint="/a2a/retrieval/tasks",
        context_id=context_id,
        limit=limit,
    )


@app.get("/a2a/retrieval/tasks/{task_id}")
async def a2a_retrieval_get_task(task_id: str, request: Request):
    """查询 Retrieval 子代理 A2A 任务状态。"""
    return await _a2a_get_task_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_retrieval,
        endpoint="/a2a/retrieval/tasks/{task_id}",
    )


@app.get("/a2a/retrieval/tasks/{task_id}/events")
async def a2a_retrieval_task_events(
    task_id: str,
    request: Request,
    cursor: int = 0,
    poll_ms: int = 500,
    timeout_seconds: int = 300,
):
    """Retrieval 子代理 A2A 任务事件流（SSE）。"""
    return await _a2a_task_events_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_retrieval,
        endpoint="/a2a/retrieval/tasks/{task_id}/events",
        cursor=cursor,
        poll_ms=poll_ms,
        timeout_seconds=timeout_seconds,
    )


@app.post("/a2a/retrieval/tasks/{task_id}/cancel")
async def a2a_retrieval_cancel_task(
    task_id: str,
    request: Request,
    data: Optional[Dict[str, Any]] = None,
):
    """取消 Retrieval 子代理 A2A 任务。"""
    return await _a2a_cancel_task_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_retrieval,
        endpoint="/a2a/retrieval/tasks/{task_id}/cancel",
        data=data,
    )


@app.post("/a2a/evaluator/rpc")
async def a2a_evaluator_rpc(data: dict, request: Request):
    """Evaluator 子代理 A2A JSON-RPC 入口。"""
    return await _a2a_rpc_dispatch(data, request, get_a2a_evaluator, "/a2a/evaluator/rpc")


@app.get("/a2a/evaluator/tasks")
async def a2a_evaluator_list_tasks(
    request: Request,
    context_id: str = "",
    limit: int = 20,
):
    """列出 Evaluator 子代理 A2A 任务。"""
    return await _a2a_list_tasks_dispatch(
        request=request,
        service_getter=get_a2a_evaluator,
        endpoint="/a2a/evaluator/tasks",
        context_id=context_id,
        limit=limit,
    )


@app.get("/a2a/evaluator/tasks/{task_id}")
async def a2a_evaluator_get_task(task_id: str, request: Request):
    """查询 Evaluator 子代理 A2A 任务状态。"""
    return await _a2a_get_task_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_evaluator,
        endpoint="/a2a/evaluator/tasks/{task_id}",
    )


@app.get("/a2a/evaluator/tasks/{task_id}/events")
async def a2a_evaluator_task_events(
    task_id: str,
    request: Request,
    cursor: int = 0,
    poll_ms: int = 500,
    timeout_seconds: int = 300,
):
    """Evaluator 子代理 A2A 任务事件流（SSE）。"""
    return await _a2a_task_events_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_evaluator,
        endpoint="/a2a/evaluator/tasks/{task_id}/events",
        cursor=cursor,
        poll_ms=poll_ms,
        timeout_seconds=timeout_seconds,
    )


@app.post("/a2a/evaluator/tasks/{task_id}/cancel")
async def a2a_evaluator_cancel_task(
    task_id: str,
    request: Request,
    data: Optional[Dict[str, Any]] = None,
):
    """取消 Evaluator 子代理 A2A 任务。"""
    return await _a2a_cancel_task_dispatch(
        task_id=task_id,
        request=request,
        service_getter=get_a2a_evaluator,
        endpoint="/a2a/evaluator/tasks/{task_id}/cancel",
        data=data,
    )


@app.get("/graph/visualization")
async def graph_visualization(keyword: str = "", limit: int = 200):
    """查询图谱可视化节点与边数据。

    Args:
        keyword: 关键词过滤条件（匹配节点名称）。
        limit: 结果数量上限。

    Returns:
        Dict[str, Any]: 图谱可视化数据。
    """
    safe_limit = max(10, min(limit, 500))
    q = (keyword or "").strip()
    try:
        if q:
            rows = await _await_call(
                neo4j_client.run_cypher,
                """
                MATCH (a)-[r]->(b)
                WHERE toString(a.name) CONTAINS $q OR toString(b.name) CONTAINS $q
                RETURN id(a) AS source_id, id(b) AS target_id,
                       labels(a) AS source_labels, labels(b) AS target_labels,
                       a.name AS source_name, b.name AS target_name,
                       type(r) AS rel_type
                LIMIT $limit
                """,
                {"q": q, "limit": safe_limit},
            )
        else:
            rows = await _await_call(
                neo4j_client.run_cypher,
                """
                MATCH (a)-[r]->(b)
                RETURN id(a) AS source_id, id(b) AS target_id,
                       labels(a) AS source_labels, labels(b) AS target_labels,
                       a.name AS source_name, b.name AS target_name,
                       type(r) AS rel_type
                LIMIT $limit
                """,
                {"limit": safe_limit},
            )
    except RuntimeError as exc:
        _log_api_error("/graph/visualization", "neo4j_query", exc, keyword=q, limit=safe_limit)
        return {"ok": False, "error": f"graph_query_error:{str(exc)}", "nodes": [], "links": []}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/graph/visualization", "unknown", exc, keyword=q, limit=safe_limit)
        return {"ok": False, "error": f"graph_unknown_error:{str(exc)}", "nodes": [], "links": []}

    node_map = {}
    links = []
    for row in rows:
        sid = str(row.get("source_id"))
        tid = str(row.get("target_id"))
        sname = row.get("source_name") or sid
        tname = row.get("target_name") or tid
        slabel = (row.get("source_labels") or ["Entity"])[0]
        tlabel = (row.get("target_labels") or ["Entity"])[0]
        if sid not in node_map:
            node_map[sid] = {"id": sid, "name": sname, "label": slabel}
        if tid not in node_map:
            node_map[tid] = {"id": tid, "name": tname, "label": tlabel}
        links.append(
            {
                "source": sid,
                "target": tid,
                "label": row.get("rel_type") or "RELATED_TO",
            }
        )

    return {"ok": True, "nodes": list(node_map.values()), "links": links, "count": len(links)}


@app.get("/graph/entity/{entity_id}")
async def get_entity_detail(entity_id: str):
    """获取实体详细信息。

    Args:
        entity_id: 实体 ID。

    Returns:
        Dict[str, Any]: 实体详情及其入边、出边关系数据。
    """
    try:
        # 获取实体基本信息
        entity_result = await _await_call(
            neo4j_client.run_cypher,
            """
            MATCH (n)
            WHERE id(n) = $entity_id
            RETURN id(n) AS id, labels(n) AS labels, n.name AS name, properties(n) AS props
            """,
            {"entity_id": int(entity_id)}
        )

        if not entity_result:
            return {"ok": False, "error": "实体不存在"}

        entity = entity_result[0]
        props = entity.get("props") or {}
        # 过滤掉 name 属性（已单独提取）
        node_properties = {k: v for k, v in props.items() if k != "name"}
        
        entity_data = {
            "id": str(entity["id"]),
            "labels": entity["labels"],
            "name": entity["name"],
            "properties": node_properties
        }

        # 获取入边关系
        incoming = await _await_call(
            neo4j_client.run_cypher,
            """
            MATCH (a)-[r]->(b)
            WHERE id(b) = $entity_id
            RETURN id(a) AS source_id, a.name AS source_name, labels(a) AS source_labels,
                   type(r) AS rel_type, properties(r) AS rel_props
            LIMIT 50
            """,
            {"entity_id": int(entity_id)}
        )

        # 获取出边关系
        outgoing = await _await_call(
            neo4j_client.run_cypher,
            """
            MATCH (a)-[r]->(b)
            WHERE id(a) = $entity_id
            RETURN id(b) AS target_id, b.name AS target_name, labels(b) AS target_labels,
                   type(r) AS rel_type, properties(r) AS rel_props
            LIMIT 50
            """,
            {"entity_id": int(entity_id)}
        )

        relationships = {
            "incoming": [
                {
                    "source_id": str(row["source_id"]),
                    "source_name": row["source_name"],
                    "source_labels": row["source_labels"],
                    "rel_type": row["rel_type"],
                    "rel_props": row.get("rel_props") or {}
                }
                for row in incoming
            ],
            "outgoing": [
                {
                    "target_id": str(row["target_id"]),
                    "target_name": row["target_name"],
                    "target_labels": row["target_labels"],
                    "rel_type": row["rel_type"],
                    "rel_props": row.get("rel_props") or {}
                }
                for row in outgoing
            ]
        }

        return {
            "ok": True,
            "entity": entity_data,
            "relationships": relationships
        }

    except ValueError as exc:
        _log_api_error("/graph/entity", "id_parse", exc, entity_id=entity_id)
        return {"ok": False, "error": f"invalid_entity_id:{str(exc)}"}
    except RuntimeError as exc:
        _log_api_error("/graph/entity", "neo4j_query", exc, entity_id=entity_id)
        return {"ok": False, "error": f"graph_query_error:{str(exc)}"}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/graph/entity", "unknown", exc, entity_id=entity_id)
        return {"ok": False, "error": f"graph_unknown_error:{str(exc)}"}


async def generate_data(input_text, user_id, session_id="", request_id: str = ""):
    """将 A2A 事件转为逐行 JSON 流输出。

    Args:
        input_text: 用户输入文本。
        user_id: 用户 ID。
        session_id: 会话 ID。
        request_id: 请求 ID。

    Returns:
        AsyncGenerator[str, None]: 每次产出一行 JSON 字符串，用于流式响应。
    """
    async for msg in _generate_a2a_stream_events(
        input_text=input_text,
        user_id=user_id,
        session_id=session_id,
        request_id=request_id,
    ):
        yield json.dumps(msg, ensure_ascii=False) + "\n"


# 会话历史 API
@app.post("/sessions")
async def create_session(data: dict):
    """创建新会话。

    Args:
        data: 请求体数据字典。

    Returns:
        Any: 操作结果。
    """
    user_id = data.get("user_id") or str(uuid.uuid4())
    title = data.get("title", "新会话")
    session = await _await_call(SessionManager.create_session, user_id, title)
    return {"ok": True, "data": session}


@app.get("/sessions")
async def list_sessions(user_id: str = "", limit: int = 50):
    """获取用户的会话列表。

    Args:
        user_id: 用户 ID。
        limit: 结果数量上限。

    Returns:
        Dict[str, Any]: 会话列表查询结果，data 字段为会话数组。
    """
    if not user_id:
        return {"ok": False, "error": "user_id_required", "data": []}
    sessions = await _await_call(SessionManager.get_sessions, user_id, limit)
    return {"ok": True, "data": sessions}


@app.get("/sessions/{session_id}")
async def get_session(session_id: str, user_id: str = ""):
    """获取单个会话详情。

    Args:
        session_id: 会话 ID。
        user_id: 用户 ID。

    Returns:
        Dict[str, Any]: 会话详情与消息列表，或未找到错误信息。
    """
    if not user_id:
        return {"ok": False, "error": "user_id_required"}

    session = await _await_call(SessionManager.get_session, session_id)
    if not session or str(session.get("user_id") or "") != user_id:
        return {"ok": False, "error": "session_not_found"}

    messages = await _await_call(MessageManager.get_messages, session_id)
    return {"ok": True, "data": {**session, "messages": messages}}


@app.put("/sessions/{session_id}")
async def update_session(session_id: str, data: dict):
    """更新会话标题。

    Args:
        session_id: 会话 ID。
        data: 请求体数据字典。

    Returns:
        Any: 操作结果。
    """
    title = data.get("title")
    user_id = str(data.get("user_id") or "").strip()
    if not user_id:
        return {"ok": False, "error": "user_id_required"}
    if not title:
        return {"ok": False, "error": "title_required"}

    if not await _is_session_owner(session_id, user_id):
        return {"ok": False, "error": "session_not_found"}

    success = await _await_call(SessionManager.update_session_title, session_id, title)
    if not success:
        return {"ok": False, "error": "session_not_found"}
    return {"ok": True}


@app.delete("/sessions/{session_id}")
async def delete_session(session_id: str, user_id: str = ""):
    """删除会话。

    Args:
        session_id: 会话 ID。
        user_id: 用户 ID。

    Returns:
        Any: 操作结果。
    """
    if not user_id:
        return {"ok": False, "error": "user_id_required"}

    if not await _is_session_owner(session_id, user_id):
        return {"ok": False, "error": "session_not_found"}

    success = await _await_call(SessionManager.delete_session, session_id)
    if not success:
        return {"ok": False, "error": "session_not_found"}
    return {"ok": True}


@app.post("/sessions/{session_id}/messages")
async def add_message(session_id: str, data: dict):
    """添加消息到会话。

    Args:
        session_id: 会话 ID。
        data: 请求体数据字典。

    Returns:
        Dict[str, Any]: 写入消息后的结果，data 字段为新消息记录。
    """
    user_id = str(data.get("user_id") or "").strip()
    role = data.get("role")
    content = data.get("content")
    think_content = data.get("think_content", "")
    
    if not user_id:
        return {"ok": False, "error": "user_id_required"}

    if not role or not content:
        return {"ok": False, "error": "role_and_content_required"}
    
    # 检查会话是否存在
    session = await _await_call(SessionManager.get_session, session_id)
    if not session or str(session.get("user_id") or "") != user_id:
        return {"ok": False, "error": "session_not_found"}
    
    message = await _await_call(MessageManager.add_message, session_id, role, content, think_content)
    return {"ok": True, "data": message}


@app.get("/sessions/{session_id}/messages")
async def get_messages(session_id: str, user_id: str = ""):
    """获取会话的所有消息。

    Args:
        session_id: 会话 ID。
        user_id: 用户 ID。

    Returns:
        Dict[str, Any]: 会话消息列表查询结果。
    """
    if not user_id:
        return {"ok": False, "error": "user_id_required"}

    session = await _await_call(SessionManager.get_session, session_id)
    if not session or str(session.get("user_id") or "") != user_id:
        return {"ok": False, "error": "session_not_found"}

    messages = await _await_call(MessageManager.get_messages, session_id)
    return {"ok": True, "data": messages}


@app.post("/sessions/{session_id}/rollback-last-question")
async def rollback_last_question(session_id: str, data: dict):
    """回撤最后一个用户问题，并删除其对应会话记忆摘要。"""
    user_id = str(data.get("user_id") or "").strip()
    if not user_id:
        return {"ok": False, "error": "user_id_required"}

    session = await _await_call(SessionManager.get_session, session_id)
    if not session or str(session.get("user_id") or "") != user_id:
        return {"ok": False, "error": "session_not_found"}

    rollback_result = await _await_call(MessageManager.rollback_last_question, session_id)
    if not bool(rollback_result.get("ok")):
        return {"ok": False, "error": "no_question_to_rollback"}

    await _await_call(SessionMemoryManager.clear_summary, session_id)
    return {
        "ok": True,
        "data": {
            "removed_count": int(rollback_result.get("removed_count") or 0),
            "question": str(rollback_result.get("question") or ""),
            "memory_cleared": True,
        },
    }


@app.post("/sessions/{session_id}/rollback-to-question")
async def rollback_to_question(session_id: str, data: dict):
    """回撤到指定用户问题，删除该问题及其之后的所有消息，并清空会话记忆摘要。"""
    user_id = str(data.get("user_id") or "").strip()
    if not user_id:
        return {"ok": False, "error": "user_id_required"}

    raw_message_id = data.get("message_id")
    if not isinstance(raw_message_id, int) or raw_message_id <= 0:
        return {"ok": False, "error": "message_id_required"}

    session = await _await_call(SessionManager.get_session, session_id)
    if not session or str(session.get("user_id") or "") != user_id:
        return {"ok": False, "error": "session_not_found"}

    rollback_result = await _await_call(MessageManager.rollback_to_question, session_id, raw_message_id)
    if not bool(rollback_result.get("ok")):
        return {"ok": False, "error": "question_not_found"}

    await _await_call(SessionMemoryManager.clear_summary, session_id)
    return {
        "ok": True,
        "data": {
            "removed_count": int(rollback_result.get("removed_count") or 0),
            "question": str(rollback_result.get("question") or ""),
            "memory_cleared": True,
        },
    }


# 文件上传接口
@app.post("/upload_image")
@app.post("/upload_file")
async def upload_image(file: UploadFile = File(...)):
    """上传聊天文件并返回文件URL；图片额外返回 base64 用于视觉问答。"""
    # 校验文件名非空
    if not file.filename:
        return {"ok": False, "error": "文件名不能为空"}

    image_extensions = {'.jpg', '.jpeg', '.png', '.gif', '.webp', '.bmp'}
    doc_extensions = {'.md', '.doc', '.docx', '.pdf', '.ppt', '.txt'}
    allowed_extensions = image_extensions | doc_extensions
    file_ext = Path(file.filename).suffix.lower()
    
    if file_ext not in allowed_extensions:
        return {"ok": False, "error": f"不支持的图片格式: {file_ext}"}
    
    try:
        import base64
        content = await file.read()
        
        file_id = str(uuid.uuid4())
        filename = f"{file_id}{file_ext}"
        file_path = UPLOAD_DIR / filename
        await _await_call(file_path.write_bytes, content)
        
        base64_data = ""
        mime_type = "application/octet-stream"
        if file_ext in image_extensions:
            base64_data = base64.b64encode(content).decode('utf-8')
            mime_type = f"image/{file_ext[1:]}"
            if file_ext == '.jpg':
                mime_type = "image/jpeg"

        return {
            "ok": True,
            "data": {
                "filename": filename,
                "original_name": str(file.filename),
                "url": f"/uploads/{filename}",
                "base64": f"data:{mime_type};base64,{base64_data}" if base64_data else "",
                "is_image": file_ext in image_extensions,
            }
        }
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/upload_image", "write_file", e, filename=file.filename)
        return {"ok": False, "error": f"上传失败: {str(e)}"}


@app.post("/process_with_session")
async def process_with_session(data: dict, request: Request):
    """处理问答并保存到会话历史"""
    trace_id = _resolve_trace_id_from_request(request)
    readiness_error = await _ensure_core_ready("/process_with_session")
    if readiness_error is not None:
        return readiness_error

    input_text = data.get("input", "")
    session_id = data.get("session_id", "")
    user_id = data.get("user_id") or ""
    image_data = data.get("image", None)
    file_name = str(data.get("file_name") or "").strip()
    
    if not user_id:
        user_id = str(uuid.uuid4())

    request.state.user_id = str(user_id)
    request.state.session_id = str(session_id or "")

    set_trace_metadata(
        {
            "trace_id": trace_id,
            "request_id": trace_id,
            "endpoint": "/process_with_session",
            "session_id": str(session_id or ""),
            "user_id": str(user_id or ""),
        },
        tags=["api", "process_with_session"],
    )
    limited = _enforce_heavy_api_rate_limit(
        endpoint="/process_with_session",
        user_id=str(user_id or ""),
        session_id=str(session_id or ""),
        request=request,
    )
    if limited is not None:
        return limited
    
    session: Optional[Dict[str, Any]] = None
    if not session_id:
        session = await _await_call(
            SessionManager.create_session,
            user_id,
            input_text[:30] + "..." if len(input_text) > 30 else input_text,
        )
        session_id = session["id"]
        request.state.session_id = str(session_id)
    else:
        session = await _await_call(SessionManager.get_session, session_id)
        if not session or str(session.get("user_id") or "") != user_id:
            return {"ok": False, "error": "session_not_found", "trace_id": trace_id}
    
    message_content = input_text
    if image_data:
        message_content = f"[图片] {input_text}"
    elif file_name:
        message_content = f"[文件:{file_name}] {input_text}".strip()
    
    user_message = await _await_call(MessageManager.add_message, session_id, "user", message_content)

    runtime_payload = {
        "original_input_text": str(input_text or ""),
        "resolved_query": {},
        "memory_context": {},
    }
    try:
        runtime_payload = await _await_call(
            get_session_memory_service().build_runtime_context,
            session_id,
            str(input_text or ""),
            int(user_message.get("id") or 0),
        )
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
        _log_api_error("/process_with_session", "build_runtime_context", exc, session_id=session_id)
    
    return StreamingResponse(
        generate_data_with_session(
            input_text,
            user_id,
            session_id,
            image_data,
            user_message_id=int(user_message.get("id") or 0),
            request_id=str(getattr(request.state, "request_id", "") or ""),
            original_input_text=str(runtime_payload.get("original_input_text") or input_text or ""),
            resolved_query=runtime_payload.get("resolved_query")
            if isinstance(runtime_payload.get("resolved_query"), dict)
            else {},
            memory_context=runtime_payload.get("memory_context")
            if isinstance(runtime_payload.get("memory_context"), dict)
            else {},
        ),
        media_type="text/plain"
    )


async def generate_data_with_session(
    input_text,
    user_id,
    session_id,
    image_data=None,
    user_message_id: int = 0,
    request_id: str = "",
    original_input_text: str = "",
    resolved_query: Optional[Dict[str, Any]] = None,
    memory_context: Optional[Dict[str, Any]] = None,
):
    """基于 A2A 生成流式响应并保存助手回复。"""
    _ = image_data
    _ = user_message_id
    full_reply = ""
    full_think = ""
    done_status = "done"
    done_processing_steps: List[str] = []
    done_tool_observations: List[Dict[str, Any]] = []
    done_evidence_summary: Dict[str, Any] = {}
    assistant_message_persisted = False
    assistant_message_id = 0
    memory_refresh_scheduled = False
    done_emitted = False

    async def persist_assistant_message_once() -> int:
        nonlocal assistant_message_persisted, assistant_message_id
        if assistant_message_persisted:
            return int(assistant_message_id or 0)
        if not full_reply:
            return 0
        assistant_message = await _await_call(
            MessageManager.add_message,
            session_id,
            "assistant",
            full_reply,
            full_think,
        )
        assistant_message_id = int((assistant_message or {}).get("id") or 0)
        assistant_message_persisted = assistant_message_id > 0
        return int(assistant_message_id or 0)

    def schedule_memory_refresh(persisted_assistant_message_id: int) -> None:
        nonlocal memory_refresh_scheduled
        if memory_refresh_scheduled or int(persisted_assistant_message_id or 0) <= 0:
            return

        turn_artifacts = {
            "processing_steps": done_processing_steps,
            "tool_observations": done_tool_observations,
            "evidence_summary": done_evidence_summary,
            "assistant_reply": full_reply,
        }

        async def _refresh_memory_task() -> None:
            await _await_call(
                get_session_memory_service().refresh_after_turn,
                session_id,
                int(persisted_assistant_message_id or 0),
                turn_artifacts,
            )

        def _on_refresh_done(task: asyncio.Task) -> None:
            try:
                task.result()
            except Exception as exc:
                _log_api_error(
                    "/process_with_session",
                    "memory_refresh_failed",
                    exc,
                    session_id=session_id,
                    user_id=user_id,
                    assistant_message_id=int(persisted_assistant_message_id or 0),
                    request_id=request_id,
                )

        try:
            refresh_task = asyncio.create_task(
                _refresh_memory_task(),
                name=f"memory-refresh-{session_id}-{persisted_assistant_message_id}",
            )
            refresh_task.add_done_callback(_on_refresh_done)
            memory_refresh_scheduled = True
        except Exception as exc:
            _log_api_error(
                "/process_with_session",
                "memory_refresh_schedule_failed",
                exc,
                session_id=session_id,
                user_id=user_id,
                assistant_message_id=int(persisted_assistant_message_id or 0),
                request_id=request_id,
            )

    try:
        async for msg in _generate_a2a_stream_events(
            input_text=input_text,
            user_id=user_id,
            session_id=session_id,
            request_id=request_id,
            original_input_text=original_input_text,
            resolved_query=resolved_query,
            memory_context=memory_context,
        ):
            if msg.get("type") == "reply":
                full_reply += msg.get("msg", "")
            elif msg.get("type") == "msg":
                full_think += msg.get("msg", "")
            elif msg.get("type") == "done":
                done_status = str(msg.get("status") or "done")
                processing_steps_obj = msg.get("processing_steps")
                if isinstance(processing_steps_obj, list):
                    done_processing_steps = [str(item or "") for item in processing_steps_obj]
                tool_observations_obj = msg.get("tool_observations")
                if isinstance(tool_observations_obj, list):
                    done_tool_observations = [
                        item if isinstance(item, dict) else {"raw": str(item or "")}
                        for item in tool_observations_obj
                    ]
                evidence_summary_obj = msg.get("evidence_summary")
                if isinstance(evidence_summary_obj, dict):
                    done_evidence_summary = evidence_summary_obj

            if msg.get("type") == "done" and not full_reply:
                full_reply = str(msg.get("msg") or "").strip()

            yield json.dumps(msg, ensure_ascii=False) + "\n"

            if msg.get("type") == "done":
                done_emitted = True
                if done_status == "done" and full_reply:
                    try:
                        persisted_id = await persist_assistant_message_once()
                        if persisted_id > 0:
                            schedule_memory_refresh(persisted_id)
                    except Exception as exc:
                        _log_api_error(
                            "/process_with_session",
                            "assistant_persist_failed_after_done",
                            exc,
                            session_id=session_id,
                            user_id=user_id,
                            request_id=request_id,
                            done_emitted=done_emitted,
                        )
                break
    except asyncio.CancelledError as exc:
        done_status = "cancelled"
        _log_api_error(
            "/process_with_session",
            "generate_data_with_session_cancelled",
            exc,
            session_id=session_id,
            user_id=user_id,
        )
        raise
    except Exception as exc:
        done_status = "error"
        _log_api_error(
            "/process_with_session",
            "generate_data_with_session_exception",
            exc,
            session_id=session_id,
            user_id=user_id,
        )
        if not done_emitted:
            yield json.dumps({"type": "done", "status": "error", "msg": "流式输出异常终止"}, ensure_ascii=False) + "\n"
            done_emitted = True
    finally:
        if full_reply and not assistant_message_persisted:
            try:
                persisted_id = await persist_assistant_message_once()
                if persisted_id > 0:
                    schedule_memory_refresh(persisted_id)
            except Exception as exc:
                _log_api_error(
                    "/process_with_session",
                    "persist_assistant_turn_finally",
                    exc,
                    session_id=session_id,
                    user_id=user_id,
                    request_id=request_id,
                    done_emitted=done_emitted,
                )


@app.post("/qa/jobs/{job_id}/cancel")
async def cancel_qa_job(job_id: str, data: dict):
    """兼容接口：取消 A2A 任务。"""
    user_id = str(data.get("user_id") or "").strip()
    session_id = str(data.get("session_id") or "").strip()
    if not user_id or not session_id:
        return {"ok": False, "error": "user_id_and_session_id_required"}

    service = get_a2a()
    task = await service.aget_task(job_id)
    if not task:
        return {"ok": False, "error": "job_not_found"}

    metadata_obj = task.get("metadata")
    metadata = metadata_obj if isinstance(metadata_obj, dict) else {}
    payload_user_id = str(metadata.get("user_id") or "")
    payload_session_id = str(task.get("contextId") or "")
    if payload_user_id != user_id or payload_session_id != session_id:
        return {"ok": False, "error": "job_not_found"}

    result = await service.acancel_task(job_id, "manual_cancel")
    if not bool(result.get("ok")) and str(result.get("error") or "") == "task_not_found":
        return {"ok": False, "error": "job_not_found_or_finished"}
    current_status = str((((result.get("task") or {}).get("status") or {}).get("state")) or "")

    return {
        "ok": True,
        "message": "cancel_requested",
        "job_id": job_id,
        "status": current_status,
        "memory_cleared": False,
    }


# 简化的管理员登录 API
@app.post("/admin/login")
async def admin_login(data: dict):
    """管理员登录。

    Args:
        data: 请求体数据字典。

    Returns:
        Dict[str, Any]: 登录结果，成功时返回管理员令牌。
    """
    username = data.get("username", "").strip()
    password = data.get("password", "")

    token = verify_login(username, password)
    if token:
        return {"ok": True, "message": "登录成功", "token": token}
    return {"ok": False, "error": "用户名或密码错误"}


@app.get("/admin/status")
async def admin_status():
    """检查登录状态。

    Returns:
        Dict[str, Any]: 当前登录状态与管理员用户名信息。
    """
    return {"ok": True, "logged_in": is_logged_in(), "user": get_current_admin_user()}


@app.get("/admin/runtime/flags")
async def admin_runtime_flags():
    """查看运行时特性开关（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    return {
        "ok": True,
        "data": {
            "flags": get_runtime_flag_snapshot(),
            "updated_at": datetime.now().isoformat(timespec="seconds"),
        },
    }


@app.get("/admin/perf/baseline")
async def admin_perf_baseline(days: int = 1):
    """生成最近 N 天性能基线报告（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    safe_days = max(1, min(int(days), 7))
    target_paths = [
        "/agent/run",
        "/process_with_session",
        "/mcp/invoke",
        "/admin/graph/",
    ]

    api_report = generate_api_baseline_report(
        safe_days,
        target_paths,
    )
    event_loop_lag = await _measure_event_loop_lag(samples=20, interval_ms=20)

    queue_snapshot = await _build_a2a_runtime_snapshot(limit_per_service=200)
    runtime_snapshot = get_runtime_flag_snapshot()

    return {
        "ok": True,
        "data": {
            "window_days": safe_days,
            "api_latency": api_report,
            "event_loop_lag_ms": event_loop_lag,
            "queue": queue_snapshot,
            "runtime_flags": runtime_snapshot,
            "active_threads": threading.active_count(),
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        },
    }


@app.get("/admin/qa/queue/status")
async def admin_qa_queue_status():
    """获取 A2A 任务状态快照（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    return {"ok": True, "data": await _build_a2a_runtime_snapshot(limit_per_service=200)}


@app.post("/admin/qa/jobs/{job_id}/cancel")
async def admin_cancel_qa_job(job_id: str):
    """取消 A2A 任务（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    service = get_a2a()
    result = await service.acancel_task(job_id, "manual_cancel")
    if not bool(result.get("ok")) and str(result.get("error") or "") == "task_not_found":
        return {"ok": False, "error": "job_not_found_or_finished"}
    state = str((((result.get("task") or {}).get("status") or {}).get("state")) or "")
    return {"ok": True, "message": "cancel_requested", "job_id": job_id, "status": state}


@app.post("/admin/logout")
async def admin_logout():
    """退出登录。

    Returns:
        Dict[str, Any]: 退出登录处理结果。
    """
    logout()
    return {"ok": True, "message": "已退出登录"}


# 管理员后台 API
@app.get("/admin/sessions")
async def admin_get_sessions(limit: int = 100, offset: int = 0):
    """获取所有会话列表（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    sessions = await _await_call(SessionManager.get_all_sessions, limit, offset)
    total = await _await_call(SessionManager.get_total_sessions_count)
    return {"ok": True, "data": sessions, "total": total}


@app.get("/admin/sessions/{session_id}/messages")
async def admin_get_session_messages(session_id: str):
    """获取指定会话的所有消息（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    session = await _await_call(SessionManager.get_session, session_id)
    if not session:
        return {"ok": False, "error": "会话不存在"}
    
    messages = await _await_call(MessageManager.get_messages, session_id)
    return {"ok": True, "session": session, "messages": messages}


@app.get("/admin/sessions/{session_id}/memory")
async def admin_get_session_memory(session_id: str):
    """获取指定会话记忆快照（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    session = await _await_call(SessionManager.get_session, session_id)
    if not session:
        return {"ok": False, "error": "会话不存在"}

    snapshot = await _await_call(get_session_memory_service().load_memory, session_id)
    return {"ok": True, "data": snapshot}


@app.post("/admin/sessions/{session_id}/memory/rebuild")
async def admin_rebuild_session_memory(session_id: str):
    """重建指定会话记忆（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    session = await _await_call(SessionManager.get_session, session_id)
    if not session:
        return {"ok": False, "error": "会话不存在"}

    await _await_call(get_session_memory_service().clear_memory, session_id)
    snapshot = await _await_call(get_session_memory_service().rebuild_from_messages, session_id)
    return {"ok": True, "data": snapshot}


@app.delete("/admin/sessions/{session_id}/messages")
async def admin_clear_session_messages(session_id: str):
    """清空指定会话的历史消息（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    session = await _await_call(SessionManager.get_session, session_id)
    if not session:
        return {"ok": False, "error": "会话不存在"}

    existing_messages = await _await_call(MessageManager.get_messages, session_id)
    removed_count = len(existing_messages)

    if removed_count > 0:
        await _await_call(MessageManager.delete_messages, session_id)

    await _await_call(SessionMemoryManager.clear_summary, session_id)
    await _await_call(SessionManager.update_session_time, session_id)

    return {
        "ok": True,
        "message": "会话历史记录已清空",
        "removed_count": removed_count
    }


@app.delete("/admin/sessions/{session_id}")
async def admin_delete_session(session_id: str):
    """删除会话（需要登录）。

    Args:
        session_id: 会话 ID。

    Returns:
        Dict[str, Any]: 删除结果，包含成功标识与提示信息。
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    success = await _await_call(SessionManager.delete_session, session_id)
    if success:
        return {"ok": True, "message": "会话已删除"}
    return {"ok": False, "error": "删除失败或会话不存在"}


@app.get("/admin/stats")
async def admin_get_stats():
    """获取统计数据（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    sessions_row = await fetch_one("SELECT COUNT(*) as count FROM sessions")
    messages_row = await fetch_one("SELECT COUNT(*) as count FROM messages")
    today_sessions_row = await fetch_one("SELECT COUNT(*) as count FROM sessions WHERE DATE(created_at) = CURDATE()")
    today_messages_row = await fetch_one("SELECT COUNT(*) as count FROM messages WHERE DATE(created_at) = CURDATE()")
    stats = {
        "session_count": int((sessions_row or {}).get("count") or 0),
        "message_count": int((messages_row or {}).get("count") or 0),
        "today_sessions": int((today_sessions_row or {}).get("count") or 0),
        "today_messages": int((today_messages_row or {}).get("count") or 0),
    }
    
    return {
        "ok": True,
        "data": {
            "total_sessions": stats["session_count"],
            "total_messages": stats["message_count"],
            "today_sessions": stats["today_sessions"],
            "today_messages": stats["today_messages"]
        }
    }


def _read_perf_event_logs(event_name: str, limit: int = 200) -> list[dict]:
    """从 logs/perf.log 中读取指定性能事件，按时间倒序返回。"""
    # 与 perf_logger.py 保持一致：固定读取项目根目录/logs/perf.log
    project_root = Path(__file__).resolve().parents[2]
    log_dir = project_root / "logs"

    candidate_paths = [
        log_dir / "perf.log",
        # 兼容日志轮转文件
        log_dir / "perf.log.1",
        log_dir / "perf.log.2",
        # 兼容旧相对路径运行方式
        Path("logs") / "perf.log",
    ]

    lines: list[str] = []
    for path in candidate_paths:
        if not path.exists():
            continue
        try:
            lines.extend(path.read_text(encoding="utf-8", errors="ignore").splitlines())
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            _log_api_error("_read_perf_event_logs", "read_file", exc, path=str(path))
            continue

    if not lines:
        return []

    events: list[dict] = []
    marker = "[性能事件]"

    for line in reversed(lines):
        if marker not in line:
            continue

        marker_index = line.find(marker)
        if marker_index == -1:
            continue

        payload_text = line[marker_index + len(marker):].strip()
        json_start = payload_text.find("{")
        if json_start == -1:
            continue

        payload_text = payload_text[json_start:]
        try:
            payload = json.loads(payload_text)
        except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as exc:
            _log_api_error("_read_perf_event_logs", "json_parse", exc)
            continue

        if payload.get("event") != event_name:
            continue

        events.append(payload)
        if len(events) >= limit:
            break

    return events


def _split_error_types(
    event: dict,
    strategy_error_keys: set[str],
) -> tuple[dict[str, int], dict[str, int]]:
    """拆分策略错误与执行错误统计。

    Args:
        event: 事件字典。
        strategy_error_keys: 策略错误类型键集合。

    Returns:
        tuple[dict[str, int], dict[str, int]]: (策略错误统计, 执行错误统计)。
    """
    strategy_errors = event.get("strategy_error_types") or {}
    execution_errors = event.get("execution_error_types") or {}

    if isinstance(strategy_errors, dict) and isinstance(execution_errors, dict):
        return (
            {str(k): int(v or 0) for k, v in strategy_errors.items()},
            {str(k): int(v or 0) for k, v in execution_errors.items()},
        )

    raw_errors = event.get("error_types") or {}
    strategy_fallback: dict[str, int] = {}
    execution_fallback: dict[str, int] = {}
    if isinstance(raw_errors, dict):
        for key, value in raw_errors.items():
            error_key = str(key)
            error_count = int(value or 0)
            if error_key in strategy_error_keys:
                strategy_fallback[error_key] = strategy_fallback.get(error_key, 0) + error_count
            else:
                execution_fallback[error_key] = execution_fallback.get(error_key, 0) + error_count
    return strategy_fallback, execution_fallback


def _to_float_metric(value: Any) -> float:
    """将监控值转换为浮点数。

    Args:
        value: 待规范化值。

    Returns:
        float: 转换后的浮点指标值。
    """
    if isinstance(value, bool):
        return float(int(value))
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            return 0.0
    return 0.0


def _to_int_metric(value: Any) -> int:
    """将监控值转换为整数。

    Args:
        value: 待规范化值。

    Returns:
        int: 转换后的整型指标值。
    """
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    if isinstance(value, str):
        try:
            return int(float(value))
        except ValueError:
            return 0
    return 0


@app.get("/admin/agent/observability-summary")
async def admin_agent_observability_summary(limit: int = 200):
    """获取 Agent 可观测性汇总数据（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    safe_limit = max(10, min(limit, 2000))
    events = _read_perf_event_logs("iterative_agent_run", limit=safe_limit)

    if not events:
        return {
            "ok": True,
            "data": {
                "total_runs": 0,
                "avg_duration_ms": 0.0,
                "avg_execution_tool_success_rate": 0.0,
                "avg_attempted_action_success_rate": 0.0,
                "avg_tool_success_rate": 0.0,
                "total_tokens": 0,
                "total_llm_calls": 0,
                "total_attempted_tool_actions": 0,
                "total_blocked_tool_actions": 0,
                "total_tool_calls": 0,
                "error_types": {},
                "strategy_error_types": {},
                "execution_error_types": {},
            },
        }

    total_runs = len(events)
    total_duration = 0.0
    total_success_rate = 0.0
    total_execution_success_rate = 0.0
    total_attempted_success_rate = 0.0
    total_tokens = 0
    total_llm_calls = 0
    total_attempted_tool_actions = 0
    total_blocked_tool_actions = 0
    total_tool_calls = 0
    error_counter: dict[str, int] = defaultdict(int)
    strategy_error_counter: dict[str, int] = defaultdict(int)
    execution_error_counter: dict[str, int] = defaultdict(int)
    strategy_error_keys = {
        "planner_error",
        "planner_upstream_error",
        "tool_not_allowed",
        "tool_temporarily_avoided",
        "duplicate_tool_call_blocked",
        "loop_guard_stop",
        "job_cancelled",
        "final_summarize_error",
    }

    for event in events:
        execution_rate_value = (
            event.get("execution_tool_success_rate")
            if event.get("execution_tool_success_rate") is not None
            else event.get("tool_success_rate")
        )
        attempted_rate_value = (
            event.get("attempted_action_success_rate")
            if event.get("attempted_action_success_rate") is not None
            else event.get("tool_success_rate")
        )
        attempted_actions_value = (
            event.get("attempted_tool_actions")
            if event.get("attempted_tool_actions") is not None
            else event.get("tool_calls")
        )

        total_duration += _to_float_metric(event.get("duration_ms"))
        total_success_rate += _to_float_metric(event.get("tool_success_rate"))
        total_execution_success_rate += _to_float_metric(execution_rate_value)
        total_attempted_success_rate += _to_float_metric(attempted_rate_value)
        total_tokens += _to_int_metric(event.get("total_tokens"))
        total_llm_calls += _to_int_metric(event.get("llm_calls"))
        total_attempted_tool_actions += _to_int_metric(attempted_actions_value)
        total_blocked_tool_actions += _to_int_metric(event.get("blocked_tool_actions"))
        total_tool_calls += _to_int_metric(event.get("tool_calls"))

        event_errors = event.get("error_types") or {}
        if isinstance(event_errors, dict):
            for key, value in event_errors.items():
                error_counter[str(key)] += int(value or 0)

        strategy_errors, execution_errors = _split_error_types(event, strategy_error_keys)
        for key, value in strategy_errors.items():
            strategy_error_counter[key] += value
        for key, value in execution_errors.items():
            execution_error_counter[key] += value

    return {
        "ok": True,
        "data": {
            "total_runs": total_runs,
            "avg_duration_ms": round(total_duration / total_runs, 2),
            "avg_execution_tool_success_rate": round(total_execution_success_rate / total_runs, 4),
            "avg_attempted_action_success_rate": round(total_attempted_success_rate / total_runs, 4),
            "avg_tool_success_rate": round(total_success_rate / total_runs, 4),
            "total_tokens": total_tokens,
            "total_llm_calls": total_llm_calls,
            "total_attempted_tool_actions": total_attempted_tool_actions,
            "total_blocked_tool_actions": total_blocked_tool_actions,
            "total_tool_calls": total_tool_calls,
            "error_types": dict(error_counter),
            "strategy_error_types": dict(strategy_error_counter),
            "execution_error_types": dict(execution_error_counter),
        },
    }


@app.get("/admin/agent/observability-runs")
async def admin_agent_observability_runs(limit: int = 50):
    """获取 Agent 最近运行可观测记录（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    safe_limit = max(1, min(limit, 200))
    events = _read_perf_event_logs("iterative_agent_run", limit=safe_limit)
    strategy_error_keys = {
        "planner_error",
        "planner_upstream_error",
        "tool_not_allowed",
        "tool_temporarily_avoided",
        "duplicate_tool_call_blocked",
        "loop_guard_stop",
        "job_cancelled",
        "final_summarize_error",
    }

    rows = []
    for index, event in enumerate(events):
        strategy_errors, execution_errors = _split_error_types(event, strategy_error_keys)
        rows.append(
            {
                "rank": index + 1,
                "timestamp": str(event.get("timestamp") or ""),
                "session_id": str(event.get("session_id") or ""),
                "user_id": str(event.get("user_id") or ""),
                "duration_ms": _to_float_metric(event.get("duration_ms")),
                "llm_calls": _to_int_metric(event.get("llm_calls")),
                "attempted_tool_actions": _to_int_metric(
                    event.get("attempted_tool_actions")
                    if event.get("attempted_tool_actions") is not None
                    else event.get("tool_calls")
                ),
                "blocked_tool_actions": _to_int_metric(event.get("blocked_tool_actions")),
                "tool_calls": _to_int_metric(event.get("tool_calls")),
                "execution_tool_success_rate": _to_float_metric(
                    event.get("execution_tool_success_rate")
                    if event.get("execution_tool_success_rate") is not None
                    else event.get("tool_success_rate")
                ),
                "attempted_action_success_rate": _to_float_metric(
                    event.get("attempted_action_success_rate")
                    if event.get("attempted_action_success_rate") is not None
                    else event.get("tool_success_rate")
                ),
                "tool_success_rate": _to_float_metric(event.get("tool_success_rate")),
                "total_tokens": _to_int_metric(event.get("total_tokens")),
                "error_types": event.get("error_types") or {},
                "strategy_error_types": strategy_errors,
                "execution_error_types": execution_errors,
                "tool_observations": event.get("tool_observations") if isinstance(event.get("tool_observations"), list) else [],
                "reflection_events": event.get("reflection_events") if isinstance(event.get("reflection_events"), list) else [],
                "turn_metrics": event.get("turn_metrics") if isinstance(event.get("turn_metrics"), list) else [],
                "processing_steps": event.get("processing_steps") if isinstance(event.get("processing_steps"), list) else [],
            }
        )

    return {
        "ok": True,
        "data": rows,
        "count": len(rows),
    }


# 文档上传和嵌入 API
@app.post("/admin/documents/upload")
async def upload_document(file: UploadFile = File(...)):
    """上传DOCX文档并进行向量嵌入（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    # 校验文件名非空
    if not file.filename:
        return {"ok": False, "error": "文件名不能为空"}

    # 检查文件类型
    if not file.filename.endswith('.docx'):
        return {"ok": False, "error": "只支持DOCX格式的文件"}
    
    try:
        # 读取文件内容
        content = await file.read()
        
        # 导入文档嵌入器
        from app.rag_core.document_embedding import embed_docx_bytes
        vector_store = await _await_call(get_shared_vector_store)

        # 分块+嵌入联动
        result = await _await_call(embed_docx_bytes, content, file.filename, vector_store)
        
        return {"ok": True, "data": result}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/documents/upload", "doc_embed", e, filename=file.filename)
        return {"ok": False, "error": f"文档处理失败: {str(e)}"}


@app.get("/admin/documents/stats")
async def get_document_stats():
    """获取文档向量库统计信息（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        from app.rag_core.document_embedding import DocumentEmbedder
        vector_store = await _await_call(get_shared_vector_store)
        embedder = DocumentEmbedder(vector_store=vector_store)
        stats = await _await_call(embedder.get_collection_stats)
        
        return stats
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/documents/stats", "get_collection_stats", e)
        return {"ok": False, "error": f"获取统计信息失败: {str(e)}"}


# Docling 文档解析 API
@app.post("/admin/docling/parse")
async def docling_parse_document(file: UploadFile = File(...)):
    """
    使用 Docling 解析文档（需要登录）
    支持 PDF、DOCX、PPTX、HTML、图片等多种格式
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    # 校验文件名非空
    if not file.filename:
        return {"ok": False, "error": "文件名不能为空"}

    # 检查文件类型
    allowed_extensions = {'.pdf', '.docx', '.doc', '.pptx', '.ppt', '.html', '.htm', '.txt', '.md', '.jpg', '.jpeg', '.png'}
    file_ext = Path(file.filename).suffix.lower()
    
    if file_ext not in allowed_extensions:
        return {"ok": False, "error": f"不支持的文件格式: {file_ext}。支持的格式: {', '.join(allowed_extensions)}"}
    
    try:
        # 读取文件内容
        content = await file.read()
        
        # 使用 Docling 解析
        from app.document_parser.docling_parser import DoclingParser
        
        parser = await _await_call(DoclingParser)
        result = await _await_call(parser.parse_bytes, content, file.filename)
        
        return {
            "ok": True,
            "data": {
                "filename": result.source,
                "format": result.format,
                "markdown": result.markdown,
                "text": result.text,
                "pages": result.pages,
                "tables_count": len(result.tables),
                "images_count": len(result.images),
                "metadata": result.metadata
            }
        }
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/docling/parse", "docling_parse", e, filename=file.filename)
        return {"ok": False, "error": f"Docling 解析失败: {str(e)}"}


@app.post("/admin/docling/parse-and-embed")
async def docling_parse_and_embed(file: UploadFile = File(...)):
    """
    使用 Docling 解析文档并嵌入向量库（需要登录）
    支持 PDF、DOCX、PPTX、HTML 等多种格式
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    # 校验文件名非空
    if not file.filename:
        return {"ok": False, "error": "文件名不能为空"}

    # 检查文件类型
    allowed_extensions = {'.pdf', '.docx', '.doc', '.pptx', '.ppt', '.html', '.htm', '.txt', '.md'}
    file_ext = Path(file.filename).suffix.lower()
    
    if file_ext not in allowed_extensions:
        return {"ok": False, "error": f"不支持的文件格式: {file_ext}。支持的格式: {', '.join(allowed_extensions)}"}
    
    try:
        # 读取文件内容
        content = await file.read()
        
        # 使用 Docling 解析
        from app.document_parser.docling_parser import DoclingParser
        from app.document_parser.document_segmentation import DocumentSegmentation
        from app.rag_core.document_embedding import DocumentEmbedder
        from app.common.config import Config
        
        # 1) 仅负责 Docling + OCR 文本抽取
        parser = await _await_call(DoclingParser)
        parsed = await _await_call(parser.parse_bytes, content, file.filename)

        # 2) 分块模块负责切块
        conf = Config()
        chunker = DocumentSegmentation(
            parent_chunk_size=conf.PARENT_CHUNK_SIZE,
            sub_chunk_size=conf.SUB_CHUNK_SIZE,
            parent_chunk_overlap=conf.PARENT_CHUNK_OVERLAP,
            sub_chunk_overlap=conf.SUB_CHUNK_OVERLAP
        )
        chunk_source_text = (parsed.text or "").strip()
        if not chunk_source_text:
            chunk_source_text = (parsed.markdown or "").strip()
        if not chunk_source_text:
            raise ValueError("Docling 解析结果为空，无法分块并嵌入")

        hierarchical_chunks = await _await_call(
            chunker.segment_text_with_page_count,
            chunk_source_text,
            file.filename,
            len(parsed.pages),
        )
        sub_chunks = await _await_call(chunker.get_all_sub_chunks, hierarchical_chunks)
        if not sub_chunks:
            raise ValueError("分块结果为空，未执行 Milvus 写入")

        # 3) 嵌入模块仅负责写向量
        vector_store = await _await_call(get_shared_vector_store)
        embedder = DocumentEmbedder(vector_store=vector_store)
        embed_result = await _await_call(embedder.embed_docling_sub_chunks, sub_chunks, file.filename, parsed)
        embed_result["parent_chunks_count"] = len(hierarchical_chunks)
        collection_stats = await _await_call(embedder.get_collection_stats)
        
        return {
            "ok": True,
            "data": {
                "filename": file.filename,
                "format": embed_result.get("format"),
                "pages": embed_result.get("pages", 0),
                "tables": embed_result.get("tables", 0),
                "images": embed_result.get("images", 0),
                "parent_chunks_count": embed_result.get("parent_chunks_count", 0),
                "sub_chunks_count": embed_result.get("sub_chunks_count", 0),
                "milvus_row_count": collection_stats.get("row_count", 0) if collection_stats.get("success") else None,
                "message": embed_result.get("message", "Docling 解析并嵌入完成")
            }
        }
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/docling/parse-and-embed", "docling_embed", e, filename=file.filename)
        return {"ok": False, "error": f"Docling 解析并嵌入失败: {str(e)}"}


# 知识图谱管理 API
# 知识图谱管理 API
@app.get("/admin/graph/labels")
async def get_graph_labels():
    """获取所有节点标签（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        query = """
        MATCH (n)
        UNWIND labels(n) AS label
        RETURN DISTINCT label
        ORDER BY label
        """
        result = await _await_call(neo4j_client.run_cypher, query)
        labels = [record["label"] for record in result]
        return {"ok": True, "data": labels}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/labels", "neo4j_query", e)
        return {"ok": False, "error": str(e)}


@app.get("/admin/graph/relationship-types")
async def get_graph_relationship_types():
    """获取所有关系类型（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        query = """
        MATCH ()-[r]->()
        RETURN DISTINCT type(r) AS rel_type
        ORDER BY rel_type
        """
        result = await _await_call(neo4j_client.run_cypher, query)
        rel_types = [record["rel_type"] for record in result]
        return {"ok": True, "data": rel_types}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/relationship-types", "neo4j_query", e)
        return {"ok": False, "error": str(e)}


@app.get("/admin/graph/nodes")
async def get_graph_nodes(label: str = "", keyword: str = "", limit: int = 100, offset: int = 0):
    """获取节点列表（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        if label:
            if keyword:
                query = f"""
                MATCH (n:{label})
                WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                ORDER BY n.name
                SKIP $offset LIMIT $limit
                """
                count_query = f"""
                MATCH (n:{label})
                WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                RETURN count(n) AS total
                """
            else:
                query = f"""
                MATCH (n:{label})
                RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                ORDER BY n.name
                SKIP $offset LIMIT $limit
                """
                count_query = f"""
                MATCH (n:{label})
                RETURN count(n) AS total
                """
        else:
            if keyword:
                query = """
                MATCH (n)
                WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                ORDER BY n.name
                SKIP $offset LIMIT $limit
                """
                count_query = """
                MATCH (n)
                WHERE n.name CONTAINS $keyword OR n.title CONTAINS $keyword
                RETURN count(n) AS total
                """
            else:
                query = """
                MATCH (n)
                RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
                ORDER BY n.name
                SKIP $offset LIMIT $limit
                """
                count_query = """
                MATCH (n)
                RETURN count(n) AS total
                """
        
        params = {"keyword": keyword, "offset": offset, "limit": limit}
        result = await _await_call(neo4j_client.run_cypher, query, params)
        
        count_result = await _await_call(neo4j_client.run_cypher, count_query, {"keyword": keyword} if keyword else {})
        total = count_result[0]["total"] if count_result else 0
        
        return {"ok": True, "data": result, "total": total}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/nodes", "neo4j_query", e, label=label, keyword=keyword)
        return {"ok": False, "error": str(e)}


@app.get("/admin/graph/nodes/{node_id}")
async def get_graph_node(node_id: int):
    """获取单个节点详情（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        # 获取节点信息
        node_query = """
        MATCH (n)
        WHERE id(n) = $node_id
        RETURN id(n) AS id, labels(n) AS labels, properties(n) AS properties
        """
        node_result = await _await_call(neo4j_client.run_cypher, node_query, {"node_id": node_id})
        
        if not node_result:
            return {"ok": False, "error": "节点不存在"}
        
        # 获取入边关系
        incoming_query = """
        MATCH (n)<-[r]-(m)
        WHERE id(n) = $node_id
        RETURN id(r) AS rel_id, type(r) AS rel_type, properties(r) AS rel_props,
               id(m) AS source_id, labels(m) AS source_labels, m.name AS source_name
        """
        incoming = await _await_call(neo4j_client.run_cypher, incoming_query, {"node_id": node_id})
        
        # 获取出边关系
        outgoing_query = """
        MATCH (n)-[r]->(m)
        WHERE id(n) = $node_id
        RETURN id(r) AS rel_id, type(r) AS rel_type, properties(r) AS rel_props,
               id(m) AS target_id, labels(m) AS target_labels, m.name AS target_name
        """
        outgoing = await _await_call(neo4j_client.run_cypher, outgoing_query, {"node_id": node_id})
        
        return {
            "ok": True,
            "data": {
                "node": node_result[0],
                "relationships": {
                    "incoming": incoming,
                    "outgoing": outgoing
                }
            }
        }
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/nodes/{node_id}", "neo4j_query", e, node_id=node_id)
        return {"ok": False, "error": str(e)}


@app.post("/admin/graph/nodes")
async def create_graph_node(data: dict):
    """创建新节点（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        raw_label = str(data.get("label", "")).strip()
        label = _validate_cypher_identifier(raw_label)
        properties = _sanitize_graph_properties(data.get("properties", {}))
        
        if not raw_label:
            return {"ok": False, "error": "节点标签不能为空"}
        if not label:
            return {"ok": False, "error": "节点标签非法，仅允许字母数字下划线且不能以数字开头"}
        
        # 构建属性字符串
        if properties:
            props_str = ", ".join([f"`{k}`: ${k}" for k in properties.keys()])
            query = f"CREATE (n:`{label}` {{{props_str}}}) RETURN id(n) AS id"
        else:
            query = f"CREATE (n:`{label}`) RETURN id(n) AS id"
        
        result = await _await_call(neo4j_client.run_cypher, query, properties)
        
        if result:
            return {"ok": True, "data": {"id": result[0]["id"], "message": "节点创建成功"}}
        return {"ok": False, "error": "创建失败"}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/nodes", "create_node", e, label=label)
        return {"ok": False, "error": str(e)}


@app.put("/admin/graph/nodes/{node_id}")
async def update_graph_node(node_id: int, data: dict):
    """更新节点（需要登录）。

    Args:
        node_id: 要更新的图谱节点内部 ID。
        data: 请求体数据字典。

    Returns:
        Any: 操作结果。
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        properties = _sanitize_graph_properties(data.get("properties", {}))
        
        if not properties:
            return {"ok": False, "error": "属性不能为空"}
        
        # 构建 SET 子句
        set_clauses = [f"n.`{k}` = ${k}" for k in properties.keys()]
        set_str = ", ".join(set_clauses)
        
        query = f"""
        MATCH (n)
        WHERE id(n) = $node_id
        SET {set_str}
        RETURN id(n) AS id, properties(n) AS properties
        """
        
        params = {"node_id": node_id, **properties}
        result = await _await_call(neo4j_client.run_cypher, query, params)
        
        if result:
            return {"ok": True, "data": {"id": result[0]["id"], "message": "节点更新成功"}}
        return {"ok": False, "error": "节点不存在或更新失败"}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/nodes/{node_id}", "update_node", e, node_id=node_id)
        return {"ok": False, "error": str(e)}


@app.delete("/admin/graph/nodes/{node_id}")
async def delete_graph_node(node_id: int):
    """删除节点及其关系（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        query = """
        MATCH (n)
        WHERE id(n) = $node_id
        DETACH DELETE n
        RETURN count(n) AS deleted_count
        """
        result = await _await_call(neo4j_client.run_cypher, query, {"node_id": node_id})
        
        if result and result[0]["deleted_count"] > 0:
            return {"ok": True, "message": "节点删除成功"}
        return {"ok": False, "error": "节点不存在"}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/nodes/{node_id}", "delete_node", e, node_id=node_id)
        return {"ok": False, "error": str(e)}


@app.post("/admin/graph/relationships")
async def create_graph_relationship(data: dict):
    """创建关系（需要登录）。

    Args:
        data: 请求体数据字典。

    Returns:
        Any: 操作结果。
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        source_id = data.get("source_id")
        target_id = data.get("target_id")
        raw_rel_type = str(data.get("rel_type", "")).strip()
        rel_type = _validate_cypher_identifier(raw_rel_type)
        properties = _sanitize_graph_properties(data.get("properties", {}))
        
        if source_id is None or target_id is None or not raw_rel_type:
            return {"ok": False, "error": "源节点、目标节点和关系类型不能为空"}
        if not rel_type:
            return {"ok": False, "error": "关系类型非法，仅允许字母数字下划线且不能以数字开头"}

        try:
            source_id_int = int(source_id)
            target_id_int = int(target_id)
        except (TypeError, ValueError):
            return {"ok": False, "error": "源节点和目标节点ID必须是整数"}
        
        # 构建属性字符串
        if properties:
            props_str = "{ " + ", ".join([f"`{k}`: ${k}" for k in properties.keys()]) + " }"
            query = f"""
            MATCH (source), (target)
            WHERE id(source) = $source_id AND id(target) = $target_id
            CREATE (source)-[r:`{rel_type}` {props_str}]->(target)
            RETURN id(r) AS rel_id
            """
            params = {"source_id": source_id_int, "target_id": target_id_int, **properties}
        else:
            query = f"""
            MATCH (source), (target)
            WHERE id(source) = $source_id AND id(target) = $target_id
            CREATE (source)-[r:`{rel_type}`]->(target)
            RETURN id(r) AS rel_id
            """
            params = {"source_id": source_id_int, "target_id": target_id_int}
        
        result = await _await_call(neo4j_client.run_cypher, query, params)
        
        if result:
            return {"ok": True, "data": {"rel_id": result[0]["rel_id"], "message": "关系创建成功"}}
        return {"ok": False, "error": "创建失败，请检查节点是否存在"}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/relationships", "create_relationship", e, rel_type=rel_type)
        return {"ok": False, "error": str(e)}


@app.delete("/admin/graph/relationships/{rel_id}")
async def delete_graph_relationship(rel_id: int):
    """删除关系（需要登录）。

    Args:
        rel_id: 关系 ID。

    Returns:
        Any: 操作结果。
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        query = """
        MATCH ()-[r]->()
        WHERE id(r) = $rel_id
        DELETE r
        RETURN count(r) AS deleted_count
        """
        result = await _await_call(neo4j_client.run_cypher, query, {"rel_id": rel_id})
        
        if result and result[0]["deleted_count"] > 0:
            return {"ok": True, "message": "关系删除成功"}
        return {"ok": False, "error": "关系不存在"}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/relationships/{rel_id}", "delete_relationship", e, rel_id=rel_id)
        return {"ok": False, "error": str(e)}


@app.post("/admin/graph/execute-cypher")
async def execute_cypher(data: dict):
    """执行 Cypher 查询（只读，需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        query = data.get("query", "").strip()
        
        if not query:
            return {"ok": False, "error": "查询语句不能为空"}
        
        # 安全检查：只允许只读查询
        forbidden_keywords = ["create", "merge", "delete", "remove", "set", "drop", "alter"]
        query_lower = query.lower()
        for keyword in forbidden_keywords:
            if keyword in query_lower:
                return {"ok": False, "error": f"出于安全考虑，不允许执行包含 '{keyword}' 的操作"}
        
        result = await _await_call(neo4j_client.run_cypher, query)
        return {"ok": True, "data": result}
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/execute-cypher", "execute_cypher", e)
        return {"ok": False, "error": str(e)}


@app.post("/admin/graph/extract-and-import")
async def extract_and_import_graph(
    file: UploadFile = File(...),
    max_chars: int = Form(12000)
):
    """从上传文件抽取并导入图谱数据。

    Args:
        file: 上传文件对象。
        max_chars: 最大处理字符数。

    Returns:
        Dict[str, Any] | JSONResponse: 导入执行结果。
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})

    filename = file.filename or "uploaded_file"
    suffix = Path(filename).suffix.lower()
    supported_suffixes = {".txt", ".md", ".pdf", ".docx", ".doc", ".pptx", ".ppt", ".html", ".htm"}
    if suffix not in supported_suffixes:
        return {
            "ok": False,
            "error": f"不支持的文件格式: {suffix}，仅支持 {', '.join(sorted(supported_suffixes))}"
        }
    chunk_limit = 20000

    def split_text_by_sentence(text: str, limit: int = 20000):
        """按句子优先策略切分长文本。

        Args:
            text: 输入文本。
            limit: 结果数量上限。

        Returns:
            list[str]: 切分后的文本片段。
        """
        chunks = []
        remaining = (text or "").strip()
        sentence_ends = ["。", ".", "!", "！", "?", "？"]

        while remaining:
            if len(remaining) <= limit:
                chunks.append(remaining)
                break

            window = remaining[:limit]
            split_at = max(window.rfind(mark) for mark in sentence_ends)

            # 找不到句号时，退化为按固定长度切分，避免死循环
            if split_at < 0:
                split_at = limit
            else:
                split_at += 1

            chunk = remaining[:split_at].strip()
            if not chunk:
                chunk = remaining[:limit]
                split_at = limit

            chunks.append(chunk)
            remaining = remaining[split_at:].strip()

        return chunks

    def extract_text_fallback(file_bytes: bytes, filename: str, suffix: str) -> str:
        """在 Docling 不可用时执行格式回退解析。

        Args:
            file_bytes: 上传文件字节内容。
            filename: 原始文件名。
            suffix: 文件扩展名（小写）。

        Returns:
            str: 回退解析得到的文本。
        """
        if suffix == ".pdf":
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            try:
                return "\n".join(page.get_text("text") for page in doc)
            finally:
                doc.close()

        if suffix == ".docx":
            from io import BytesIO
            from docx import Document

            doc = Document(BytesIO(file_bytes))
            return "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text and p.text.strip())

        if suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            html = file_bytes.decode("utf-8", errors="replace")
            return BeautifulSoup(html, "html.parser").get_text(separator="\n", strip=True)

        raise RuntimeError(f"当前环境无法使用 Docling，且该格式无内置回退解析器: {suffix}")

    try:
        with log_duration("graph.extract_import.read_upload", filename=filename):
            file_bytes = await file.read()
        if not file_bytes:
            return {"ok": False, "error": "上传文件为空"}

        parser_used = "plain_text"
        parse_warning = ""

        if suffix in {".txt", ".md"}:
            with log_duration("graph.extract_import.parse_text", filename=filename):
                text = file_bytes.decode("utf-8", errors="replace")
        else:
            parsed = None
            try:
                from app.document_parser.docling_parser import DoclingParser

                parser = DoclingParser(enable_image_ocr=True)
                with log_duration("graph.extract_import.parse_docling", filename=filename):
                    parsed = await _await_call(parser.parse_bytes, file_bytes, filename)
                text = parsed.text
                parser_used = "docling"
            except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as docling_error:
                _log_api_error("/admin/graph/extract-and-import", "docling_parse_fallback", docling_error, filename=filename)
                with log_duration("graph.extract_import.parse_fallback", filename=filename, suffix=suffix):
                    text = await _await_call(extract_text_fallback, file_bytes, filename, suffix)
                parser_used = "fallback"
                parse_warning = f"Docling 不可用，已自动回退解析: {docling_error}"

        if not text or not text.strip():
            return {"ok": False, "error": "文件解析后文本为空，无法抽取图谱"}

        from app.knowledge_graph.extract_graph_data_utils import extract_text_to_graph_queries
        with log_duration("graph.extract_import.split_chunks", filename=filename, text_len=len(text), chunk_limit=chunk_limit):
            chunks = split_text_by_sentence(text, chunk_limit)

        total_entities = 0
        total_relations = 0
        total_node_batches = 0
        total_relation_batches = 0

        for idx, chunk in enumerate(chunks, start=1):
            with log_duration("graph.extract_import.extract_chunk", filename=filename, chunk_index=idx, chunk_len=len(chunk)):
                graph_data = await _await_call(extract_text_to_graph_queries, chunk, len(chunk))

            node_queries = graph_data["node_queries"]
            relation_queries = graph_data["relation_queries"]

            if node_queries:
                with log_duration("graph.extract_import.write_nodes", filename=filename, chunk_index=idx, batches=len(node_queries)):
                    await _await_call(neo4j_client.run_multiple_cypher, node_queries)
            if relation_queries:
                with log_duration("graph.extract_import.write_relations", filename=filename, chunk_index=idx, batches=len(relation_queries)):
                    await _await_call(neo4j_client.run_multiple_cypher, relation_queries)

            stats = graph_data["stats"]
            total_entities += stats["entities"]
            total_relations += stats["relations"]
            total_node_batches += len(node_queries)
            total_relation_batches += len(relation_queries)

        # 构建返回数据
        response_data = {
            "file_name": filename,
            "text_length": len(text),
            "parser_used": parser_used,
            "parse_warning": parse_warning,
            "chunk_limit": chunk_limit,
            "chunks": len(chunks),
            "entity_count": total_entities,
            "relation_count": total_relations,
            "node_batches": total_node_batches,
            "relation_batches": total_relation_batches,
            "message": f"抽取并导入完成，共分 {len(chunks)} 段，实体 {total_entities} 个，关系 {total_relations} 条"
        }

        # 如果使用了 Docling，添加解析详情
        if parsed and parser_used == "docling":
            response_data["parse_details"] = {
                "format": parsed.format,
                "pages": parsed.pages,
                "tables_count": len(parsed.tables),
                "images_count": len(parsed.images),
                "metadata": parsed.metadata,
                "text_preview": parsed.text[:2000] if parsed.text else ""  # 前2000字符预览
            }

        return {
            "ok": True,
            "data": response_data
        }
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        import traceback
        _log_api_error("/admin/graph/extract-and-import", "extract_import", e, filename=filename)
        return {"ok": False, "error": f"抽取写入失败: {str(e)}\n{traceback.format_exc()}"}


# 知识图谱嵌入 Milvus API
@app.post("/admin/graph/embed-to-milvus")
async def embed_graph_to_milvus(data: Optional[Dict[str, Any]] = None):
    """
    将知识图谱的实体和关系一键嵌入到 Milvus（需要登录）
    
    Args:
        data: 可选参数
            - labels: 指定要嵌入的节点标签列表，默认为所有标签
            - batch_size: 批处理大小，默认 100
    """
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        labels = data.get("labels") if data else None
        batch_size = data.get("batch_size", 100) if data else 100

        kg_store = await _await_call(get_shared_kg_milvus_store)
        stats = await _await_call(
            kg_store.rebuild_from_neo4j,
            neo4j_client,
            labels,
            batch_size,
        )

        return {
            "ok": True,
            "data": {
                "message": f"知识图谱嵌入成功，共 {stats['total_documents']} 条记录",
                "stats": stats,
                "timestamp": stats.get("timestamp")
            }
        }
        
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        import traceback
        _log_api_error("/admin/graph/embed-to-milvus", "embed_graph", e)
        return {"ok": False, "error": f"嵌入失败: {str(e)}\n{traceback.format_exc()}"}


@app.get("/admin/graph/embed-status")
async def get_graph_embed_status():
    """获取知识图谱嵌入状态（需要登录）"""
    if not is_logged_in():
        return JSONResponse(status_code=401, content={"ok": False, "error": "未登录"})
    
    try:
        kg_store = await _await_call(get_shared_kg_milvus_store)
        stats = await _await_call(kg_store.get_stats)
        return {
            "ok": True,
            "data": stats,
        }
            
    except (RuntimeError, ValueError, TypeError, KeyError, AttributeError, OSError, ImportError) as e:
        _log_api_error("/admin/graph/embed-status", "embed_status", e)
        return {"ok": False, "error": str(e)}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
